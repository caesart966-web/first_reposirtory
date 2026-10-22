"""И05.09.26-ОВ.ВР. Ведомость воздухораспределителей — данные для подбора решёток и диффузоров (А3).

Расходы — те же, что у решёток на листах 9, 10, 12 и в таблице подбора листа 10; сечения присоединений — по сетям
листов 9–11. Скорость в присоединении считается, живое сечение — по допустимой скорости в живом сечении.
python3 sheets/ov_vr.py -> out/ov_vr.dxf + PDF, out/И05.09.26-ОВ.ВР Ведомость воздухораспределителей.xlsx
"""
import json
import math
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)
from rdbase import new_doc, RDSheet, save_and_pdf, fmt_num, CODE  # noqa: E402

OUT = os.path.join(ROOT, 'out')
RD = json.load(open(os.path.join(ROOT, '..', 'rd', 'rd_calc.json')))
C = json.load(open(os.path.join(ROOT, '..', 'ovk', 'calc.json')))
ROOM = {r['no']: r['name'] for r in C['rooms']}
CODE_VR = CODE + '.ВР'
TITLE = 'Ведомость воздухораспределителей (данные для подбора)'
V_RES = 2.0             # м/с в живом сечении — жилые и вспомогательные помещения (как в таблице листа 10)
V_OUT = 2.0             # наружный воздухозабор — без захвата снега и дождя
NOISE_25 = ('2.3', '2.4', '2.5', '2.7', 'б5')


def f(v, d=1):
    return fmt_num(v, d)


def v_round(L, d):
    return L / 3600 / (math.pi * (d / 1000) ** 2 / 4)


def v_rect(L, a, b):
    return L / 3600 / (a / 1000 * b / 1000)


def rows():
    import ov_vent1 as v1
    import ov_vent2 as v2
    import ov_bath as ba
    AC = v1.AC
    brv = RD['boiler_room_vent']
    R = []

    def add(room, sysk, kind, typ, L, conn, vconn, vmax=V_RES, note='', Lmax=None):
        Lc = Lmax or L
        noise = 25 if room.split()[0] in NOISE_25 else 30
        if room.startswith('1.6') or room.startswith('снаружи'):
            noise = '—'
        R.append([room, sysk, kind, typ, f'{L}' + (f' ({Lmax})' if Lmax else ''), conn,
                  f(vconn, 1) if vconn is not None else '—', f(vmax, 1), f(Lc / 3600 / vmax, 3), str(noise), note])

    def rn(no):
        return f'{no} {ROOM[no]}'
    # ---------------------------------------------------------------- 1 этаж, П1 (лист 9)
    for x, nm in ((7.60, 'восток'), (4.60, 'центр'), (0.90, 'запад')):
        add(rn('1.10'), 'П1', 'приток', f'решётка (щелевой диффузор) в торце подшивки П-1, через КСД — {nm}', 70,
            'Ø125', v_round(70, 125))
    for nm in ('восток', 'запад'):
        add(rn('1.11'), 'П1', 'приток', f'диффузор потолочный — {nm}', 60, 'Ø100', v_round(60, 100))
    add(rn('1.2'), 'П1', 'приток', 'решётка в стене на КСД', 30, 'Ø100', v_round(30, 100))
    add(rn('1.10') + ' (кухня)', 'П1', 'вытяжка', 'решётка в стене на КСД', 90, 'Ø125', v_round(90, 125),
        note='зонт кухни — система В3 отдельно')
    add(rn('1.8'), 'П1', 'вытяжка', 'решётка в стене на КСД', 30, 'Ø100', v_round(30, 100))
    add(rn('1.4'), 'П1', 'вытяжка', 'решётка потолочная', 50, 'Ø100', v_round(50, 100))
    add(rn('1.3'), 'П1', 'вытяжка', 'решётка потолочная', 40, 'Ø100', v_round(40, 100))
    # ---------------------------------------------------------------- 2 этаж, П1 (лист 10)
    add(rn('2.2'), 'П1', 'приток', 'диффузор потолочный', 30, 'Ø100', v_round(30, 100))
    for no in ('2.3', '2.4', '2.5', '2.7'):
        add(rn(no), 'П1', 'приток (свежий)', 'в КСД щелевой решётки ЩП блока К (см. ниже)', 60, 'Ø100',
            v_round(60, 100), note='через ЩП вместе с воздухом блока К')
    add(rn('2.1'), 'П1', 'вытяжка', 'решётка потолочная', 90, 'Ø125', v_round(90, 125))
    add(rn('2.6'), 'П1', 'вытяжка', 'решётка потолочная', 60, 'Ø100', v_round(60, 100))
    add(rn('2.9'), 'П1', 'вытяжка', 'решётка потолочная', 75, 'Ø125', v_round(75, 125))
    add(rn('2.8'), 'П1', 'вытяжка', 'решётка потолочная', 75, 'Ø125', v_round(75, 125))
    add(rn('1.10') + ', верх 2-го света', 'П1', 'вытяжка', 'решётка в стене на КСД, низ +6,400', 120, 'Ø160',
        v_round(120, 160))
    # ---------------------------------------------------------------- кондиционирование (листы 9, 10)
    for n, ln in (('К-1', 1000), ('К-2', 1000), ('К-3', 1100)):
        a = AC[n]['air']
        add(rn(AC[n]['room']), n, 'подача', f'щелевой диффузор длиной {ln} мм в торце подшивки, через КСД', a,
            'КСД от фланца блока', None)
        add(rn(AC[n]['room']), n, 'возврат', 'решётка-люк 600×600 под блоком', a, 'блок', None)
    for n in ('К-4', 'К-5', 'К-6', 'К-7'):
        ls, lr = v2.slot_L('ЩП', n), v2.slot_L('ЩВ', n)
        add(rn(AC[n]['room']), n, 'подача', 'щелевая решётка ЩП 800×150 в стене, через КСД', ls,
            '250×150', v_rect(ls, 250, 150), note='блок К + свежий П1' if n != 'К-7' else '')
        add(rn(AC[n]['room']), n, 'возврат', 'щелевая решётка ЩВ 800×150 в стене, через КСД', lr,
            '250×150', v_rect(lr, 250, 150), note='с возвратом части П1' if n == 'К-7' else '')
    # ---------------------------------------------------------------- топочная, наружный воздух (листы 7, 9)
    add('1.6 Топочная', 'В2', 'вытяжка', 'решётка под потолком в западной стене → канал Ø250 в Ш-2', brv['L_exh'],
        'Ø250', v_round(brv['L_exh'], 250), vmax=1.0, note='естественная тяга — решётка с малым сопротивлением')
    add('1.6 Топочная', '—', 'приток', 'решётка в нижней части наружной двери', brv['L_in'], '—', None, vmax=1.0,
        note=f'живое сечение не менее {f(brv["A_in_free"], 3)} м² (лист 7)')
    add('снаружи, вост. стена 1.6', 'П1', 'наружный воздух', 'решётка воздухозабора Н1 500×400 с камерой', 700,
        'Ø250', v_round(700, 250), vmax=V_OUT, note='жалюзи от осадков, сетка')
    # ---------------------------------------------------------------- баня (лист 12)
    for k, fa in ba.FANS.items():
        add(f'{fa["room"]} {ROOM[fa["room"]]} (баня)', k, 'вытяжка', 'решётка под потолком → канальный вентилятор',
            fa['L'], 'Ø100', v_round(fa['Lmax'], 100), Lmax=fa['Lmax'],
            note='в скобках — форсированный режим; скорость — при нём')
    return R


HEAD = ['№', 'Помещение', 'Система', 'Назначение', 'Воздухораспределитель (по проекту)', 'L, м³/ч',
        'Присоеди-нение', 'v в присо-единении, м/с', 'v в живом сечении не более, м/с', 'Живое сечение не менее, м²',
        'Шум, дБА, не более', 'Примечание']
W = [8, 42, 14, 22, 92, 15, 22, 17, 19, 19, 14, 63]


def build():
    data = rows()
    doc = new_doc()
    per = [24, 30]
    pages, i = [], 0
    while i < len(data):
        n = per[0] if not pages else per[1]
        pages.append(data[i:i + n])
        i += n
    total = len(pages)
    k = 1
    for pi, chunk in enumerate(pages, 1):
        sh = RDSheet(doc, f'Лист {pi}', TITLE, pi, total, code=CODE_VR)
        tb = [HEAD] + [[str(k + j)] + r for j, r in enumerate(chunk)]
        k += len(chunk)
        sh.s.text(TITLE, 25 + sum(W) / 2, 288, h=3.0, align='center')
        yb = sh.table(25, 284, W, tb, h=1.8, heights=[12.0] + [7.6] * len(chunk),
                      aligns=['center', 'left', 'center', 'center', 'left', 'center', 'center', 'center', 'center',
                              'center', 'center', 'left'])
        if pi == total:
            sh.notes(25, yb - 5, 200, [
                'Расходы — по листам 9, 10, 12 и таблице подбора листа 10 (кондиционирование). Скорость в '
                'присоединении — по сечению воздуховода или камеры статического давления (КСД) перед решёткой.',
                f'Живое сечение решётки — по скорости в живом сечении: F = L / (3600 · v); для жилых и '
                f'вспомогательных помещений v не более {f(V_RES, 1)} м/с; в топочной (естественная вентиляция) — не '
                'более 1,0 м/с. При подборе по каталогу проверить уровень шума при расчётном расходе (не выше графы «Шум») '
                'и дальнобойность струи: подвижность воздуха в рабочей зоне — не более 0,2 м/с.',
                'Регулирование расходов — дроссель-клапанами и регуляторами расхода на ответвлениях (лист 11); '
                'решётки — с регулировочным клапаном или без него, по каталогу.',
                'Типоразмеры решёток подбираются позже по каталогу производителя по этим данным (решение Заказчицы '
                '22.10.2026).',
            ], h=2.0, spacing=1.0)
    return doc, data


def xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side, Font
    wb = Workbook()
    ws = wb.active
    ws.title = 'Воздухораспределители'
    ws.append([f'{TITLE} — {CODE_VR}'])
    ws['A1'].font = Font(bold=True)
    ws.append([h.replace('-', '') for h in HEAD])
    for j, r in enumerate(data, 1):
        ws.append([j] + r)
    thin = Side(style='thin')
    for col, w in zip('ABCDEFGHIJKL', [5, 30, 9, 14, 52, 10, 14, 12, 14, 14, 10, 40]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='center',
                                    horizontal='left' if c.column in (2, 5, 12) else 'center')
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    p = os.path.join(OUT, 'И05.09.26-ОВ.ВР Ведомость воздухораспределителей.xlsx')
    wb.save(p)
    return p


if __name__ == '__main__':
    doc, data = build()
    print(save_and_pdf(doc, os.path.join(OUT, 'ov_vr.dxf')))
    print(xlsx(data), len(data), 'строк')
