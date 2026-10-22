"""Печать спецификации (форма 1 ГОСТ 21.110-2013) в PDF А3 и XLSX — общая для ОВ.СО и ТМ.СО."""
import os
import subprocess
import sys

sys.path.insert(0, os.path.dirname(__file__))
from gost_html import CSS, side_strip, form5, form6  # noqa: E402


def esc(s):
    return str(s).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def render(rows, code, title, OUT, base, sheet_title='Спецификация', bld=None):
    """rows — как spec_data.numbered(); bld — наименование раздела в штампе (иначе из rd_stamp)."""
    import rd_stamp
    old = rd_stamp.BLD
    if bld:
        rd_stamp.BLD = bld
    try:
        W = [20, 130, 60, 35, 45, 20, 20, 25, 40]
        HEAD = ['Поз.', 'Наименование и техническая характеristика', 'Тип, марка, обозначение документа, опросного листа',
                'Код продукции', 'Поставщик', 'Ед. измерения', 'Количество', 'Масса 1 ед., кг', 'Примечание']
        HEAD[1] = 'Наименование и техническая характеристика'


        def esc(s):
            return str(s).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


        trs = []
        for r in rows:
            if r[0] == 'section':
                trs.append(f'<tr class="sec"><td></td><td class="l"><u>{esc(r[1])}</u></td><td></td><td></td><td></td>'
                           f'<td></td><td></td><td></td><td></td></tr>')
            else:
                _, n, name, typ, kod, sup, unit, qty, note = r
                trs.append(f'<tr><td>{n}</td><td class="l">{esc(name)}</td><td>{esc(typ)}</td><td>{esc(kod)}</td>'
                           f'<td>{esc(sup)}</td><td>{unit}</td><td>{qty}</td><td></td><td>{esc(note)}</td></tr>')

        colg = '<colgroup>' + ''.join(f'<col style="width:{w}mm">' for w in W) + '</colgroup>'
        thead = '<tr class="hd">' + ''.join(f'<td>{h}</td>' for h in HEAD) + '</tr>' + \
                '<tr class="nums">' + ''.join(f'<td>{i}</td>' for i in range(1, 10)) + '</tr>'
        note = ('Допускается применение оборудования, изделий и материалов других изготовителей с техническими '
                'характеристиками не ниже указанных в настоящей спецификации.')

        extra_css = """
        table.sp { position: absolute; left: 20mm; top: 5mm; border-collapse: collapse; table-layout: fixed;
          font: 9pt 'Times New Roman', 'Liberation Serif', serif; }
        table.sp td { border-left: 0.6mm solid #000; border-right: 0.6mm solid #000; padding: 0.6mm 1.2mm;
          text-align: center; vertical-align: middle; height: 8mm; line-height: 1.15 }
        table.sp td.l { text-align: left }
        table.sp tr.hd td { height: 15mm; border-bottom: 0.6mm solid #000; border-top: 0.6mm solid #000; font-size: 8.5pt }
        table.sp tr.nums td { height: 5mm; border-bottom: 0.6mm solid #000; font-size: 8pt }
        table.sp tr.item td, table.sp tr td { border-bottom: 0.25mm solid #000 }
        .pnote { position: absolute; left: 22mm; font: italic 9pt 'Times New Roman', 'Liberation Serif', serif; width: 200mm }
        """

        html = f"""<!doctype html><html lang="ru"><head><meta charset="utf-8"><style>{CSS % dict(w=420, h=297)}{extra_css}</style></head>
        <body><div id="src" style="display:none"><table class="sp">{colg}<tbody>{''.join(trs)}</tbody></table></div>
        <script>
        const thead = `{thead}`, colg = `{colg}`;
        const F5 = `{form5(title, '@S', '@N', code)}`, F6 = `{form6('@S', code)}`, SIDE = `{side_strip()}`;
        const NOTE = `{note}`;
        const MM = 96 / 25.4;
        const srcRows = Array.from(document.querySelectorAll('#src tr'));
        const pages = [];
        function newPage(first) {{
          const p = document.createElement('div'); p.className = 'page';
          p.innerHTML = '<div class="frame"></div>' + SIDE + '<table class="sp">' + colg + '<tbody>' + thead + '</tbody></table>' +
            (first ? F5 : F6) + '<div class="fmt">Формат А3</div>';
          document.body.appendChild(p); pages.push(p);
          return {{ p, tb: p.querySelector('table.sp tbody'), limit: (297 - 5 - (first ? 55 : 15) - 2) * MM }};
        }}
        let cur = newPage(true);
        for (let i = 0; i < srcRows.length; i++) {{
          const r = srcRows[i].cloneNode(true);
          cur.tb.appendChild(r);
          const bottom = r.getBoundingClientRect().bottom - cur.p.getBoundingClientRect().top;
          if (bottom > cur.limit) {{
            cur.tb.removeChild(r);
            // заголовок раздела не оставляем последней строкой листа
            const last = cur.tb.lastElementChild;
            const moveSec = last && last.classList.contains('sec');
            if (moveSec) cur.tb.removeChild(last);
            cur = newPage(false);
            if (moveSec) cur.tb.appendChild(last);
            cur.tb.appendChild(r);
          }}
        }}
        // примечание под таблицей последнего листа
        const t = cur.p.querySelector('table.sp');
        const y = (t.getBoundingClientRect().bottom - cur.p.getBoundingClientRect().top) / MM + 4;
        const n = document.createElement('div'); n.className = 'pnote'; n.style.top = y + 'mm';
        n.textContent = 'Примечание — ' + NOTE; cur.p.appendChild(n);
        // дорисовать пустые строки до низа листа (форма заполняется по высоте)
        pages.forEach((pg, k) => {{
          const tb = pg.querySelector('table.sp tbody');
          const lim = (297 - 5 - (k === 0 ? 55 : 15)) * MM;
          if (k === pages.length - 1) return;
          for (let j = 0; j < 40; j++) {{
            const e = document.createElement('tr'); e.innerHTML = '<td></td>'.repeat(9); tb.appendChild(e);
            if (e.getBoundingClientRect().bottom - pg.getBoundingClientRect().top > lim) {{ tb.removeChild(e); break; }}
          }}
        }});
        pages.forEach((pg, k) => {{ pg.innerHTML = pg.innerHTML.replaceAll('@S', String(k + 1)).replaceAll('@N', String(pages.length)); }});
        document.getElementById('src').remove();
        document.title = 'ready';
        </script></body></html>"""
        hp = os.path.join(OUT, base + '.html')
        open(hp, 'w').write(html)
        pp = hp[:-5] + '.pdf'
        env = dict(os.environ, NODE_PATH=subprocess.check_output(['npm', 'root', '-g'], text=True).strip())
        subprocess.check_call(['node', os.path.join(os.path.dirname(__file__), 'html2pdf.js'), hp, pp, '420', '297'], env=env)

        # XLSX
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Side, Font
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_title
        ws.append([title + ' — ' + code])
        ws.append(HEAD)
        thin = Side(style='thin')
        for r in rows:
            if r[0] == 'section':
                ws.append(['', r[1]])
                ws.cell(ws.max_row, 2).font = Font(underline='single', bold=True)
            else:
                _, n, name, typ, kod, sup, unit, qty, note_ = r
                ws.append([n, name, typ, kod, sup, unit, qty, '', note_])
        ws.append([])
        ws.append(['', 'Примечание — ' + note])
        for col, w in zip('ABCDEFGHI', [6, 70, 22, 12, 14, 9, 9, 9, 16]):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row - 2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left' if c.column == 2 else 'center')
                c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        wb.save(os.path.join(OUT, base + '.xlsx'))

    finally:
        rd_stamp.BLD = old
    return os.path.join(OUT, base + '.pdf'), os.path.join(OUT, base + '.xlsx')
