#!/usr/bin/env python3
"""Строительные компании Москвы с ОКВЭД 41–43 и членством в СРО — в один Excel.

Откуда данные
-------------
* ``mosstroybase.sqlite3`` — база, собранная командой ``python -m mosstroybase
  build`` из Единого реестра субъектов МСП ФНС: ИНН, ОГРН, название, код
  региона юридического адреса, ОКВЭД. Код 77 — город Москва, включая
  Зеленоград и Новую Москву. Московская область — код 50 — сюда не попадает.
* Реестр НОСТРОЙ (reestr.nostroy.ru) — членство в строительных СРО.

Почему реестр НОСТРОЙ качается целиком
--------------------------------------
Московских строительных компаний десятки тысяч. Проверка по одному ИНН —
около секунды на компанию, то есть сутки. Весь реестр — это около 840
запросов по 500 записей, полчаса. Сопоставление потом идёт локально.

Цена такого пути — риск молча недокачать: так уже было с НОПРИЗом, когда
постраничная выгрузка отдала 147 тысяч записей из 212 и остановилась.
Поэтому выгрузка сверяется трижды, и результат сверки попадает в таблицу:

1. сколько записей получено против заявленного реестром;
2. компании, которые раньше проверялись по одному ИНН и числились членами,
   должны найтись и в выгрузке;
3. случайные компании со статусом «Нет» перепроверяются поиском по ИНН.

Команды
-------
    python москва_стройка.py всё          всё по очереди (так и запускать)
    python москва_стройка.py нострой      выкачать реестр НОСТРОЙ
    python москва_стройка.py выборка      перепроверить случайные «Нет» по ИНН
    python москва_стройка.py таблица      собрать Excel

Выгрузку можно прервать Ctrl+C и продолжить той же командой.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import random
import re
import sqlite3
import sys
import time
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any, Iterable

try:
    import requests
except ImportError:  # pragma: no cover
    requests = None

BASE = "https://reestr.nostroy.ru"
LIST_URL = f"{BASE}/api/sro/all/member/list"

ДАМП = "нострой_реестр.jsonl"
ВЫБОРКА = "нострой_выборка.json"

РАЗДЕЛЫ = {
    "41": "Строительство зданий",
    "42": "Строительство инженерных сооружений",
    "43": "Работы строительные специализированные",
}

# Наименования по ОКВЭД 2 (ОК 029-2014) до подкласса. Код длиннее, чем есть
# в словаре, расшифровывается по ближайшему старшему.
ОКВЭД = {
    "41.1": "Разработка строительных проектов",
    "41.10": "Разработка строительных проектов",
    "41.2": "Строительство жилых и нежилых зданий",
    "41.20": "Строительство жилых и нежилых зданий",
    "42.1": "Строительство автомобильных и железных дорог",
    "42.11": "Строительство автомобильных дорог и автомагистралей",
    "42.12": "Строительство железных дорог и метро",
    "42.13": "Строительство мостов и тоннелей",
    "42.2": "Строительство инженерных коммуникаций",
    "42.21": "Строительство инженерных коммуникаций для водоснабжения, "
             "водоотведения и газоснабжения",
    "42.22": "Строительство коммунальных объектов для обеспечения "
             "электроэнергией и телекоммуникациями",
    "42.9": "Строительство прочих инженерных сооружений",
    "42.91": "Строительство водных сооружений",
    "42.99": "Строительство прочих инженерных сооружений, не включённых "
             "в другие группировки",
    "43.1": "Разборка и снос зданий, подготовка строительного участка",
    "43.11": "Разборка и снос зданий",
    "43.12": "Подготовка строительной площадки",
    "43.13": "Разведочное бурение",
    "43.2": "Электромонтажные, санитарно-технические и прочие "
            "строительно-монтажные работы",
    "43.21": "Производство электромонтажных работ",
    "43.22": "Производство санитарно-технических работ, монтаж отопительных "
             "систем и систем кондиционирования воздуха",
    "43.29": "Производство прочих строительно-монтажных работ",
    "43.3": "Работы строительные отделочные",
    "43.31": "Производство штукатурных работ",
    "43.32": "Работы столярные и плотничные",
    "43.33": "Работы по устройству покрытий полов и облицовке стен",
    "43.34": "Производство малярных и стекольных работ",
    "43.39": "Производство прочих отделочных и завершающих работ",
    "43.9": "Работы строительные специализированные прочие",
    "43.91": "Производство кровельных работ",
    "43.99": "Работы строительные специализированные прочие, не включённые "
             "в другие группировки",
}

# В базе категория хранится словом («микро»), в сыром реестре — цифрой
КАТЕГОРИИ_МСП = {"1": "Микропредприятие", "2": "Малое предприятие",
                 "3": "Среднее предприятие", "микро": "Микропредприятие",
                 "малое": "Малое предприятие", "среднее": "Среднее предприятие"}


# -- мелочи -------------------------------------------------------------


def тыс(n: int) -> str:
    """12 345 — разделитель тысяч пробелом, как принято в русском тексте."""
    return f"{n:,}".replace(",", "\u00a0")


def норм_инн(value: Any) -> str:
    """ИНН к сравнимому виду: только цифры и восстановленный ведущий ноль."""
    digits = re.sub(r"\D", "", "" if value is None else str(value))
    if len(digits) == 9:
        return digits.zfill(10)
    if len(digits) == 11:
        return digits.zfill(12)
    return digits


def разобрать_дату(text: Any) -> date | None:
    text = str(text or "").strip()
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", text)
    if m:
        year, month, day = map(int, m.groups())
    else:
        m = re.match(r"(\d{2})\.(\d{2})\.(\d{4})", text)
        if not m:
            return None
        day, month, year = map(int, m.groups())
    try:
        return date(year, month, day)
    except ValueError:
        return None


def раздел_оквэд(code: str | None) -> str | None:
    """«41», «42» или «43» — если код строительный, иначе None."""
    head = (code or "").strip()[:2]
    return head if head in РАЗДЕЛЫ else None


def название_оквэд(code: str | None) -> str:
    code = (code or "").strip()
    while code:
        if code in ОКВЭД:
            return ОКВЭД[code]
        code = code[:-1].rstrip(".")
    return ""


def москва_по_адресу(address: str | None) -> bool:
    """Не противоречит ли свежий адрес коду региона 77.

    Код региона приходит из реестра МСП и обновляется раз в месяц, а адрес
    в базе бывает свежее — его приносит Checko. Переехавшую в область
    компанию по коду было бы не отличить. Пустой адрес ничему не противоречит.
    """
    text = (address or "").lower().replace("ё", "е")
    if not text.strip():
        return True
    if re.search(r"московск\w*\s+обл|обл\w*\.?\s+московск", text):
        return False
    if re.search(r"\bмосква\b|\bг\.?\s*москва\b|\bмосквы\b", text):
        return True
    # другой субъект назван явно — а Москва нет
    return not re.search(r"\bобл\b|област|\bкрай\b|\bресп\b|республик|автономн", text)


# -- реестр НОСТРОЙ -----------------------------------------------------


def тело(page: int, size: int, search: str = "") -> dict:
    """Тело запроса в том виде, как его шлёт сам сайт реестра.

    Размер страницы — pageCount и строкой. Прежний pageSize реестр молча
    игнорировал и отдавал по 20 записей.
    """
    return {"filters": {}, "page": page, "pageCount": str(size),
            "searchString": search, "sortBy": {}}


def пауза(status: int | None, attempt: int) -> float:
    """500 у реестра случайная — повтор через секунду. 429/503 — просьба
    сбавить темп, тут ждём по-настоящему. Обрыв сети — тоже."""
    if status in (429, 503):
        return min(2 ** attempt, 30)
    if status is not None and 500 <= status < 600:
        return min(attempt, 4)
    return min(2 ** attempt, 20)


def сессия():
    s = requests.Session()
    s.headers.update({"User-Agent": "mosstroybase/0.1 (+open-data harvester)",
                      "Content-Type": "application/json",
                      "Accept": "application/json"})
    return s


def запрос(session, body: dict, что: str, attempts: int = 5,
           timeout: int = 60) -> dict | None:
    for attempt in range(1, attempts + 1):
        status = None
        try:
            resp = session.post(LIST_URL, json=body, timeout=timeout)
            if resp.status_code == 200:
                return resp.json()
            status = resp.status_code
            print(f"[нострой] {что}: HTTP {status} (попытка {attempt}/{attempts})",
                  flush=True)
        except (requests.RequestException, ValueError) as exc:
            print(f"[нострой] {что}: {type(exc).__name__} "
                  f"(попытка {attempt}/{attempts})", flush=True)
        time.sleep(пауза(status, attempt))
    return None


def записи_ответа(payload: Any) -> list[dict]:
    if not isinstance(payload, dict):
        return []
    data = payload.get("data")
    if isinstance(data, dict):
        data = data.get("data")
    if isinstance(data, list):
        return [r for r in data if isinstance(r, dict)]
    return []


def всего_в_ответе(payload: Any) -> int | None:
    if not isinstance(payload, dict):
        return None
    data = payload.get("data")
    for container in ([data, payload] if isinstance(data, dict) else [payload]):
        for key in ("count", "total", "totalCount"):
            value = container.get(key)
            if isinstance(value, int):
                return value
    return None


def _строка(obj: dict, *keys: str) -> str:
    for key in keys:
        value = obj.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


_ПРИЗНАКИ_БЫВШЕГО = ("исключ", "прекращ")
_ПРИЗНАКИ_ЧЛЕНА = ("является членом", "действ", "член ")
# Имена полей подсмотрены в сыром ответе реестра: дата вступления —
# registry_registration_date, дата выхода — suspension_date (у записи
# со статусом «Исключен» других дат нет вовсе)
_ПОЛЯ_ВСТУПЛЕНИЯ = ("registry_registration_date", "member_right_start_date",
                    "right_start_date", "start_date")
_ПОЛЯ_ВЫХОДА = ("suspension_date", "member_right_stop_date", "right_stop_date",
                "stop_date", "exclusion_date")


def статус_члена(record: dict) -> str:
    status = record.get("member_status") or record.get("status")
    if isinstance(status, dict):
        return _строка(status, "title", "name")
    return status.strip() if isinstance(status, str) else ""


def _первая_дата(record: dict, поля: Iterable[str]) -> date | None:
    for поле in поля:
        d = разобрать_дату(record.get(поле))
        if d:
            return d
    return None


def бывший(record: dict) -> bool:
    """Решает текст статуса: у действующих дата окончания права бывает
    плановой, в будущем, и считать её признаком выхода нельзя."""
    status = статус_члена(record).lower()
    if any(m in status for m in _ПРИЗНАКИ_БЫВШЕГО):
        return True
    if any(m in status for m in _ПРИЗНАКИ_ЧЛЕНА):
        return False
    stop = _первая_дата(record, _ПОЛЯ_ВЫХОДА)
    return bool(stop and stop <= date.today())


def сжать(record: dict) -> dict:
    """Из записи реестра — только то, что нужно таблице."""
    sro = record.get("sro") if isinstance(record.get("sro"), dict) else {}
    start = _первая_дата(record, _ПОЛЯ_ВСТУПЛЕНИЯ)
    stop = _первая_дата(record, _ПОЛЯ_ВЫХОДА)
    return {
        "id": record.get("id"),
        "inn": норм_инн(record.get("inn")),
        "status": статус_члена(record),
        "former": бывший(record),
        "sro": _строка(sro, "title", "full_description", "short_description", "name"),
        "sro_num": _строка(sro, "registration_number"),
        "start": start.isoformat() if start else None,
        "stop": stop.isoformat() if stop else None,
    }


def ключ_записи(r: dict) -> str:
    if r.get("id") is not None:
        return f"id:{r['id']}"
    return f"{r['inn']}|{r['sro_num']}|{r['start']}|{r['status']}"


def свести(записи: list[dict]) -> dict:
    """Членства одной компании — в одну строку таблицы.

    Действующее важнее любого прошлого. У действующего берётся самое раннее
    вступление, у прекращённого — самый поздний выход: от него считается
    годичный запрет на вступление в другую СРО.
    """
    действующие = [r for r in записи if not r["former"]]
    прошлые = [r for r in записи if r["former"]]
    if действующие:
        главная = min(действующие, key=lambda r: r["start"] or "9999")
        имена = list(dict.fromkeys(r["sro"] for r in действующие if r["sro"]))
        return {"В СРО": "Да", "СРО": "; ".join(имена),
                "Рег. номер СРО": главная["sro_num"],
                "В СРО с": главная["start"], "Исключена": None}
    if прошлые:
        последняя = max(прошлые, key=lambda r: r["stop"] or r["start"] or "")
        return {"В СРО": "Исключена", "СРО": последняя["sro"],
                "Рег. номер СРО": последняя["sro_num"],
                "В СРО с": последняя["start"], "Исключена": последняя["stop"]}
    return {"В СРО": "Нет", "СРО": "", "Рег. номер СРО": "", "В СРО с": None,
            "Исключена": None}


# -- команда «нострой» ---------------------------------------------------


def _состояние(дамп: Path) -> Path:
    return дамп.with_name(дамп.name + ".state.json")


def cmd_нострой(args) -> int:
    дамп = Path(args.дамп)
    state_path = _состояние(дамп)
    state = {}
    if state_path.exists() and not args.заново:
        state = json.loads(state_path.read_text(encoding="utf-8"))
        if state.get("готово"):
            print(f"[нострой] выгрузка уже есть: {дамп} ({state.get('получено')} "
                  f"записей от {state.get('дата')}). Перекачать: --заново")
            return 0
    if args.заново or not state:
        дамп.write_text("", encoding="utf-8")
        state = {}

    session = сессия()
    размер = state.get("размер") or args.размер
    if not state:
        print("[нострой] спрашиваю реестр, сколько в нём записей …", flush=True)
        первая = запрос(session, тело(1, args.размер), "страница 1")
        if первая is None:
            print("[нострой] реестр не ответил. Проверьте интернет и повторите.",
                  file=sys.stderr)
            return 1
        всего = всего_в_ответе(первая)
        получено = записи_ответа(первая)
        if not получено or not всего:
            print("[нострой] реестр ответил без записей — сырой ответ:",
                  json.dumps(первая, ensure_ascii=False)[:1500], file=sys.stderr)
            return 1
        # Реестр вправе урезать страницу — тогда число страниц считаем
        # по фактическому размеру. На этом уже ошибались: считали по
        # запрошенному и занижали число страниц впятеро
        размер = len(получено) if len(получено) < args.размер else args.размер
        state = {"всего": всего, "размер": размер, "страница": 1,
                 "дата": date.today().isoformat(), "провалы": []}
        with дамп.open("a", encoding="utf-8") as f:
            for r in получено:
                f.write(json.dumps(сжать(r), ensure_ascii=False) + "\n")
        state_path.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")

    всего, размер = state["всего"], state["размер"]
    страниц = math.ceil(всего / размер)
    print(f"[нострой] в реестре {тыс(всего)} записей, по {размер} на странице — "
          f"{тыс(страниц)} страниц", flush=True)

    t0, сделано = time.time(), 0
    пустых_подряд = 0
    страница = state["страница"] + 1
    try:
        while страница <= страниц + 2:
            payload = запрос(session, тело(страница, размер), f"страница {страница}")
            if payload is None:
                state["провалы"].append(страница)
                print(f"[нострой] страница {страница} не пришла — вернусь к ней в конце",
                      file=sys.stderr)
            else:
                recs = записи_ответа(payload)
                if not recs and страница <= страниц:
                    # Пустая страница посреди реестра — не конец, а сбой.
                    # На этом НОПРИЗ и потерял треть записей
                    for _ in range(3):
                        time.sleep(3)
                        payload = запрос(session, тело(страница, размер),
                                         f"страница {страница} (пустая, ещё раз)")
                        recs = записи_ответа(payload)
                        if recs:
                            break
                if recs:
                    пустых_подряд = 0
                    with дамп.open("a", encoding="utf-8") as f:
                        for r in recs:
                            f.write(json.dumps(сжать(r), ensure_ascii=False) + "\n")
                else:
                    пустых_подряд += 1
                    if страница <= страниц:
                        state["провалы"].append(страница)
                    if пустых_подряд >= 5:
                        break
            state["страница"] = страница
            state_path.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")
            сделано += 1
            if страница % 20 == 0 or страница == страниц:
                темп = (time.time() - t0) / сделано
                осталось = max(страниц - страница, 0) * темп / 60
                print(f"[нострой] страница {страница}/{страниц}, "
                      f"осталось ~{осталось:.0f} мин", flush=True)
            страница += 1
            time.sleep(args.пауза)

        # вторая попытка по провалившимся страницам
        провалы = sorted(set(state["провалы"]))
        state["провалы"] = []
        for стр in провалы:
            payload = запрос(session, тело(стр, размер), f"страница {стр} (повтор)")
            recs = записи_ответа(payload)
            if recs:
                with дамп.open("a", encoding="utf-8") as f:
                    for r in recs:
                        f.write(json.dumps(сжать(r), ensure_ascii=False) + "\n")
            else:
                state["провалы"].append(стр)
            time.sleep(args.пауза)
    except KeyboardInterrupt:
        state_path.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")
        print("\n[нострой] прервано — продолжить: та же команда", file=sys.stderr)
        return 1

    уникальных = len({ключ_записи(r) for r in читать_дамп(дамп)})
    state.update(готово=True, получено=уникальных, дата=date.today().isoformat())
    state_path.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")
    доля = уникальных / всего * 100
    print(f"[нострой] получено {тыс(уникальных)} записей из заявленных {тыс(всего)} "
          f"({доля:.1f} %)")
    if state["провалы"]:
        print(f"[нострой] ВНИМАНИЕ: не пришли страницы {state['провалы'][:20]}"
              f"{' …' if len(state['провалы']) > 20 else ''}", file=sys.stderr)
    if доля < 99:
        print("[нострой] ВНИМАНИЕ: выгрузка неполная. Колонка «В СРО» может "
              "ошибаться в сторону «Нет» — насколько, покажет «выборка».",
              file=sys.stderr)
    return 0


def читать_дамп(дамп: Path) -> Iterable[dict]:
    with дамп.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                yield json.loads(line)


def членства_по_инн(дамп: Path, нужные: set[str] | None = None
                    ) -> tuple[dict[str, list[dict]], int]:
    """ИНН -> записи реестра. Дубли (от продолжения после обрыва) схлопываются."""
    по_инн: dict[str, list[dict]] = defaultdict(list)
    видели: set[str] = set()
    for r in читать_дамп(дамп):
        k = ключ_записи(r)
        if k in видели:
            continue
        видели.add(k)
        if нужные is None or r["inn"] in нужные:
            по_инн[r["inn"]].append(r)
    return по_инн, len(видели)


# -- база mosstroybase ---------------------------------------------------


def _json_list(value: Any) -> list:
    if isinstance(value, list):
        return value
    try:
        parsed = json.loads(value or "[]")
    except (TypeError, ValueError):
        return []
    return parsed if isinstance(parsed, list) else []


def читать_базу(путь: Path, с_ип: bool = False) -> tuple[list[dict], list[dict], Counter]:
    """(основной ОКВЭД 41–43, 41–43 только в дополнительных, счётчик отсеянного)."""
    con = sqlite3.connect(f"file:{путь}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    основные, доп, отсеяно = [], [], Counter()
    for row in con.execute("SELECT * FROM companies"):
        c = dict(row)
        if str(c.get("region_code") or "").strip() != "77":
            отсеяно["не Москва по коду региона"] += 1
            continue
        if (c.get("kind") or "ЮЛ") != "ЮЛ" and not с_ип:
            отсеяно["индивидуальные предприниматели"] += 1
            continue
        if c.get("is_active") == 0:
            отсеяно["ликвидированы или в ликвидации"] += 1
            continue
        if not москва_по_адресу(c.get("address")):
            отсеяно["свежий адрес — не Москва"] += 1
            continue
        c["okved_add"] = _json_list(c.get("okved_add"))
        c["phones"] = _json_list(c.get("phones"))
        c["emails"] = _json_list(c.get("emails"))
        c["sro_info"] = _json_list(c.get("sro_info"))
        c["inn"] = норм_инн(c.get("inn"))
        if раздел_оквэд(c.get("okved_main")):
            основные.append(c)
        elif any(раздел_оквэд(code) for code in c["okved_add"]):
            доп.append(c)
        else:
            отсеяно["нет ОКВЭД 41–43"] += 1
    con.close()
    return основные, доп, отсеяно


def дата_базы(путь: Path) -> str:
    try:
        con = sqlite3.connect(f"file:{путь}?mode=ro", uri=True)
        value = con.execute("SELECT MAX(updated_at) FROM companies").fetchone()[0]
        con.close()
        d = разобрать_дату(value)
        return f"{d:%d.%m.%Y}" if d else ""
    except sqlite3.Error:
        return ""


def был_членом_по_прежней_проверке(c: dict) -> bool:
    """Прошлая проверка по одному ИНН (enrich-sro / daily) нашла действующее
    членство. Это независимый от постраничной выгрузки способ — на нём
    и проверяется, не потеряла ли выгрузка записи."""
    return any("член строительной СРО" in str(x) for x in c.get("sro_info") or [])


# -- команда «выборка» ---------------------------------------------------


def cmd_выборка(args) -> int:
    дамп = Path(args.дамп)
    if not дамп.exists():
        print(f"Нет выгрузки {дамп} — сначала: python москва_стройка.py нострой",
              file=sys.stderr)
        return 1
    основные, _доп, _о = читать_базу(Path(args.db), args.с_ип)
    инн = {c["inn"] for c in основные}
    по_инн, _ = членства_по_инн(дамп, инн)
    нет = sorted(i for i in инн if свести(по_инн.get(i, []))["В СРО"] == "Нет")
    random.seed(args.зерно)
    выборка = random.sample(нет, min(args.n, len(нет)))
    print(f"[выборка] перепроверяю поиском по ИНН {len(выборка)} случайных "
          f"компаний со статусом «Нет» (из {len(нет)})", flush=True)

    session = сессия()
    нашлись, не_ответил = [], 0
    for n, i in enumerate(выборка, 1):
        payload = запрос(session, тело(1, 50, i), f"ИНН {i}")
        if payload is None:
            не_ответил += 1
            continue
        свои = [сжать(r) for r in записи_ответа(payload)
                if норм_инн(r.get("inn")) == i]
        итог = свести(свои)["В СРО"]
        if итог != "Нет":
            нашлись.append({"ИНН": i, "на самом деле": итог})
            print(f"[выборка] ИНН {i}: выгрузка сказала «Нет», поиск — «{итог}»",
                  flush=True)
        if n % 25 == 0:
            print(f"[выборка] {n}/{len(выборка)}", flush=True)
        time.sleep(args.пауза)

    проверено = len(выборка) - не_ответил
    итог = {"проверено": проверено, "ошибок": len(нашлись), "ошибки": нашлись,
            "дата": date.today().isoformat()}
    Path(args.результат).write_text(json.dumps(итог, ensure_ascii=False, indent=1),
                                    encoding="utf-8")
    print(f"[выборка] проверено {проверено}, выгрузка ошиблась в {len(нашлись)}")
    if нашлись:
        print("[выборка] ВНИМАНИЕ: выгрузка теряет членов СРО. Пришлите этот вывод.",
              file=sys.stderr)
    return 0


# -- команда «таблица» ---------------------------------------------------


КОЛОНКИ = [
    ("№", 7), ("Наименование", 44), ("ИНН", 13), ("ОГРН", 16),
    ("ОКВЭД", 9), ("Вид деятельности (основной ОКВЭД)", 46), ("Раздел", 30),
    ("В СРО", 11), ("СРО", 46), ("Рег. номер СРО", 22), ("В СРО с", 12),
    ("Исключена из СРО", 12), ("Размер (реестр МСП)", 20),
    ("Дата регистрации", 12), ("Руководитель", 30), ("Адрес", 50),
    ("Телефон", 18), ("E-mail", 28),
]
# Лист «доп.»: основной ОКВЭД у этих компаний не строительный, и его надо
# видеть — иначе магазин стройматериалов (46.73) с 43.21 в дополнительных
# выглядит электромонтажной фирмой. Так и было в первой версии: колонка
# называлась «основной ОКВЭД», а стоял в ней код из дополнительных
КОЛОНКИ_ДОП = [
    ("№", 7), ("Наименование", 44), ("ИНН", 13), ("ОГРН", 16),
    ("Основной ОКВЭД", 11), ("Строительный ОКВЭД (доп.)", 13),
    ("Строительный вид деятельности (доп.)", 46), ("Раздел", 30),
] + КОЛОНКИ[7:]
ДАТЫ = {"В СРО с", "Исключена из СРО", "Дата регистрации"}


def строка_таблицы(c: dict, членство: dict) -> dict:
    code = (c.get("okved_main") or "").strip()
    if not раздел_оквэд(code):
        # на листе «доп. ОКВЭД» показываем строительный код из дополнительных
        code = next((x for x in c["okved_add"] if раздел_оквэд(x)), code)
    раздел = раздел_оквэд(code)
    return {
        "Наименование": c.get("name_short") or c.get("name") or "",
        "ИНН": c["inn"], "ОГРН": c.get("ogrn") or "",
        "ОКВЭД": code, "Вид деятельности (основной ОКВЭД)": название_оквэд(code),
        "Основной ОКВЭД": (c.get("okved_main") or "").strip(),
        "Строительный ОКВЭД (доп.)": code,
        "Строительный вид деятельности (доп.)": название_оквэд(code),
        "Раздел": f"{раздел} — {РАЗДЕЛЫ[раздел]}" if раздел else "",
        "В СРО": членство["В СРО"], "СРО": членство["СРО"],
        "Рег. номер СРО": членство["Рег. номер СРО"],
        "В СРО с": разобрать_дату(членство["В СРО с"]),
        "Исключена из СРО": разобрать_дату(членство["Исключена"]),
        "Размер (реестр МСП)": КАТЕГОРИИ_МСП.get(
            str(c.get("msp_category") or "").strip().lower(), c.get("msp_category") or ""),
        "Дата регистрации": разобрать_дату(c.get("reg_date")),
        "Руководитель": c.get("director") or "",
        "Адрес": c.get("address") or "",
        "Телефон": ", ".join(str(p) for p in c["phones"][:2]),
        "E-mail": ", ".join(str(e) for e in c["emails"][:2]),
    }


def _стили():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    return {
        "заголовок": Font(name="Calibri", size=16, bold=True, color="1F3864"),
        "подзаголовок": Font(name="Calibri", size=10, italic=True, color="595959"),
        "шапка_шрифт": Font(name="Calibri", size=10, bold=True, color="FFFFFF"),
        "шапка_заливка": PatternFill("solid", fgColor="1F3864"),
        "шапка_выравн": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "шапка_рамка": Border(bottom=Side(style="medium", color="FFFFFF")),
        "раздел": Font(name="Calibri", size=12, bold=True, color="1F3864"),
        "жирный": Font(name="Calibri", size=11, bold=True),
        "число": Font(name="Calibri", size=11, bold=True, color="1F3864"),
        "серый": Font(name="Calibri", size=10, color="595959"),
        "красный": Font(name="Calibri", size=11, bold=True, color="C00000"),
    }


def _лист_компаний(wb, имя: str, заголовок: str, пояснение: str,
                   строки: list[dict], стили: dict, колонки=None) -> None:
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    КОЛОНКИ_ЛИСТА = колонки or КОЛОНКИ
    ws = wb.create_sheet(имя)
    for n, (_к, ширина) in enumerate(КОЛОНКИ_ЛИСТА, 1):
        ws.column_dimensions[get_column_letter(n)].width = ширина
    ws.freeze_panes = "C4"   # шапка и название видны при прокрутке

    a1 = WriteOnlyCell(ws, value=заголовок); a1.font = стили["заголовок"]
    ws.append([a1])
    a2 = WriteOnlyCell(ws, value=пояснение); a2.font = стили["подзаголовок"]
    ws.append([a2])
    шапка = []
    for к, _ш in КОЛОНКИ_ЛИСТА:
        cell = WriteOnlyCell(ws, value=к)
        cell.font, cell.fill = стили["шапка_шрифт"], стили["шапка_заливка"]
        cell.alignment, cell.border = стили["шапка_выравн"], стили["шапка_рамка"]
        шапка.append(cell)
    ws.append(шапка)
    ws.row_dimensions[3].height = 32

    for n, s in enumerate(строки, 1):
        ряд = []
        for к, _ш in КОЛОНКИ_ЛИСТА:
            value = n if к == "№" else s.get(к)
            if к in ДАТЫ and value:
                cell = WriteOnlyCell(ws, value=value)
                cell.number_format = "DD.MM.YYYY"
                ряд.append(cell)
            else:
                ряд.append(value if value not in (None, "") else None)
        ws.append(ряд)

    последняя = 3 + len(строки)
    if строки:
        буквы = get_column_letter(len(КОЛОНКИ_ЛИСТА))
        ws.auto_filter.ref = f"A3:{буквы}{последняя}"
        имена = [к for к, _ in КОЛОНКИ_ЛИСТА]
        статус = next((к for к in ("В СРО", "В строительной СРО") if к in имена), None)
        кол_сро = get_column_letter(имена.index(статус) + 1) if статус else None
        диапазон = f"{кол_сро}4:{кол_сро}{последняя}"
        for значение, фон, цвет in ((("Да", "C6EFCE", "006100"),
                                     ("Нет", "FFC7CE", "9C0006"),
                                     ("Исключена", "FFEB9C", "9C5700")) if статус else ()):
            ws.conditional_formatting.add(диапазон, FormulaRule(
                formula=[f'${кол_сро}4="{значение}"'],
                fill=PatternFill("solid", fgColor=фон),
                font=Font(bold=True, color=цвет), stopIfTrue=True))
        ws.conditional_formatting.add(f"A4:{буквы}{последняя}", FormulaRule(
            formula=["MOD(ROW(),2)=0"], fill=PatternFill("solid", fgColor="F3F6FB")))


def _лист_сводки(wb, стили: dict, итоги: dict) -> None:
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, PatternFill
    ws = wb.create_sheet("Сводка", 0)
    for буква, ширина in zip("ABCDE", (92, 13, 13, 13, 13)):
        ws.column_dimensions[буква].width = ширина

    def ячейка(value, стиль=None, формат=None, заливка=None):
        c = WriteOnlyCell(ws, value=value)
        if стиль:
            c.font = стили[стиль]
        if формат:
            c.number_format = формат
        if заливка:
            c.fill = PatternFill("solid", fgColor=заливка)
        c.alignment = Alignment(vertical="center", wrap_text=isinstance(value, str)
                                and len(value) > 88)
        return c

    def строка(*vals):
        ws.append(list(vals))

    строка(ячейка("Строительные компании Москвы — ОКВЭД 41, 42, 43 и членство в СРО",
                  "заголовок"))
    строка(ячейка(итоги["подзаголовок"], "подзаголовок"))
    строка()

    всего = итоги["всего"]
    строка(ячейка("Итого", "раздел"))
    строка(ячейка("Компаний с основным ОКВЭД 41–43", "жирный"),
           ячейка(всего, "число", "# ##0"))
    for метка, ключ, фон in (("Состоят в строительной СРО", "Да", "C6EFCE"),
                             ("Не состоят в СРО", "Нет", "FFC7CE"),
                             ("Были в СРО, исключены", "Исключена", "FFEB9C"),
                             ("Не проверено", "Не проверено", "EDEDED")):
        n = итоги["по_сро"].get(ключ, 0)
        if ключ == "Не проверено" and not n:
            continue
        строка(ячейка(метка, заливка=фон), ячейка(n, "число", "# ##0"),
               ячейка(n / всего if всего else 0, "серый", "0.0%"))
    строка()

    for заголовок, ключ, подписи in (
            ("По видам деятельности", "по_разделам",
             {k: f"{k} — {v}" for k, v in РАЗДЕЛЫ.items()}),
            ("По размеру предприятия (реестр МСП)", "по_размеру", None)):
        строка(ячейка(заголовок, "раздел"))
        шапка = [ячейка(x, "шапка_шрифт", заливка="1F3864")
                 for x in ("", "Всего", "В СРО", "Не в СРО", "Исключены")]
        строка(*шапка)
        for группа, счёт in итоги[ключ].items():
            строка(ячейка(подписи.get(группа, группа) if подписи else группа or "не указан"),
                   *[ячейка(счёт.get(x, 0), формат="# ##0")
                     for x in ("всего", "Да", "Нет", "Исключена")])
        строка()

    строка(ячейка("Дополнительно", "раздел"))
    строка(ячейка("Строительный ОКВЭД только среди дополнительных — отдельный лист"),
           ячейка(итоги["доп"], "число", "# ##0"))
    for причина, n in итоги["отсеяно"].most_common():
        строка(ячейка(f"Не вошли: {причина}", "серый"), ячейка(n, "серый", "# ##0"))
    строка()

    строка(ячейка("Откуда данные", "раздел"))
    for t in итоги["источники"]:
        строка(ячейка(t))
    строка()

    строка(ячейка("Насколько можно доверять колонке «В СРО»", "раздел"))
    for t, тревога in итоги["доверие"]:
        строка(ячейка(t, "красный" if тревога else None))
    строка()

    строка(ячейка("Что важно знать", "раздел"))
    for t in итоги["оговорки"]:
        строка(ячейка(t, "серый"))


def cmd_таблица(args) -> int:
    try:
        from openpyxl import Workbook
    except ImportError:
        print("Нужен openpyxl: pip install openpyxl", file=sys.stderr)
        return 1
    db = Path(args.db)
    if not db.exists():
        print(f"Не найдена база {db}. Запускайте из папки, где лежит "
              "mosstroybase.sqlite3, или укажите --db путь.", file=sys.stderr)
        return 1
    print(f"[таблица] читаю {db} …", flush=True)
    основные, доп, отсеяно = читать_базу(db, args.с_ип)
    print(f"[таблица] Москва, основной ОКВЭД 41–43: {len(основные)}; "
          f"только в дополнительных: {len(доп)}", flush=True)

    дамп = Path(args.дамп)
    доверие: list[tuple[str, bool]] = []
    источники = [
        "Список компаний, код региона и ОКВЭД — Единый реестр субъектов МСП ФНС "
        f"(открытые данные), база mosstroybase{', обновлялась ' + дата_базы(db) if дата_базы(db) else ''}.",
    ]
    if дамп.exists():
        нужные = {c["inn"] for c in основные} | {c["inn"] for c in доп}
        по_инн, уникальных = членства_по_инн(дамп, нужные)
        state_path = _состояние(дамп)
        state = json.loads(state_path.read_text(encoding="utf-8")) if state_path.exists() else {}
        заявлено = state.get("всего")
        дата_дампа = разобрать_дату(state.get("дата"))
        источники.append(
            "Членство в СРО — реестр НОСТРОЙ (reestr.nostroy.ru), выгружен целиком"
            f"{f' {дата_дампа:%d.%m.%Y}' if дата_дампа else ''}.")
        if заявлено:
            доля = уникальных / заявлено * 100
            доверие.append((f"Полнота выгрузки НОСТРОЙ: {тыс(уникальных)} записей из "
                            f"заявленных реестром {тыс(заявлено)} ({доля:.1f} %).",
                            доля < 99))
        членство = lambda c: свести(по_инн.get(c["inn"], []))
    else:
        print("[таблица] ВНИМАНИЕ: выгрузки НОСТРОЙ нет — колонка «В СРО» только "
              "по прежним проверкам из базы. Полная: python москва_стройка.py всё",
              file=sys.stderr)
        источники.append("Членство в СРО — только прежние проверки по ИНН из базы "
                         "(выгрузка НОСТРОЙ не делалась).")
        доверие.append(("Выгрузки реестра НОСТРОЙ нет: где прежней проверки не было, "
                        "стоит «Не проверено».", True))

        def членство(c):
            if "sro" not in (c.get("sources") or "") and not c.get("sro_info") \
                    and c.get("sro_member") is None:
                return {"В СРО": "Не проверено", "СРО": "", "Рег. номер СРО": "",
                        "В СРО с": None, "Исключена": None}
            if был_членом_по_прежней_проверке(c):
                return {"В СРО": "Да", "СРО": "", "Рег. номер СРО": "",
                        "В СРО с": None, "Исключена": None}
            if any("исключ" in str(x).lower() for x in c.get("sro_info") or []):
                return {"В СРО": "Исключена", "СРО": "", "Рег. номер СРО": "",
                        "В СРО с": None, "Исключена": None}
            return {"В СРО": "Нет", "СРО": "", "Рег. номер СРО": "",
                    "В СРО с": None, "Исключена": None}

    строки = [строка_таблицы(c, членство(c)) for c in основные]
    строки_доп = [строка_таблицы(c, членство(c)) for c in доп]
    строки.sort(key=lambda s: (s["ОКВЭД"][:2], s["Наименование"].lower()))
    строки_доп.sort(key=lambda s: s["Наименование"].lower())

    if дамп.exists():
        # Сверка 2: прежние проверки по одному ИНН — независимый способ
        прежние = [c for c in основные if был_членом_по_прежней_проверке(c)]
        if прежние:
            потеряны = [c for c in прежние
                        if свести(по_инн.get(c["inn"], []))["В СРО"] == "Нет"]
            доверие.append((
                f"Сверка с прежними проверками по одному ИНН: {len(прежние)} компаний "
                f"тогда числились членами СРО; в выгрузке не нашлось "
                f"{len(потеряны)} ({len(потеряны) / len(прежние) * 100:.1f} %).",
                len(потеряны) / len(прежние) > 0.02))
        # Сверка 3: случайные «Нет», перепроверенные поиском
        выб = Path(args.результат)
        if выб.exists():
            v = json.loads(выб.read_text(encoding="utf-8"))
            доверие.append((
                f"Выборочная перепроверка: {v['проверено']} случайных компаний со "
                f"статусом «Нет» проверены поиском по ИНН, ошибок — {v['ошибок']}.",
                v["ошибок"] > 0))
        else:
            доверие.append(("Выборочная перепроверка «Нет» не проводилась: "
                            "python москва_стройка.py выборка", True))

    по_сро = Counter(s["В СРО"] for s in строки)
    по_разделам: dict[str, Counter] = {}
    по_размеру: dict[str, Counter] = {}
    for s in строки:
        for словарь, ключ in ((по_разделам, s["ОКВЭД"][:2]),
                              (по_размеру, s["Размер (реестр МСП)"] or "не указан")):
            счёт = словарь.setdefault(ключ, Counter())
            счёт["всего"] += 1
            счёт[s["В СРО"]] += 1
    итоги = {
        "подзаголовок": (f"Сформировано {date.today():%d.%m.%Y}. Юридический адрес — "
                         "город Москва (код региона 77, включая Зеленоград и Новую "
                         "Москву). Московская область (код 50) не входит."),
        "всего": len(строки), "по_сро": по_сро,
        "по_разделам": dict(sorted(по_разделам.items())),
        "по_размеру": dict(sorted(по_размеру.items(), key=lambda kv: (
            ["Микропредприятие", "Малое предприятие", "Среднее предприятие"].index(kv[0])
            if kv[0] in ("Микропредприятие", "Малое предприятие", "Среднее предприятие")
            else 9, kv[0]))),
        "доп": len(строки_доп), "отсеяно": отсеяно,
        "источники": источники, "доверие": доверие,
        "оговорки": [
            "В реестр МСП входят только малые и средние предприятия: крупных "
            "строительных компаний в нём нет, и в таблицу они не попали.",
            "«В СРО» — членство в строительной СРО по реестру НОСТРОЙ. "
            "Проектные и изыскательские СРО (НОПРИЗ) здесь не проверялись.",
            "«Исключена» — компания была членом строительной СРО и сейчас не состоит. "
            "Вступить в другую СРО исключённая компания может не раньше чем через год "
            "(дата — в колонке «Исключена из СРО»).",
            "Отобраны компании, у которых строительный ОКВЭД основной. У кого он только "
            "среди дополнительных — на отдельном листе.",
        ],
    }

    стили = _стили()
    wb = Workbook(write_only=True)
    _лист_сводки(wb, стили, итоги)
    пояснение = ("Юрлица с адресом в г. Москве и основным ОКВЭД 41–43. "
                 "Колонка «В СРО» — по реестру НОСТРОЙ.")
    листы = [
        ("Все компании", "Все строительные компании Москвы", строки),
        ("Не в СРО", "Не состоят в строительной СРО",
         [s for s in строки if s["В СРО"] == "Нет"]),
        ("Исключены из СРО", "Были в СРО и исключены",
         sorted([s for s in строки if s["В СРО"] == "Исключена"],
                key=lambda s: s["Исключена из СРО"] or date.min, reverse=True)),
        ("В СРО", "Состоят в строительной СРО",
         [s for s in строки if s["В СРО"] == "Да"]),
    ]
    if по_сро.get("Не проверено"):
        листы.append(("Не проверено", "СРО не проверялась",
                      [s for s in строки if s["В СРО"] == "Не проверено"]))
    for имя, заголовок, данные in листы:
        print(f"[таблица] лист «{имя}»: {len(данные)}", flush=True)
        _лист_компаний(wb, f"{имя} ({len(данные)})", заголовок,
                       f"{пояснение} Строк: {len(данные)}.", данные, стили)
    _лист_компаний(wb, f"ОКВЭД 41–43 доп. ({len(строки_доп)})",
                   "Строительный ОКВЭД только среди дополнительных",
                   "Основной вид деятельности у этих компаний не строительный "
                   "(колонка «Основной ОКВЭД»), а 41–43 есть только среди "
                   "дополнительных. Менее целевая часть базы.", строки_доп, стили,
                   КОЛОНКИ_ДОП)

    out = Path(args.out or f"Москва_стройка_СРО_{date.today():%d.%m.%Y}.xlsx")
    print(f"[таблица] сохраняю {out} …", flush=True)
    wb.save(out)
    print(f"[таблица] готово: {out}")
    print(f"[таблица]   всего {len(строки)}: в СРО {по_сро.get('Да', 0)}, "
          f"не в СРО {по_сро.get('Нет', 0)}, исключены {по_сро.get('Исключена', 0)}"
          + (f", не проверено {по_сро['Не проверено']}" if по_сро.get("Не проверено") else ""))
    for t, тревога in доверие:
        print(("[таблица] ВНИМАНИЕ: " if тревога else "[таблица] ") + t)
    return 0


# -- команда «реестры»: члены заданных СРО -> Москва, ОКВЭД 41–43, СРО --------

# Реестры, о которых просил заказчик. Сравнение по началу названия: «СФЕРА
# проект» ловит и «СФЕРА проектировщиков», а строительная «СИС» (все 374 её
# члена — Петербург) НЕ ловится «СИС проект» — их путать нельзя
РЕЕСТРЫ_ПО_УМОЛЧАНИЮ = ("СФЕРА изыскатели", "СФЕРА проект", "СИС проект", "ЯРД")

ОКВЭД.update({
    "71.11": "Деятельность в области архитектуры",
    "71.12": "Деятельность в области инженерных изысканий, инженерно-технического "
             "проектирования, управления проектами строительства, строительного "
             "контроля и авторского надзора",
})


def _норм_текст(value: Any) -> str:
    text = str(value or "").lower().replace("ё", "е")
    text = re.sub(r"[«»\"'`“”„]", "", text)
    text = re.sub(r"\s*[-–—]\s*", "-", text)      # «системы – проект» = «системы-проект»
    return re.sub(r"\s+", " ", text).strip()


# Как реестр может называться в колонке «Тип организации», в карточке СРО
# внутри выгрузки или в имени файла. «СИС» без «проект» — это строительная
# СРО Петербурга «Строительство. Инженерные системы», и сюда она попасть
# не должна; «\bярд\b» — словом, иначе «миллиард» оказался бы «Ярдом»
ПСЕВДОНИМЫ = {
    "СФЕРА изыскатели": (r"\bсфера изыскат",),
    "СФЕРА проект": (r"\bсфера проект",),
    "СИС проект": (r"инженерные системы-проект", r"\bсис-проект", r"\bсис проект"),
    "ЯРД": (r"\bярд\b",),
}


def реестр_подходит(тип: Any, нужные: Iterable[str]) -> str | None:
    """Какой из нужных реестров назван в тексте — или None."""
    t = _норм_текст(тип)
    if not t:
        return None
    for n in нужные:
        шаблоны = ПСЕВДОНИМЫ.get(n, (r"^" + re.escape(_норм_текст(n)),))
        if any(re.search(ш, t) for ш in шаблоны):
            return n
    return None


def _статус(value: Any) -> str:
    t = _норм_текст(value)
    if not t:
        return ""
    if "является членом" in t or t.startswith("действ"):
        return "действующий"
    if "исключ" in t or "прекращ" in t:
        return "исключён"
    return t


def _сро_из_карточки(text: str) -> str:
    """Краткое имя СРО из карточки выгрузки: вторая строка карточки."""
    части = [p.strip() for p in re.split(r"[;\n]+", text or "") if p.strip()]
    части = [p for p in части if not re.match(r"СРО-[СПИ]-\d", p) and ":" not in p]
    return части[1] if len(части) > 1 else (части[0] if части else "")


def читать_реестр(путь: Path, нужные: Iterable[str]) -> tuple[list[dict], Counter, list[dict]]:
    """(строки нужных реестров, счётчик типов в файле, строки без ИНН).

    Понимает два вида файлов:
    * выгрузку вашей базы с колонкой «Тип организации» — берутся только
      строки нужных реестров;
    * выгрузку реестра одной СРО с сайта (НОПРИЗ/НОСТРОЙ): шапка в две
      строки, затем карточка СРО, затем члены. Реестр определяется по
      карточке, а если её нет — по имени файла.

    Регион берётся из колонки «Регион», а если её нет — из адреса места
    нахождения. Строка с названием, но без российского ИНН не выбрасывается
    молча: она уходит в отдельный список, чтобы счёт сходился.
    """
    from openpyxl import load_workbook
    wb = load_workbook(путь, read_only=True, data_only=True)
    строки, типы, без_инн = [], Counter(), []
    for ws in wb.worksheets:
        все = list(ws.iter_rows(values_only=True))
        i_шапки = next((i for i, row in enumerate(все[:15])
                        if any(_норм_текст(x) == "инн" or _норм_текст(x).startswith("инн ")
                               for x in row)), None)
        if i_шапки is None:
            continue
        шапка = [_норм_текст(x) for x in все[i_шапки]]
        признаки = (
            ("инн", lambda x: x == "инн" or x.startswith("инн ")),
            ("кратко", lambda x: x.startswith("сокращенное наименование")),
            ("полно", lambda x: x.startswith("полное наименование")),
            # не «наименование страховой компании» из подстрок о страховании
            ("имя", lambda x: (x.startswith("наименован") or x == "название")
                              and "саморегулируем" not in x and "страхов" not in x),
            ("номер", lambda x: x in ("n п/п", "№ п/п", "№", "n") or x.startswith("n п/п")
                                or x.startswith("№ п/п")),
            ("тип", lambda x: x.startswith("тип организац")),
            ("статус", lambda x: x.startswith("статус") and "прав" not in x),
            ("регион", lambda x: x.startswith("регион")),
            # «саморегулируем» — колонка со сведениями о самой СРО: её длинный
            # заголовок тоже упоминает адрес юрлица, но у членов она пустая
            ("адрес", lambda x: (x.startswith("адрес места нахождения")
                                 or ("адрес" in x and "юридическ" in x))
                                and "саморегулируем" not in x),
            ("адрес_ип", lambda x: x.startswith("адрес места фактического")),
        )
        кол = {}
        for n, x in enumerate(шапка):
            for ключ, признак in признаки:
                if ключ not in кол and признак(x):
                    кол[ключ] = n

        # Карточка СРО — строка сразу под шапкой с текстом в первой колонке
        # (у строк членов первая колонка пуста). Номер СРО в ней может и не
        # быть: у файла, скачанного командой «сро», он бывает неизвестен
        карточка = next((str(row[0]) for row in все[i_шапки + 1:i_шапки + 6]
                         if row and isinstance(row[0], str) and row[0].strip()
                         and (re.search(r"СРО-[СПИ]-\d", row[0])
                              or re.search(r"ассоциац|союз|саморегулир|партнерств",
                                           row[0], re.I))), "")
        if "тип" not in кол:
            реестр_файла = (реестр_подходит(карточка, нужные)
                            or реестр_подходит(путь.stem, нужные))
            if not реестр_файла:
                # Выгрузка чужой СРО: рядом с нужными есть похожие имена
                # («Центр объединения проектировщиков «СФЕРА-А» рядом с
                # «СФЕРА проектировщиков»), и взять её молча — значит
                # подмешать в таблицу компании не из тех реестров
                print(f"[реестры] ВНИМАНИЕ: {путь.name} — это реестр "
                      f"«{_сро_из_карточки(карточка) or 'не распознан'}», его в списке "
                      f"нет. Файл пропускаю.", file=sys.stderr)
                типы["(не тот реестр — пропущен)"] += 1
                continue
            типы[реестр_файла] += 1

        for row in все[i_шапки + 1:]:
            get = lambda k: row[кол[k]] if k in кол and кол[k] < len(row) else None
            # В выгрузке с сайта под строкой члена идут подстроки — договоры
            # страхования, проверки. У них пусто в «N п/п» и в ИНН, зато
            # заполнено, например, название страховой компании: без этой
            # проверки 832 подстроки СИС записались бы членами без ИНН
            if "номер" in кол and not str(get("номер") or "").strip().isdigit():
                continue
            имя = next((str(v).strip() for v in (get("кратко"), get("полно"), get("имя"))
                        if v is not None and str(v).strip()), "")
            if "тип" in кол:
                тип = str(get("тип") or "").strip()
                реестр = реестр_подходит(тип, нужные)
                if тип:
                    типы[тип] += 1
                if not реестр:
                    continue
            else:
                реестр = реестр_файла
            регион = next((str(v).strip() for v in (get("регион"), get("адрес"),
                                                    get("адрес_ип"))
                           if v is not None and str(v).strip()), "")
            запись = {"inn": норм_инн(get("инн")), "имя": имя, "реестр": реестр,
                      "статус": _статус(get("статус")), "регион": регион,
                      "из_выгрузки": "тип" not in кол}
            if len(запись["inn"]) in (10, 12):
                строки.append(запись)
            elif имя:
                запись["инн_как_есть"] = str(get("инн") or "").strip()
                без_инн.append(запись)
    wb.close()
    return строки, типы, без_инн


def похоже_на_москву(текст: str) -> bool:
    """Регион или адрес из файла реестра — похоже ли на Москву.

    Пустое — «неизвестно», и тоже да: такую компанию надо проверить.
    """
    t = _норм_текст(текст)
    if not t:
        return True
    return bool(re.search(r"\bмосква\b|\bмосквы\b", t)) and москва_по_адресу(текст)


def строки_базы_по_инн(путь: Path, инны: set[str]) -> dict[str, dict]:
    if not путь.exists():
        return {}
    con = sqlite3.connect(f"file:{путь}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    найдено, список = {}, sorted(инны)
    for i in range(0, len(список), 500):
        часть = список[i:i + 500]
        q = f"SELECT * FROM companies WHERE inn IN ({','.join('?' * len(часть))})"
        for row in con.execute(q, часть):
            c = dict(row)
            c["okved_add"] = _json_list(c.get("okved_add"))
            найдено[норм_инн(c["inn"])] = c
    con.close()
    return найдено


def _оквэд_checko(value: Any) -> str:
    if isinstance(value, dict):
        return str(value.get("Код") or "").strip()
    return str(value or "").strip()


def разобрать_checko(data: dict) -> dict | None:
    """Регион, адрес, ОКВЭД и статус из ответа Checko /company.

    Форма ответа снята с настоящего: «Регион»: {"Код": "77"}, «ЮрАдрес»:
    {"АдресРФ": …}, «ОКВЭД»: {"Код": …}, «Статус»: {"Наим": "Действует"};
    дополнительные ОКВЭД — список в «ОКВЭДДоп». Без основного ОКВЭД ответ
    считается неразобранным: гадать нельзя.
    """
    main = _оквэд_checko(data.get("ОКВЭД"))
    if not main:
        return None
    доп = []
    for key, value in data.items():
        if key.startswith("ОКВЭДДоп") and isinstance(value, list):
            доп += [x for x in (_оквэд_checko(v) for v in value) if x]
    регион = data.get("Регион")
    регион = str(регион.get("Код") if isinstance(регион, dict) else регион or "").strip()
    адрес = ""
    юр = data.get("ЮрАдрес")
    if isinstance(юр, dict):
        адрес = str(юр.get("АдресРФ") or юр.get("НасПункт") or "").strip()
    статус = data.get("Статус")
    статус = str(статус.get("Наим") if isinstance(статус, dict) else статус or "")
    return {"region_code": регион, "address": адрес, "okved_main": main,
            "okved_add": доп, "is_active": 0 if re.search(r"ликвид|прекращ|исключ",
                                                          статус.lower()) else 1,
            "источник": "Checko"}


def адрес_checko(инн: str) -> str:
    """У Checko организации и ИП спрашиваются по разным адресам. Спросив ИП
    как организацию, получаешь «не найдена» — так 16 предпринимателей из
    реестров остались непроверенными."""
    вид = "entrepreneur" if len(инн) == 12 else "company"
    return f"https://api.checko.ru/v2/{вид}"


def спросить_checko(session, инн: str, ключ: str) -> tuple[dict | None, str]:
    try:
        resp = session.get(адрес_checko(инн),
                           params={"key": ключ, "inn": инн}, timeout=60)
    except requests.RequestException as exc:
        return None, f"Checko: сеть ({type(exc).__name__})"
    if resp.status_code != 200:
        return None, f"Checko: HTTP {resp.status_code}"
    try:
        payload = resp.json()
    except ValueError:
        return None, "Checko: ответ не разобрался"
    meta = payload.get("meta") if isinstance(payload, dict) else None
    if isinstance(meta, dict) and str(meta.get("status", "")).lower() == "error":
        return None, f"Checko: {meta.get('message') or 'ошибка'}"
    data = payload.get("data") if isinstance(payload, dict) else None
    if not isinstance(data, dict) or not data:
        return None, "Checko: компания не найдена"
    разбор = разобрать_checko(data)
    if разбор is None:
        print(f"[реестры] Checko по ИНН {инн}: не нашёл ОКВЭД в ответе. Ключи ответа: "
              f"{sorted(data)[:30]}", file=sys.stderr)
        return None, "Checko: в ответе нет ОКВЭД"
    return разбор, ""


def оценить(c: dict | None) -> tuple[str, str, str]:
    """(вердикт, строительный ОКВЭД, как) для данных о компании.

    вердикт: «подходит» или причина, по которой не подходит.
    """
    if c is None:
        return "нет данных", "", ""
    if c.get("is_active") == 0:
        return "ликвидирована", "", ""
    регион = str(c.get("region_code") or "").strip()
    if регион and регион != "77":
        return ("Московская область" if регион == "50"
                else f"не Москва (код региона {регион})"), "", ""
    if not москва_по_адресу(c.get("address")):
        return "не Москва (по адресу)", "", ""
    if not регион and not re.search(r"москв", (c.get("address") or "").lower()):
        return "регион не определён", "", ""
    main = (c.get("okved_main") or "").strip()
    if раздел_оквэд(main):
        return "подходит", main, "основной"
    доп = next((x for x in c.get("okved_add") or [] if раздел_оквэд(x)), "")
    if доп:
        return "подходит", доп, "дополнительный"
    return f"нет ОКВЭД 41–43 (основной {main or '—'})", "", ""


def проверить_сро(session, инн: str) -> dict | None:
    payload = запрос(session, тело(1, 50, инн), f"ИНН {инн}")
    if payload is None:
        return None
    свои = [сжать(r) for r in записи_ответа(payload) if норм_инн(r.get("inn")) == инн]
    return свести(свои)


КОЛОНКИ_РЕЕСТРЫ = [
    ("№", 6), ("Наименование", 44), ("ИНН", 13), ("Реестр (статус)", 36),
    ("Юр. адрес", 46), ("ОКВЭД 41–43", 12), ("Вид деятельности", 44),
    ("В строительной СРО", 12), ("Строительная СРО", 46),
]


def cmd_реестры(args) -> int:
    from openpyxl import Workbook
    файлы = [Path(f) for f in args.файлы]
    нет = [str(f) for f in файлы if not f.exists()]
    if нет:
        print(f"Не найдены файлы: {', '.join(нет)}", file=sys.stderr)
        return 1

    # 1. члены заданных реестров
    по_инн: dict[str, dict] = {}
    без_инн: list[dict] = []
    прочитано = [(f, *читать_реестр(f, args.реестры)) for f in файлы]
    # Реестр, выгруженный с сайта, полнее базы: у «СФЕРА изыскатели» в базе
    # 223 компании, в реестре — 643. Если реестр есть и там, и там, берём
    # его только из выгрузки — иначе у одной компании спорили бы два статуса
    из_выгрузок = {s["реестр"] for _f, строки, _т, _п in прочитано
                   for s in строки if s["из_выгрузки"]}
    for f, строки, типы, пропуски in прочитано:
        вытеснены = Counter(s["реестр"] for s in строки
                            if not s["из_выгрузки"] and s["реестр"] in из_выгрузок)
        for реестр, n in вытеснены.items():
            print(f"[реестры] {f.name}: «{реестр}» ({n} строк) не беру — есть "
                  f"выгрузка этого реестра с сайта, она полнее")
        строки = [s for s in строки
                  if s["из_выгрузки"] or s["реестр"] not in из_выгрузок]
        пропуски = [s for s in пропуски
                    if s["из_выгрузки"] or s["реестр"] not in из_выгрузок]
        for s in пропуски:
            метка = f"{s['реестр']} ({s['статус']})" if s["статус"] else s["реестр"]
            без_инн.append({"inn": s["инн_как_есть"] or "—", "имя": s["имя"],
                            "реестры": [метка], "регион": s["регион"], "данные": None,
                            "вердикт": ("нет российского ИНН", "", "")})
        if пропуски:
            print(f"[реестры] {f.name}: без российского ИНН — {len(пропуски)} "
                  f"(лист «Не подошли»)")
        if типы:
            print(f"[реестры] {f.name}: реестры в файле — "
                  + "; ".join(f"{т or '(пусто)'}: {n}" for т, n in типы.most_common()))
        взяты = Counter(s["реестр"] for s in строки)
        print(f"[реестры] {f.name}: взято строк {len(строки)} — "
              + ("; ".join(f"{т}: {n}" for т, n in взяты.most_common()) or "ни одной"))
        for s in строки:
            d = по_инн.setdefault(s["inn"], {"inn": s["inn"], "имя": s["имя"],
                                             "реестры": [], "регион": s["регион"]})
            d["имя"] = d["имя"] or s["имя"]
            # полный адрес из выгрузки сайта важнее «г. Москва» из базы
            if s["регион"] and (s["из_выгрузки"] or not d["регион"]):
                d["регион"] = s["регион"]
            метка = f"{s['реестр']} ({s['статус']})" if s["статус"] else s["реестр"]
            if метка not in d["реестры"]:
                d["реестры"].append(метка)
    for нужный in args.реестры:
        if not any(_норм_текст(r).startswith(_норм_текст(нужный))
                   for d in по_инн.values() for r in d["реестры"]):
            print(f"[реестры] ВНИМАНИЕ: реестра «{нужный}» нет ни в одном файле",
                  file=sys.stderr)
    # Реестр из базы, а не с сайта: в базе он неполный и не всегда чистый —
    # у «СФЕРА проект» в ней члены четырёх разных проектных СРО
    из_базы = sorted({s["реестр"] for _f, строки, _т, _п in прочитано for s in строки
                      if not s["из_выгрузки"] and s["реестр"] not in из_выгрузок})
    for реестр in из_базы:
        print(f"[реестры] ВНИМАНИЕ: «{реестр}» взят из вашей базы — там может быть "
              f"не весь реестр и не только он. Точнее — выгрузка с сайта реестра.",
              file=sys.stderr)
    print(f"[реестры] уникальных компаний: {len(по_инн)}", flush=True)

    # 2. регион и ОКВЭД: база mosstroybase (реестр МСП ФНС), остальным — Checko
    база = строки_базы_по_инн(Path(args.db), set(по_инн))
    print(f"[реестры] нашлись в базе mosstroybase (Москва, ОКВЭД 41–43): {len(база)}")
    ключ = args.checko_key or os.environ.get("CHECKO_API_KEY", "")
    почему_без_checko = "нет ключа Checko"
    # Ответы Checko запоминаются: лимит — около 100 в сутки, и без запоминания
    # повторный запуск назавтра спрашивал бы тех же первых и снова упирался
    # в лимит, так и не дойдя до остальных
    кэш_путь = Path(args.checko_кэш)
    кэш = json.loads(кэш_путь.read_text(encoding="utf-8")) if кэш_путь.exists() else {}
    из_кэша = 0
    session = сессия()
    checko_сделано = checko_осечки = 0
    for инн, d in по_инн.items():
        c = база.get(инн)
        if c is not None and not (c.get("okved_main") or c.get("okved_add")):
            c = None   # добавлена списком ИНН без данных реестра МСП — ОКВЭД неизвестен
        if c is not None:
            c["источник"] = "реестр МСП ФНС"
            d["данные"], d["вердикт"] = c, оценить(c)
            continue
        # В базе все московские компании с 41–43 из реестра МСП. Нет в базе —
        # значит у компании нет 41–43, или она не в Москве, или не входит
        # в реестр МСП (крупная). Последнее и проверяем через Checko,
        # но только тем, кого ваша база числит в Москве или без региона
        if not похоже_на_москву(d["регион"]):
            d["данные"], d["вердикт"] = None, ("не Москва (по адресу в реестре)", "", "")
            continue
        # Адреса нет (список, скачанный командой «сро», его может не дать).
        # Тогда в Checko только тех, чей ИНН выдан в Москве (77, 97): остальных
        # в базе ФНС среди московских строительных нет, а гонять через Checko
        # всех — это сотни запросов и несколько дней лимита
        if not d["регион"].strip() and not d["inn"].startswith(("77", "97")):
            d["данные"], d["вердикт"] = None, (
                "не Москва (адреса в реестре нет, ИНН выдан не в Москве)", "", "")
            continue
        if инн in кэш:
            из_кэша += 1
            d["данные"], d["вердикт"] = кэш[инн], оценить(кэш[инн])
            continue
        if not ключ:
            d["данные"], d["вердикт"] = None, (f"не проверено: {почему_без_checko}", "", "")
            continue
        data, ошибка = спросить_checko(session, инн, ключ)
        checko_сделано += 1
        if data is not None:
            кэш[инн] = data
            кэш_путь.write_text(json.dumps(кэш, ensure_ascii=False), encoding="utf-8")
        if data is None:
            checko_осечки += 1
            d["данные"], d["вердикт"] = None, (
                "не проверено: " + ("лимит Checko исчерпан"
                                    if re.search(r"лимит|limit|429|402", ошибка.lower())
                                    else ошибка), "", "")
            if re.search(r"лимит|limit|429|402", ошибка.lower()):
                ключ = ""   # дальше спрашивать бесполезно
                почему_без_checko = "лимит Checko исчерпан"
        else:
            d["данные"], d["вердикт"] = data, оценить(data)
        time.sleep(args.пауза)
    if checko_сделано or из_кэша:
        print(f"[реестры] Checko: запросов {checko_сделано}, не получилось {checko_осечки}, "
              f"взято из прошлых запусков {из_кэша}")

    итог = [d for d in по_инн.values() if d["вердикт"][0] == "подходит"]
    проверить = [d for d in по_инн.values() if d["вердикт"][0].startswith("не проверено")]
    мимо = [d for d in по_инн.values()
            if d["вердикт"][0] != "подходит" and d not in проверить] + без_инн
    print(f"[реестры] Москва + ОКВЭД 41–43: {len(итог)}; не подошли: {len(мимо)}; "
          f"не удалось проверить: {len(проверить)}", flush=True)

    # 3. строительная СРО — каждой итоговой компании поиском по ИНН
    дамп = Path(args.дамп)
    из_дампа = {}
    if дамп.exists():
        из_дампа, _ = членства_по_инн(дамп, {d["inn"] for d in итог})
    расхождений = 0
    for n, d in enumerate(итог, 1):
        сро = проверить_сро(session, d["inn"])
        if сро is None:
            сро = свести(из_дампа.get(d["inn"], [])) if дамп.exists() else None
            d["сро_как"] = "по выгрузке реестра" if сро else "не проверено"
        else:
            d["сро_как"] = "поиск по ИНН"
            if дамп.exists() and свести(из_дампа.get(d["inn"], []))["В СРО"] != сро["В СРО"]:
                расхождений += 1
        d["сро"] = сро or {"В СРО": "Не проверено", "СРО": "", "Исключена": None}
        if n % 25 == 0:
            print(f"[реестры] СРО проверено {n}/{len(итог)}", flush=True)
        time.sleep(args.пауза)
    if дамп.exists() and итог:
        print(f"[реестры] выгрузка НОСТРОЙ разошлась с поиском по ИНН в {расхождений} "
              f"из {len(итог)} случаев")

    # 4. Excel — только нужное: кто, откуда, где, какой ОКВЭД, есть ли СРО
    def ряд(d):
        c = d.get("данные") or {}
        вердикт, код, как = d["вердикт"]
        сро = d.get("сро") or {}
        адрес = c.get("address") or ""
        if not re.search(r"\d", адрес) and d["регион"] and похоже_на_москву(d["регион"]) \
                and re.search(r"\d", d["регион"]):
            адрес = d["регион"]          # полный адрес из файла реестра
        if not адрес and str(c.get("region_code")) == "77":
            адрес = "г. Москва"
        return {"Наименование": c.get("name_short") or c.get("name") or d["имя"],
                "ИНН": d["inn"], "Реестр (статус)": "; ".join(d["реестры"]),
                "Юр. адрес": адрес,
                "ОКВЭД 41–43": (код + (" (доп.)" if как == "дополнительный" else "")) if код else "",
                "Вид деятельности": название_оквэд(код),
                "В строительной СРО": сро.get("В СРО", ""),
                "Строительная СРО": сро.get("СРО", ""),
                "Причина": re.sub(r"^не проверено:\s*", "", вердикт)}

    порядок_сро = {"Нет": 0, "Исключена": 1, "Да": 2, "Не проверено": 3}
    строки_итог = sorted((ряд(d) for d in итог),
                         key=lambda s: (порядок_сро.get(s["В строительной СРО"], 9),
                                        s["Наименование"].lower()))
    по_сро = Counter(s["В строительной СРО"] for s in строки_итог)
    стили = _стили()
    wb = Workbook(write_only=True)

    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import PatternFill
    ws = wb.create_sheet("Сводка")
    for буква, ширина in zip("ABC", (46, 14, 18)):
        ws.column_dimensions[буква].width = ширина

    def w(*ячейки, стиль=None, фон=None, шапка=False):
        ряд_ = []
        for v in ячейки:
            a = WriteOnlyCell(ws, value=v)
            if шапка:
                a.font, a.fill = стили["шапка_шрифт"], стили["шапка_заливка"]
            elif стиль:
                a.font = стили[стиль]
            if фон:
                a.fill = PatternFill("solid", fgColor=фон)
            ряд_.append(a)
        ws.append(ряд_)

    w("Москва, ОКВЭД 41–43: члены СФЕРА изыскатели, СФЕРА проект, СИС проект, ЯРД",
      стиль="заголовок")
    w(f"на {date.today():%d.%m.%Y}", стиль="подзаголовок")
    w()
    w("", "Компаний", шапка=True)
    w("Всего компаний (без повторов)", len(по_инн) + len(без_инн))
    w("Юр. адрес в Москве и ОКВЭД 41–43", len(строки_итог), стиль="жирный")
    for метка, ключ_, фон in (("   в строительной СРО", "Да", "C6EFCE"),
                              ("   не в строительной СРО", "Нет", "FFC7CE"),
                              ("   исключены из строительной СРО", "Исключена", "FFEB9C"),
                              ("   СРО не проверена", "Не проверено", "EDEDED")):
        if по_сро.get(ключ_) or ключ_ in ("Да", "Нет"):
            w(метка, по_сро.get(ключ_, 0), фон=фон)
    w()
    w("Реестр", "Всего", "Москва, 41–43", шапка=True)
    def реестры_компании(d):          # каждый реестр — один раз на компанию
        return {r.split(" (")[0] for r in d["реестры"]}
    всего_по = Counter(р for d in по_инн.values() for р in реестры_компании(d))
    итог_по = Counter(р for d in итог for р in реестры_компании(d))
    for реестр, n in всего_по.most_common():
        w(реестр, n, итог_по.get(реестр, 0))
    for нужный in args.реестры:
        if not any(_норм_текст(r).startswith(_норм_текст(нужный))
                   for d in по_инн.values() for r in d["реестры"]):
            w(f"«{нужный}» — списка нет, не проверялся", стиль="красный")
    for реестр in из_базы:
        w(f"«{реестр}» — из вашей базы, не с сайта реестра: может быть неполным",
          стиль="красный")
    if проверить:
        w()
        w(f"Не проверено {len(проверить)} — лист «Проверить»", стиль="красный")
    w()
    w("ОКВЭД и адрес — ФНС (реестр МСП) и Checko. Строительная СРО — реестр "
      "НОСТРОЙ, каждая компания проверена по ИНН.", стиль="серый")

    _лист_компаний(wb, f"Компании ({len(строки_итог)})",
                   "Москва, ОКВЭД 41–43",
                   "Сначала — кто не в строительной СРО. «(доп.)» — ОКВЭД 41–43 "
                   "не основной, а дополнительный.",
                   строки_итог, стили, КОЛОНКИ_РЕЕСТРЫ)
    if проверить:
        _лист_компаний(wb, f"Проверить ({len(проверить)})", "Не проверено",
                       "Checko не ответил или кончился лимит. Запустите команду ещё "
                       "раз — со вторым ключом или завтра.",
                       [ряд(d) for d in проверить], стили, КОЛОНКИ_МИМО)
    _лист_компаний(wb, f"Не подошли ({len(мимо)})", "Не подошли",
                   "Почему компания не вошла в список.",
                   sorted((ряд(d) for d in мимо), key=lambda s: (s["Причина"],
                                                                s["Наименование"].lower())),
                   стили, КОЛОНКИ_МИМО)
    out = Path(args.out or f"Реестры_Москва_41-43_СРО_{date.today():%d.%m.%Y}.xlsx")
    wb.save(out)
    print(f"[реестры] готово: {out}")
    print(f"[реестры]   Москва + 41–43: {len(строки_итог)} — в СРО {по_сро.get('Да', 0)}, "
          f"не в СРО {по_сро.get('Нет', 0)}, исключены {по_сро.get('Исключена', 0)}"
          + (f", не проверено {по_сро['Не проверено']}" if по_сро.get("Не проверено") else ""))
    return 0


КОЛОНКИ_МИМО = [("№", 6), ("Наименование", 46), ("ИНН", 13),
                ("Реестр (статус)", 36), ("Причина", 52)]


# -- команда «сро»: члены одной СРО прямо из реестра, когда кнопка выгрузки висит --

РЕЕСТРЫ_САЙТЫ = {"ноприз": "https://reestr.nopriz.ru", "нострой": "https://reestr.nostroy.ru"}


def _адрес_записи(obj: Any, путь: str = "") -> str:
    """Адрес места нахождения члена — где бы он ни лежал в записи.

    Ветку «sro» пропускаем: там адрес самой СРО, а не компании.
    """
    if isinstance(obj, dict):
        for key, value in obj.items():
            k = key.lower()
            if k in ("sro", "sro_info"):
                continue
            if isinstance(value, str) and value.strip() and "address" in k \
                    and "email" not in k and "site" not in k:
                return value.strip()
            if isinstance(value, (dict, list)):
                найдено = _адрес_записи(value, k)
                if найдено:
                    return найдено
    elif isinstance(obj, list):
        for item in obj:
            найдено = _адрес_записи(item, путь)
            if найдено:
                return найдено
    return ""


def _сро_записи(record: dict) -> tuple[str, str, str]:
    sro = record.get("sro") if isinstance(record.get("sro"), dict) else {}
    return (str(sro.get("id") or ""),
            _строка(sro, "full_description", "title", "short_description", "name"),
            _строка(sro, "registration_number"))


def _страницы(session, url: str, что: str, размер: int, пауза_: float
              ) -> tuple[list[dict], int | None]:
    """Все записи постранично с защитой от молчаливой недокачки.

    Размер страницы — по факту первого ответа; пустая страница посреди
    списка — сбой, её переспрашиваем; не пришедшие страницы повторяем в
    конце. Возвращает (записи, заявлено реестром).
    """
    def страница(n, size):
        for attempt in range(1, 6):
            status = None
            try:
                resp = session.post(url, json=тело(n, size), timeout=90)
                if resp.status_code == 200:
                    return resp.json()
                status = resp.status_code
                if status in (404, 405):
                    return "нет пути"
                print(f"[сро] {что}, страница {n}: HTTP {status} (попытка {attempt}/5)",
                      flush=True)
            except (requests.RequestException, ValueError) as exc:
                print(f"[сро] {что}, страница {n}: {type(exc).__name__} "
                      f"(попытка {attempt}/5)", flush=True)
            time.sleep(пауза(status, attempt))
        return None

    первая = страница(1, размер)
    if первая in (None, "нет пути") or not записи_ответа(первая):
        return [], None
    всего = всего_в_ответе(первая)
    получено = записи_ответа(первая)
    size = len(получено) if всего and len(получено) < min(размер, всего) else размер
    страниц = math.ceil(всего / size) if всего else 1
    print(f"[сро] {что}: заявлено {всего}, по {size} на странице — {страниц} стр.",
          flush=True)
    записи, провалы = list(получено), []
    for n in range(2, страниц + 1):
        payload = страница(n, size)
        recs = записи_ответа(payload) if isinstance(payload, dict) else []
        for _ in range(3):            # пустая посреди списка — сбой, а не конец
            if recs:
                break
            time.sleep(3)
            payload = страница(n, size)
            recs = записи_ответа(payload) if isinstance(payload, dict) else []
        if recs:
            записи += recs
        else:
            провалы.append(n)
        if n % 20 == 0:
            print(f"[сро] {что}: страница {n}/{страниц}", flush=True)
        time.sleep(пауза_)
    for n in провалы[:]:
        payload = страница(n, size)
        recs = записи_ответа(payload) if isinstance(payload, dict) else []
        if recs:
            записи += recs
            провалы.remove(n)
    if провалы:
        print(f"[сро] ВНИМАНИЕ: не пришли страницы {провалы}", file=sys.stderr)
    return записи, всего


def cmd_сро(args) -> int:
    from openpyxl import Workbook
    m = re.search(r"(?:/sro/)?(\d+)\s*/?$", args.id.strip())
    if not m:
        print("Нужен номер СРО из адреса страницы, например 397 "
              "(из https://reestr.nopriz.ru/sro/397)", file=sys.stderr)
        return 1
    ид = m.group(1)
    сайт = РЕЕСТРЫ_САЙТЫ["нострой" if "nostroy" in args.id or args.нострой else "ноприз"]
    session = сессия()

    # быстрый путь — список одной СРО; если его нет — весь реестр и отбор
    записи, заявлено = _страницы(session, f"{сайт}/api/sro/{ид}/member/list",
                                 f"СРО {ид}", args.размер, args.пауза)
    чужие = [r for r in записи if _сро_записи(r)[0] not in ("", ид)]
    if записи and чужие:
        print(f"[сро] список СРО {ид} вернул и чужих ({len(чужие)}) — путь не тот, "
              f"иду через весь реестр", flush=True)
        записи = []
    имя_сро = номер_сро = ""
    название = _норм_текст(args.название or "")

    def та_сро(имя: str, сро_ид: str) -> bool:
        """Та ли это СРО: по названию, если оно задано, иначе по номеру."""
        return название in _норм_текст(имя) if название else сро_ид == ид

    if записи:
        # В списке одной СРО реестр не пишет у членов, чья это СРО. Название
        # и проверку, что это правда её члены, берём из общего реестра:
        # несколько членов из списка ищутся там по ИНН. Номер СРО из адреса
        # страницы сверять нельзя: у «СФЕРА проектировщиков» (/sro/397) ни один
        # образец не числился в СРО с номером 397 — нумерации разные. Поэтому
        # смотрим, в одной ли СРО все образцы, и сверяем её по названию
        образцы = []
        for r in записи[:: max(1, len(записи) // 5)][:5]:
            инн = норм_инн(r.get("inn"))
            payload = запрос(session, тело(1, 50, инн), f"проверка ИНН {инн}") if инн else None
            if payload is None:
                continue                  # реестр не ответил — образец не в счёт
            сро = {_сро_записи(x) for x in записи_ответа(payload)
                   if норм_инн(x.get("inn")) == инн}
            образцы.append((инн, сро))
            print(f"[сро]   образец ИНН {инн or '(нет ИНН в списке)'}: "
                  + ("; ".join(f"{имя or '—'} ({номер or '—'}, №{сид or '—'})"
                               for сид, имя, номер in сро) or "в общем реестре не найден"),
                  flush=True)
            time.sleep(args.пауза)
        # СРО, в которой состоят ВСЕ образцы: иначе у списка «всех подряд»
        # случайно попавший наш член сошёл бы за подтверждение
        общие = set.intersection(*(s_ for _i, s_ in образцы)) if образцы else set()
        подходящие = [x for x in общие if x[1] and (not название or та_сро(x[1], x[0]))]
        if len(образцы) >= 3 and len(подходящие) == 1 and (заявлено or 0) <= 50000:
            _сид, имя_сро, номер_сро = подходящие[0]
            print(f"[сро] проверка по образцам: все {len(образцы)} — члены «{имя_сро}»",
                  flush=True)
        else:
            print("[сро] проверка по образцам не подтвердила список — иду через весь реестр",
                  flush=True)
            записи = []
    if not записи:
        if not название:
            print("[сро] ВНИМАНИЕ: без --название отбор по всему реестру идёт по номеру "
                  f"{ид}, а он может не совпадать с номером у членов. Надёжнее: "
                  f'--название "СФЕРА проектировщиков"', file=sys.stderr)
        print("[сро] прохожу весь реестр и отбираю членов этой СРО "
              "(это дольше — 15–30 минут) …", flush=True)
        все, заявлено_всего = _страницы(session, f"{сайт}/api/sro/all/member/list",
                                        "весь реестр", args.размер, args.пауза)
        if заявлено_всего and len(все) < заявлено_всего * 0.99:
            print(f"[сро] ВНИМАНИЕ: реестр отдал {len(все)} из {заявлено_всего} — "
                  f"список СРО может быть неполным", file=sys.stderr)
        записи = [r for r in все if та_сро(_сро_записи(r)[1], _сро_записи(r)[0])]
        имена = Counter(_сро_записи(r)[1:] for r in записи)
        if len(имена) > 1:
            print(f"[сро] ВНИМАНИЕ: под отбор попало несколько СРО: "
                  f"{[n for n, _ in имена.most_common(5)]}", file=sys.stderr)
        if записи:
            имя_сро, номер_сро = имена.most_common(1)[0][0]
        заявлено = None
    if not записи:
        print(f"[сро] членов СРО {ид} не нашлось. Пришлите этот вывод.", file=sys.stderr)
        return 1

    # одна запись — одна строка; дубли (после повторов страниц) схлопываем
    уникальные = {}
    for r in записи:
        уникальные.setdefault(r.get("id") or (r.get("inn"), r.get("registry_registration_date")), r)
    записи = list(уникальные.values())
    if not имя_сро:
        _ид, имя_сро, номер_сро = _сро_записи(записи[0])
    if not имя_сро:
        print(f"[сро] ВНИМАНИЕ: название СРО {ид} узнать не удалось. Переименуйте "
              f"файл по названию СРО, например «СФЕРА проект.xlsx»", file=sys.stderr)

    с_адресом = sum(1 for r in записи if _адрес_записи(r))
    print(f"[сро] {имя_сро} ({номер_сро}): членов {len(записи)}"
          + (f" из заявленных {заявлено}" if заявлено else "")
          + f"; с адресом {с_адресом}", flush=True)
    if заявлено and len(записи) < заявлено:
        print(f"[сро] ВНИМАНИЕ: получено меньше заявленного", file=sys.stderr)

    wb = Workbook()
    ws = wb.active
    ws.append(["Сведения о СРО", "Сведения о члене саморегулируемой организации"])
    ws.append(["Сведения о СРО", "N п/п", "Сокращенное наименование", "Полное наименование",
               "Статус (члена СРО", "Дата вступления", "Дата прекращения членства",
               "ИНН", "ОГРН", "Адрес места нахождения юридического лица"])
    ws.append([f"{имя_сро};\n{номер_сро}"])
    for n, r in enumerate(записи, 1):
        z = сжать(r)
        ws.append([None, str(n), _строка(r, "short_description"), _строка(r, "full_description"),
                   статус_члена(r), z["start"], z["stop"], str(r.get("inn") or ""),
                   _строка(r, "ogrnip", "ogrn"), _адрес_записи(r)])
    out = Path(args.out or f"сро_{ид}.xlsx")
    wb.save(out)
    print(f"[сро] готово: {out} — добавьте его в --файлы команды «реестры»")
    return 0


def cmd_всё(args) -> int:
    for шаг in (cmd_нострой, cmd_выборка, cmd_таблица):
        код = шаг(args)
        if код:
            return код
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Строительные компании Москвы (ОКВЭД 41–43) и членство в СРО")
    общие = argparse.ArgumentParser(add_help=False)
    общие.add_argument("--db", default="mosstroybase.sqlite3", help="база mosstroybase")
    общие.add_argument("--дамп", default=ДАМП, help="куда класть выгрузку НОСТРОЙ")
    общие.add_argument("--результат", default=ВЫБОРКА, help="файл итога выборки")
    общие.add_argument("--out", default=None, help="имя Excel-файла")
    общие.add_argument("--с-ип", dest="с_ип", action="store_true",
                       help="включить индивидуальных предпринимателей")
    общие.add_argument("--размер", type=int, default=500, help="записей на страницу")
    общие.add_argument("--пауза", type=float, default=0.3, help="пауза между запросами, сек")
    общие.add_argument("--заново", action="store_true", help="перекачать НОСТРОЙ заново")
    общие.add_argument("--n", type=int, default=150, help="размер выборки для перепроверки")
    общие.add_argument("--зерно", type=int, default=17, help="зерно случайной выборки")
    sub = parser.add_subparsers(dest="команда")
    for имя, func, справка in (
            ("всё", cmd_всё, "выгрузка НОСТРОЙ, выборочная перепроверка и Excel"),
            ("нострой", cmd_нострой, "выкачать реестр НОСТРОЙ целиком"),
            ("выборка", cmd_выборка, "перепроверить случайные «Нет» поиском по ИНН"),
            ("таблица", cmd_таблица, "собрать Excel")):
        sub.add_parser(имя, parents=[общие], help=справка).set_defaults(func=func)
    p = sub.add_parser("реестры", parents=[общие],
                       help="члены заданных СРО: Москва, ОКВЭД 41–43, строительная СРО")
    p.add_argument("--файлы", nargs="+", required=True,
                   help="выгрузки реестров (xlsx): ваша база с колонкой «Тип "
                        "организации» и/или списки отдельных СРО")
    p.add_argument("--реестры", nargs="+", default=list(РЕЕСТРЫ_ПО_УМОЛЧАНИЮ),
                   help="какие реестры брать из колонки «Тип организации»")
    p.add_argument("--checko-key", dest="checko_key", default=None,
                   help="ключ Checko (или переменная CHECKO_API_KEY)")
    p.add_argument("--checko-кэш", dest="checko_кэш", default="checko_реестры.json",
                   help="где хранить ответы Checko между запусками")
    p.set_defaults(func=cmd_реестры)
    p = sub.add_parser("сро", parents=[общие],
                       help="скачать членов одной СРО из реестра (номер из адреса страницы)")
    p.add_argument("--id", required=True,
                   help="номер СРО или адрес страницы: https://reestr.nopriz.ru/sro/397")
    p.add_argument("--нострой", action="store_true", help="СРО из реестра НОСТРОЙ")
    p.add_argument("--название", default="",
                   help='название СРО, как на сайте: "СФЕРА проектировщиков"')
    p.set_defaults(func=cmd_сро)
    args = parser.parse_args(argv)
    if not getattr(args, "func", None):
        parser.print_help()
        return 1
    if requests is None and args.команда in ("всё", "нострой", "выборка", "реестры", "сро"):
        print("Нужен requests: pip install requests", file=sys.stderr)
        return 1
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
