"""отсев_до_checko.py: не тратить лимит Checko на тех, кто по базе не подходит."""

import importlib.util
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location("отсев", ROOT / "отсев_до_checko.py")
отсев = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(отсев)
checko = отсев.checko


def база(путь: Path, строки: list[tuple]) -> None:
    con = sqlite3.connect(путь)
    con.execute("CREATE TABLE companies (inn TEXT PRIMARY KEY, region_code TEXT, okved_main TEXT, "
                "is_active INTEGER, egrul_status TEXT, updated_at TEXT)")
    con.executemany("INSERT INTO companies VALUES (?, ?, ?, ?, ?, '2026-08-19T10:00:00')", строки)
    con.commit()
    con.close()


def очередь(путь: Path, строки: list[dict]) -> None:
    шапка = ["Наименование", "ИНН", "Телефоны (Checko)", "Статус (Checko)"]
    wb = Workbook()
    ws = wb.active
    ws.append(шапка)
    for с in строки:
        ws.append([с.get(к, "") for к in шапка])
    wb.save(путь)


class TestПричина(unittest.TestCase):
    def test_стройка_москвы_пробивать(self):
        self.assertEqual(отсев.причина({"region_code": "77", "okved_main": "41.20", "is_active": None}, False), "")

    def test_основной_не_стройка(self):
        self.assertIn("не стройка", отсев.причина({"region_code": "77", "okved_main": "71.12"}, False))

    def test_ликвидирована_по_прошлой_проверке(self):
        self.assertIn("ликвидирована", отсев.причина(
            {"region_code": "77", "okved_main": "43.21", "is_active": 0, "egrul_status": "Ликвидировано"}, False))

    def test_нет_в_базе(self):
        self.assertEqual(отсев.причина(None, False), отсев.НЕТ_В_БАЗЕ)
        self.assertEqual(отсев.причина(None, True), "")


class TestФайл(unittest.TestCase):
    def test_отметки_и_пропуск_в_checko(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            база(tmp / "б.sqlite3", [("7701000001", "77", "41.20", 1, ""),
                                     ("7701000002", "77", "71.12", 1, "")])
            путь = tmp / "очередь.xlsx"
            очередь(путь, [{"Наименование": "Стройка", "ИНН": "7701000001"},
                           {"Наименование": "Проект", "ИНН": "7701000002"},
                           {"Наименование": "Нет в базе", "ИНН": "7701000003"},
                           {"Наименование": "Уже пробита", "ИНН": "7701000004",
                            "Статус (Checko)": "Действует"}])
            self.assertEqual(отсев.main(["--file", str(путь), "--база", str(tmp / "б.sqlite3")]), 0)
            ws = load_workbook(путь).active
            шапка = [c.value for c in ws[1]]
            ряды = {r[шапка.index("ИНН")]: dict(zip(шапка, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
            self.assertFalse(ряды["7701000001"]["Не пробивать"])
            self.assertIn("не стройка", ряды["7701000002"]["Не пробивать"])
            self.assertEqual(ряды["7701000003"]["Не пробивать"], отсев.НЕТ_В_БАЗЕ)
            self.assertFalse(ряды["7701000004"]["Не пробивать"])          # пробитую не трогаем
            self.assertEqual(шапка.index("Не пробивать"), шапка.index("ИНН") + 1)

            спрошены = []

            def fetch(inn, key, session, timeout=60, учёт=None):
                спрошены.append(inn)
                return {"Статус": {"Наим": "Действует"}, "ОКВЭД": {"Код": "41.20"}}, ""

            with mock.patch.object(checko, "fetch", side_effect=fetch), \
                    mock.patch.object(checko.time, "sleep"):
                self.assertEqual(checko.main(["--file", str(путь), "--key", "к", "--без-порции"]), 0)
            self.assertEqual(спрошены, ["7701000001"])                     # лимит — только на стройку
            self.assertEqual(checko.осталось_пробить(checko.read_rows(путь)[1]), set())

    def test_без_базы_понятная_ошибка(self):
        with tempfile.TemporaryDirectory() as tmp:
            путь = Path(tmp) / "очередь.xlsx"
            очередь(путь, [{"ИНН": "7701000001"}])
            self.assertEqual(отсев.main(["--file", str(путь), "--база", str(Path(tmp) / "нет.sqlite3")]), 1)


if __name__ == "__main__":
    unittest.main()
