"""отбор_стройка.py: стройка Москвы, которая действует, — из файла после Checko."""

import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location("отбор", ROOT / "отбор_стройка.py")
отбор = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(отбор)


def строка(**поля):
    база = {"Наименование": "ООО «Тест»", "ИНН": "7701234567", "Статус (Checko)": "Действует",
            "Регион (Checko)": "Москва", "ОКВЭД основной (Checko)": "41.20 Строительство жилых и нежилых зданий",
            "Адрес (Checko)": "101000, г. Москва, ул. Мясницкая, д. 5", "Сотрудников (СЧР)": "12",
            "Налоги уплачено, ₽": "350000.0", "Банкротство": ""}
    база.update(поля)
    return база


class TestОтбор(unittest.TestCase):
    def лист(self, **поля):
        return отбор.разобрать(строка(**поля))[0]

    def test_подходит(self):
        self.assertEqual(self.лист(), "Подходят")
        self.assertEqual(self.лист(**{"ОКВЭД основной (Checko)": "43.21 Производство электромонтажных работ"}),
                         "Подходят")

    def test_не_стройка(self):
        self.assertEqual(self.лист(**{"ОКВЭД основной (Checko)": "71.12 Деятельность в области инженерных изысканий"}),
                         "Не стройка")

    def test_московская_область_не_москва(self):
        self.assertEqual(self.лист(**{"Регион (Checko)": "Московская область"}), "Не Москва")

    def test_ликвидирована(self):
        self.assertEqual(self.лист(**{"Статус (Checko)": "Ликвидировано"}), "Не действуют")
        self.assertEqual(self.лист(**{"Статус (Checko)": "В процессе ликвидации"}), "Не действуют")

    def test_ноль_налогов_и_сотрудников_мусор(self):
        self.assertEqual(self.лист(**{"Сотрудников (СЧР)": "0", "Налоги уплачено, ₽": "0"}), "Не действуют")

    def test_нет_данных_о_налогах_не_ноль(self):
        # крупные и бюджетные организации сотрудников и налоги не раскрывают
        лист, _, внимание = отбор.разобрать(строка(**{"Сотрудников (СЧР)": "", "Налоги уплачено, ₽": ""}))
        self.assertEqual(лист, "Подходят")
        self.assertTrue(внимание)

    def test_несколько_причин_все_в_колонке(self):
        лист, почему, _ = отбор.разобрать(строка(**{"Статус (Checko)": "Ликвидировано",
                                                    "Регион (Checko)": "Тверская область",
                                                    "ОКВЭД основной (Checko)": "47.11 Торговля"}))
        self.assertEqual(лист, "Не действуют")
        self.assertEqual(len(почему), 3)

    def test_не_пробита(self):
        self.assertEqual(отбор.разобрать({"ИНН": "7701234567"})[0], "Не проверены")
        self.assertEqual(self.лист(**{"Статус (Checko)": "HTTP 403: лимит или ключ"}), "Не проверены")

    def test_под_вопросом(self):
        self.assertEqual(self.лист(**{"Статус (Checko)": "В отношении юридического лица в деле о "
                                      "несостоятельности (банкротстве) введено наблюдение"}), "Под вопросом")
        self.assertEqual(self.лист(**{"Статус (Checko)": "Юридическим лицом принято решение об "
                                      "изменении места нахождения"}), "Под вопросом")
        self.assertEqual(self.лист(**{"Банкротство": "банкротство"}), "Под вопросом")

    def test_коротко_телефоны_и_почта_вместе(self):
        с = отбор.коротко(строка(**{"Телефоны (Checko)": "+74951490123",
                                   "Телефон (из файла)": "(495) 149-01-23, 8-916-123-45-67",
                                   "E-mail (Checko)": "Info@Firma.ru", "Почта (из файла)": "info@firma.ru; buh@firma.ru"}))
        self.assertEqual(с["Телефон"], "+7 (495) 149-01-23; +7 (916) 123-45-67")
        self.assertEqual(с["Почта"], "info@firma.ru; buh@firma.ru")
        self.assertEqual(с["ОКВЭД основной"], "41.20 Строительство жилых и нежилых зданий")

    def test_отсеяна_до_checko(self):
        лист, почему, _ = отбор.разобрать({"ИНН": "7701234567",
                                            "Не пробивать": "не Москва (база: регион 50)"})
        self.assertEqual((лист, почему), ("Отсеяны до Checko", ["не Москва (база: регион 50)"]))

    def test_старая_версия_без_оквэд_и_региона(self):
        # пробита до 23.10.2026: статус есть, ОКВЭД и региона нет — регион по адресу, ОКВЭД перепробить
        лист, почему, _ = отбор.разобрать(строка(**{"ОКВЭД основной (Checko)": "", "Регион (Checko)": ""}))
        self.assertEqual(лист, "Не проверены")
        self.assertIn("перепробить", почему[0])

    def test_москва_по_адресу_запасным_путём(self):
        self.assertTrue(отбор.москва("", "124482, г. Москва, г. Зеленоград, корп. 1"))
        self.assertFalse(отбор.москва("", "143409, Московская обл., г. Красногорск"))
        self.assertIsNone(отбор.москва("", ""))

    def test_файл(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "москва_для_checko.xlsx"
            wb = Workbook()
            ws = wb.active
            ряды = [строка(), строка(**{"ИНН": "7701234568", "Регион (Checko)": "Московская область"}),
                    {"Наименование": "ООО «Ещё»", "ИНН": "7701234569"}]
            шапка = list(ряды[0])
            ws.append(шапка)
            for р in ряды:
                ws.append([р.get(к, "") for к in шапка])
            wb.save(src)
            self.assertEqual(отбор.main(["--file", str(src)]), 0)
            out = load_workbook(Path(tmp) / "москва_для_checko_отбор.xlsx")
            self.assertEqual(out.sheetnames, ["Сводка", *отбор.ЛИСТЫ])
            self.assertEqual(out["Подходят"].max_row - 1, 1)
            self.assertEqual([c.value for c in out["Подходят"][1]], отбор.КОРОТКО)
            self.assertEqual(out["Не Москва"].max_row - 1, 1)
            self.assertEqual(out["Не проверены"].max_row - 1, 1)


if __name__ == "__main__":
    unittest.main()
