"""checko_контакты.py: ИП, лимит, повторы одной компании — на подменённой сети."""

import importlib.util
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location("checko_k", ROOT / "checko_контакты.py")
checko = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(checko)

НЕ_ТАМ = "Не найдено ни одной организации с указанными реквизитами"


class Ответ:
    def __init__(self, code, payload=None):
        self.status_code, self._payload = code, payload

    def json(self):
        if self._payload is None:
            raise ValueError("не JSON")
        return self._payload


class Сеть:
    """Подменённый Checko: ответы по ИНН, журнал запросов."""

    def __init__(self, ответы):
        self.ответы, self.журнал, self.headers = ответы, [], {}

    def get(self, url, params=None, timeout=None):
        self.журнал.append((url.rsplit("/", 1)[1], params["inn"]))
        ответ = self.ответы[params["inn"]]
        return ответ.pop(0) if isinstance(ответ, list) else ответ


def организация(тел="", фио="", должн="", статус="Действует"):
    data = {"Статус": {"Код": "001", "Наим": статус}}
    if тел:
        data["Контакты"] = {"Тел": [тел]}
    if фио:
        data["Руковод"] = [{"ФИО": фио, "НаимДолжн": должн}]
    return Ответ(200, {"data": data, "meta": {"status": "ok", "today_request_count": 7}})


class TestАдрес(unittest.TestCase):
    def test_ип_по_своему_адресу(self):
        self.assertTrue(checko.адрес_checko("610211153094").endswith("/v2/entrepreneur"))
        self.assertTrue(checko.адрес_checko("6150041102").endswith("/v2/company"))

    def test_старый_ответ_про_ип_повторяется(self):
        # Так отвечали версии скрипта, спрашивавшие ИП как организацию
        self.assertTrue(checko.нужно_повторить(
            {"ИНН": "610211153094", "Статус (Checko)": НЕ_ТАМ}))
        # У организации тот же ответ — окончательный
        self.assertFalse(checko.нужно_повторить(
            {"ИНН": "6150041102", "Статус (Checko)": НЕ_ТАМ}))
        # Ответ, полученный уже по адресу ИП, тоже окончательный
        self.assertFalse(checko.нужно_повторить(
            {"ИНН": "610211153094", "Статус (Checko)": "ИП: " + НЕ_ТАМ}))

    def test_своя_колонка_инн(self):
        self.assertTrue(checko.нужно_повторить(
            {"ИНН компании": "610211153094", "Статус (Checko)": НЕ_ТАМ}, "ИНН компании"))


class TestFetch(unittest.TestCase):
    def спросить(self, инн, ответ):
        сеть = Сеть({инн: ответ})
        return checko.fetch(инн, "ключ", сеть), сеть.журнал

    def test_403_это_лимит(self):
        (data, note), _ = self.спросить("6150041102", Ответ(403, {"message": "Исчерпан лимит"}))
        self.assertIsNone(data)
        self.assertTrue(checko.лимит_кончился(note), note)
        self.assertTrue(checko.нужно_повторить({"Статус (Checko)": note}), note)
        self.assertIn("Исчерпан лимит", note)

    def test_лимит_с_кодом_200(self):
        (data, note), _ = self.спросить("6150041102", Ответ(200, {
            "meta": {"status": "error", "message": "Превышен лимит запросов"}}))
        self.assertTrue(checko.лимит_кончился(note), note)
        self.assertTrue(checko.нужно_повторить({"Статус (Checko)": note}), note)

    def test_исключена_не_лимит(self):
        # «исключена» содержит «ключ», но к ключу API отношения не имеет
        (data, note), _ = self.спросить("6150041102", Ответ(200, {
            "meta": {"status": "error", "message": "Организация исключена из ЕГРЮЛ"}}))
        self.assertFalse(checko.лимит_кончился(note), note)
        (data, note), _ = self.спросить("6150041102", Ответ(200, {
            "meta": {"status": "error", "message": "Неверный ключ"}}))
        self.assertTrue(checko.лимит_кончился(note), note)

    def test_ненайденный_ип_не_повторяется(self):
        (data, note), журнал = self.спросить("610211153094", Ответ(200, {
            "meta": {"status": "error", "message": НЕ_ТАМ}}))
        self.assertEqual(журнал, [("entrepreneur", "610211153094")])
        self.assertTrue(note.startswith("ИП: "))
        self.assertFalse(checko.лимит_кончился(note))
        self.assertFalse(checko.нужно_повторить(
            {"ИНН": "610211153094", "Статус (Checko)": note}))

    def test_500_без_json(self):
        (data, note), _ = self.спросить("6150041102", Ответ(502))
        self.assertEqual(note, "HTTP 502")
        self.assertFalse(checko.лимит_кончился(note))


class TestExtract(unittest.TestCase):
    def test_ип_сам_себе_руководитель(self):
        info = checko.extract({"ФИО": "Алексанян Артур Ашотович",
                               "Статус": {"Наим": "Действует"},
                               "Контакты": {"Тел": ["+7 (928) 111-22-33"]}})
        self.assertEqual(info["Руководитель"], "Алексанян Артур Ашотович")
        self.assertEqual(info["Должность"], "Индивидуальный предприниматель")
        self.assertEqual(info["Телефоны (Checko)"], "+79281112233")

    def test_фио_по_частям(self):
        info = checko.extract({"ФИО": {"Фамилия": "Иванов", "Имя": "Иван", "Отчество": "Иванович"}})
        self.assertEqual(info["Руководитель"], "Иванов Иван Иванович")

    def test_у_организации_должность_руководителя(self):
        info = checko.extract({"Руковод": [{"ФИО": "Петров П. П.", "НаимДолжн": "ДИРЕКТОР"}],
                               "ФИО": "не то"})
        self.assertEqual((info["Руководитель"], info["Должность"]), ("Петров П. П.", "Директор"))

    def test_статус_банкрота_не_теряется(self):
        # Отбор по словам «действ/ликвид/…» оставлял такую колонку пустой
        статус = ("Юридическое лицо признано несостоятельным (банкротом) "
                  "и в отношении него открыто конкурсное производство")
        info = checko.extract({"Статус": {"Наим": статус}})
        self.assertEqual(info["Статус (Checko)"], статус)

    def test_статус_строкой_и_вложенный(self):
        self.assertEqual(checko.extract({"Статус": "Действует"})["Статус (Checko)"], "Действует")
        self.assertEqual(checko.extract({"Ликвид": {"Статус": "Ликвидировано"}})["Статус (Checko)"],
                         "Ликвидировано")


КОЛОНКИ = ["Наименование", "ИНН", "СРО", "Руководитель",
           "Телефоны (Checko)", "Есть телефон", "Статус (Checko)"]


class TestПрогон(unittest.TestCase):
    """Сквозной прогон main() на файле, как у пользователя."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "ростов_для_checko.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def записать(self, строки):
        wb = Workbook()
        ws = wb.active
        ws.append(КОЛОНКИ)
        for s in строки:
            ws.append([s.get(к) for к in КОЛОНКИ])
        wb.save(self.path)

    def прочитать(self):
        rows = list(load_workbook(self.path).active.iter_rows(values_only=True))
        шапка = rows[0]
        return [{к: ("" if v is None else str(v)) for к, v in zip(шапка, r)} for r in rows[1:]]

    def запустить(self, сеть, *доп):
        with mock.patch.object(checko.requests, "Session", return_value=сеть), \
                mock.patch.object(checko.time, "sleep"):
            return checko.main(["--file", str(self.path), "--key", "ключ",
                                "--без-порции", *доп])

    def test_ип_повтор_дубль_и_лимит(self):
        self.записать([
            # уже пробита; ниже в файле — она же из другой СРО
            {"Наименование": "ООО «Альфа»", "ИНН": "6154555042", "СРО": "АСЮО",
             "Руководитель": "Старый", "Телефоны (Checko)": "+78634613665",
             "Есть телефон": "да", "Статус (Checko)": "Действует"},
            # ИП, которого старая версия спросила как организацию
            {"Наименование": "ИП Алексанян", "ИНН": "610211153094", "СРО": "АСЮО",
             "Статус (Checko)": НЕ_ТАМ},
            {"Наименование": "ООО «Альфа»", "ИНН": "6154555042", "СРО": "ЮгСевКавСтрой"},
            # руководитель из выгрузки СРО; у Checko его нет — не затирать
            {"Наименование": "ООО «Бета»", "ИНН": "6150041102", "СРО": "АСЮО",
             "Руководитель": "Камфорин Николай Дмитриевич"},
            # на ней кончится лимит
            {"Наименование": "ООО «Гамма»", "ИНН": "6125014183", "СРО": "АСЮО"},
            {"Наименование": "ООО «Дельта»", "ИНН": "6125027619", "СРО": "АСЮО"},
        ])
        сеть = Сеть({
            "610211153094": Ответ(200, {"data": {"ФИО": "Алексанян Артур Ашотович",
                                                 "Статус": {"Наим": "Действует"},
                                                 "Контакты": {"Тел": ["89281112233"]}}}),
            "6150041102": организация(),
            "6125014183": [Ответ(403, {"message": "Исчерпан лимит"}), организация(тел="+78636023254")],
            "6125027619": организация(тел="+78636223663", фио="Иванов И. И.", должн="ДИРЕКТОР"),
        })
        self.assertEqual(self.запустить(сеть), 0)
        # дубль «Альфы» не запрашивался; на «Гамме» прогон остановился
        self.assertEqual(сеть.журнал, [("entrepreneur", "610211153094"),
                                       ("company", "6150041102"),
                                       ("company", "6125014183")])
        r = self.прочитать()
        self.assertEqual((r[1]["Руководитель"], r[1]["Есть телефон"]),
                         ("Алексанян Артур Ашотович", "да"))
        self.assertEqual(r[2]["Телефоны (Checko)"], "+78634613665")    # копия
        self.assertEqual(r[2]["Руководитель"], "Старый")
        self.assertEqual(r[3]["Руководитель"], "Камфорин Николай Дмитриевич")
        self.assertEqual(r[3]["Есть телефон"], "нет")
        self.assertTrue(r[4]["Статус (Checko)"].startswith("HTTP 403"))
        self.assertEqual(r[4]["Есть телефон"], "")                   # не пробита
        self.assertEqual(checko.осталось_пробить(r), {"6125014183", "6125027619"})

        # назавтра: «Гамма» повторяется, остальные не трогаются
        self.assertEqual(self.запустить(сеть), 0)
        self.assertEqual(сеть.журнал[3:], [("company", "6125014183"), ("company", "6125027619")])
        r = self.прочитать()
        self.assertEqual(r[4]["Телефоны (Checko)"], "+78636023254")
        self.assertEqual(r[5]["Руководитель"], "Иванов И. И.")
        self.assertEqual(checko.осталось_пробить(r), set())

        # и третий раз — запросов нет вовсе
        self.assertEqual(self.запустить(сеть), 0)
        self.assertEqual(len(сеть.журнал), 5)

    def test_пустой_ответ_не_запрашивается_снова(self):
        # Ни телефона, ни почты, ни статуса: раньше такая строка считалась
        # непробитой и съедала запрос каждый день
        self.записать([{"Наименование": "ООО «Пусто»", "ИНН": "6150041102"}])
        сеть = Сеть({"6150041102": Ответ(200, {"data": {"ОГРН": "1"}})})
        self.запустить(сеть)
        self.запустить(сеть)
        self.assertEqual(len(сеть.журнал), 1)
        self.assertEqual(self.прочитать()[0]["Есть телефон"], "нет")

    def test_десять_осечек_подряд_останавливают(self):
        self.записать([{"Наименование": f"ООО {n}", "ИНН": f"61500411{n:02d}"} for n in range(15)])
        сеть = Сеть({f"61500411{n:02d}": Ответ(502) for n in range(15)})
        self.запустить(сеть)
        self.assertEqual(len(сеть.журнал), 10)
        # осечки временные — все 15 остаются в работе
        self.assertEqual(len(checko.осталось_пробить(self.прочитать())), 15)

    def test_limit_не_считает_копии(self):
        self.записать([
            {"Наименование": "А", "ИНН": "6154555042", "Есть телефон": "да",
             "Телефоны (Checko)": "+78634613665", "Статус (Checko)": "Действует"},
            {"Наименование": "А", "ИНН": "6154555042"},
            {"Наименование": "Б", "ИНН": "6150041102"},
        ])
        сеть = Сеть({"6150041102": организация(тел="+78635265327")})
        self.запустить(сеть, "--limit", "1")
        self.assertEqual(сеть.журнал, [("company", "6150041102")])
        self.assertEqual(checko.осталось_пробить(self.прочитать()), set())


if __name__ == "__main__":
    unittest.main()
