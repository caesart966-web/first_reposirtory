"""Отбор после Checko: строительные компании Москвы, которые действуют.

Берёт файл, который пополняет checko_контакты.py, и раскладывает компании
по листам. Подходит компания, у которой одновременно:

* основной ОКВЭД по ЕГРЮЛ — из разделов 41, 42 или 43 (стройка);
* регион юридического адреса по ЕГРЮЛ — Москва (Зеленоград и Новая Москва —
  тоже Москва; Московская область — нет);
* она действует — по тем же правилам, что в москва_стройка.py
  (оценить_живость): «не действует» только по твёрдым фактам — ликвидирована
  или ликвидируется, деятельность прекращена, банкрот в конкурсном
  производстве, либо за год 0 налогов при нуле сотрудников. Отсутствие
  данных — не ноль: у крупных и бюджетных организаций сотрудники и налоги
  в открытых данных не раскрываются.

ЗАПУСК (из папки со скриптами):

    python отбор_стройка.py --file москва_для_checko.xlsx --out москва_стройка_отбор.xlsx

На удалённом компьютере — python\\python.exe вместо python. Запросов к Checko
не делает: всё берётся из уже заполненного файла, запускать можно хоть каждый
день, пока идёт пробивка.

Листы: «Подходят», «Не действуют», «Не Москва», «Не стройка», «Отсеяны
до Checko» (по базе реестра МСП, см. отсев_до_checko.py), «Не проверены».
Компания, не прошедшая сразу по нескольким условиям, лежит на листе первого
из них, а колонка «Почему не подходит» перечисляет все.
"""

from __future__ import annotations

import argparse
import importlib.util
import re
import sys
from collections import Counter
from pathlib import Path

_ЗДЕСЬ = Path(__file__).resolve().parent


def _модуль(имя: str, файл: str):
    spec = importlib.util.spec_from_file_location(имя, _ЗДЕСЬ / файл)
    модуль = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(модуль)
    return модуль


стройка = _модуль("москва_стройка", "москва_стройка.py")

СТРОЙКА = ("41", "42", "43")
ЛИСТЫ = ("Подходят", "Не действуют", "Не Москва", "Не стройка", "Отсеяны до Checko",
         "Не проверены")
ДОБАВИТЬ = ["Отбор", "Почему не подходит", "Обратить внимание"]


def _число(значение) -> float | None:
    текст = str(значение if значение is not None else "").strip().replace(",", ".").replace(" ", "")
    try:
        return float(текст) if текст else None
    except ValueError:
        return None


def москва(регион: str, адрес: str) -> bool | None:
    """True / False, None — данных нет. Регион из ЕГРЮЛ надёжнее адреса;
    адрес — запасной путь для строк без региона."""
    р = (регион or "").strip().lower()
    if р:
        return р in ("77", "москва", "г. москва", "город москва") or (
            "москва" in р and "московск" not in р)
    а = (адрес or "").lower()
    if а:
        return bool(re.search(r"(^|[\s,])(г\.?\s*)?москва\b", а)) and "московская обл" not in а
    return None


def разобрать(строка: dict) -> tuple[str, list[str], str]:
    """(лист, почему не подходит, обратить внимание)."""
    статус = строка.get("Статус (Checko)") or ""
    оквэд = (строка.get("ОКВЭД основной (Checko)") or "").strip()
    регион = строка.get("Регион (Checko)") or ""
    адрес = строка.get("Адрес (Checko)") or ""
    if not (статус or оквэд or регион or адрес):
        отсев = (строка.get("Не пробивать") or "").strip()
        if отсев:
            # отсев_до_checko.py: по базе реестра МСП не подходит — запрос не тратили
            return "Отсеяны до Checko", [отсев], ""
        return "Не проверены", ["ещё не пробита в Checko"], ""
    if статус.startswith(("лимит", "HTTP", "сеть", "ответ не", "не найдено")):
        return "Не проверены", [f"Checko: {статус}"], ""

    почему, листы = [], []
    живость = {"статус_егрюл": статус,
               "сотрудники": _число(строка.get("Сотрудников (СЧР)")),
               "налоги": _число(строка.get("Налоги уплачено, ₽")),
               "банкротство": bool((строка.get("Банкротство") or "").strip())}
    вердикт, причина, внимание = стройка.оценить_живость(живость, строка.get("ИНН") or "")
    if вердикт == "Не действует":
        почему.append(f"не действует: {причина}")
        листы.append("Не действуют")

    м = москва(регион, адрес)
    if м is False:
        почему.append(f"не Москва: {регион or адрес}")
        листы.append("Не Москва")
    код = оквэд.split()[0] if оквэд else ""
    if код and not код.startswith(СТРОЙКА):
        почему.append(f"основной ОКВЭД не стройка: {оквэд}")
        листы.append("Не стройка")
    if not код or м is None:
        # Строка пробита версией скрипта до 23.10.2026 — ОКВЭД и регион она не писала
        почему.append("нет ОКВЭД или региона в ответе — перепробить")
        листы.append("Не проверены")
    if not листы:
        return "Подходят", [], внимание
    порядок = {л: i for i, л in enumerate(ЛИСТЫ)}
    return min(листы, key=порядок.get), почему, внимание


def прочитать(путь: Path) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook
    wb = load_workbook(путь, read_only=True)
    строки = wb.active.iter_rows(values_only=True)
    шапка = [str(x) if x is not None else "" for x in next(строки)]
    записи = []
    for значения in строки:
        запись = {шапка[i]: ("" if v is None else str(v)) for i, v in enumerate(значения)
                  if i < len(шапка)}
        if any(запись.values()):
            записи.append(запись)
    wb.close()
    return шапка, записи


def записать(путь: Path, шапка: list[str], по_листам: dict[str, list[dict]], итог: Counter) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws.append(["Строительные компании Москвы, которые действуют — отбор по данным Checko (ЕГРЮЛ)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Подходят: основной ОКВЭД 41/42/43, юридический адрес в Москве, компания действует."])
    ws.append([])
    for лист in ЛИСТЫ:
        ws.append([лист, итог.get(лист, 0)])
    ws.append(["Всего", sum(итог.values())])
    ws.column_dimensions["A"].width = 40

    колонки = ДОБАВИТЬ + [к for к in шапка if к not in ДОБАВИТЬ]
    for лист in ЛИСТЫ:
        w = wb.create_sheet(лист)
        w.append(колонки)
        for c in w[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A5F")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for запись in по_листам.get(лист, []):
            w.append([запись.get(к, "") for к in колонки])
        for i, к in enumerate(колонки, 1):
            w.column_dimensions[get_column_letter(i)].width = (
                36 if к in ("Наименование", "Почему не подходит", "ОКВЭД основной (Checko)",
                            "Адрес (Checko)", "Адрес (из файла)") else 18)
        w.freeze_panes = "B2"
        w.auto_filter.ref = w.dimensions
    wb.save(путь)


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Стройка Москвы, которая действует — из файла после Checko")
    p.add_argument("--file", required=True, help="файл, который пополняет checko_контакты.py")
    p.add_argument("--out", default=None, help="куда сохранить (по умолчанию <файл>_отбор.xlsx)")
    args = p.parse_args(argv)
    путь = Path(args.file)
    if not путь.exists():
        print(f"Нет файла {путь}", file=sys.stderr)
        return 1
    out = Path(args.out) if args.out else путь.with_name(f"{путь.stem}_отбор.xlsx")
    шапка, записи = прочитать(путь)
    по_листам: dict[str, list[dict]] = {л: [] for л in ЛИСТЫ}
    итог: Counter = Counter()
    for запись in записи:
        лист, почему, внимание = разобрать(запись)
        запись = dict(запись, **{"Отбор": лист, "Почему не подходит": "; ".join(почему),
                                 "Обратить внимание": внимание})
        по_листам[лист].append(запись)
        итог[лист] += 1
    записать(out, шапка, по_листам, итог)
    print(" · ".join(f"{л}: {итог.get(л, 0)}" for л in ЛИСТЫ) + f" → {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
