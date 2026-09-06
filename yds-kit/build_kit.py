#!/usr/bin/env python3
"""YDS deneme kitapçığı üreticisi / YDS practice-kit builder.

Bir kit JSON dosyasını okur, kuralları doğrular ve ÖSYM kitapçığı düzeninde
bir .docx üretir.

    python3 build_kit.py kits/kit-01.json
    python3 build_kit.py kits/kit-01.json --out dist/

Kural kaynağı: spec.md
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# --------------------------------------------------------------------------
# Kit kuralları (spec.md ile birebir aynı olmak zorunda)
# --------------------------------------------------------------------------

TOTAL_KITS = 12

#: Kit numarası -> seviye. Aralıklar kapalıdır (her iki uç dahil).
LEVEL_BANDS = [
    (1, 3, "B2 Beginner"),
    (4, 7, "B2 Advanced"),
    (8, 12, "C1 Entry"),
]

LETTERS = ["A", "B", "C", "D", "E"]
ROMAN = ["I", "II", "III", "IV", "V"]

FIRST_Q, LAST_Q = 27, 80

SECTIONS = [
    {
        "group": "sentence_completion",
        "tr": "CÜMLE TAMAMLAMA",
        "first": 27,
        "last": 36,
        "instr": "For these questions, choose the best option to complete the given sentence.",
    },
    {
        "group": "translation",
        "tr": "ÇEVİRİ",
        "first": 37,
        "last": 42,
        "instr": "For these questions, choose the most accurate Turkish translation of the "
                 "sentences in English, and the most accurate English translation of the "
                 "sentences in Turkish.",
    },
    {
        "group": "paragraph",
        "tr": "PARAGRAF SORULARI",
        "first": 43,
        "last": 62,
        "instr": "Answer these questions according to the passage below.",
    },
    {
        "group": "dialogue",
        "tr": "DİYALOG TAMAMLAMA",
        "first": 63,
        "last": 67,
        "instr": "For these questions, choose the best option to complete the dialogue.",
    },
    {
        "group": "restatement",
        "tr": "YAKIN ANLAMLI CÜMLE",
        "first": 68,
        "last": 71,
        "instr": "For these questions, choose the best rephrased form of the given sentence.",
    },
    {
        "group": "paragraph_completion",
        "tr": "PARAGRAF TAMAMLAMA",
        "first": 72,
        "last": 75,
        "instr": "For these questions, choose the best option to complete the missing part "
                 "of the passage.",
    },
    {
        "group": "odd_one_out",
        "tr": "ANLAM BÜTÜNLÜĞÜNÜ BOZAN CÜMLE",
        "first": 76,
        "last": 80,
        "instr": "For these questions, choose the irrelevant sentence in the passage.",
    },
]

MIN_VOCAB = 40

#: Her paragrafın dört sorusunda bu tiplerin her birinden en az biri bulunmalı.
PARAGRAPH_QTYPES = ("according_to", "it_is_clear", "inference", "main_idea")

# Görsel dil: ÖSYM kitapçığı — lacivert metin, ince çerçeveler, beyaz kâğıt.
NAVY = RGBColor(0x1B, 0x33, 0x5F)
INK = RGBColor(0x14, 0x1C, 0x2B)
RULE = "8FA3BF"
RULE_SOFT = "C3CEDD"
BAR_FILL = "1B335F"
FONT = "Arial"


def level_for_kit(kit: int) -> str:
    """Kit numarasına bağlı seviye. LEVEL_BANDS dışında kit üretilmez."""
    for first, last, name in LEVEL_BANDS:
        if first <= kit <= last:
            return name
    raise ValueError(
        f"Kit {kit} tanımlı aralıkların dışında; geçerli aralık 1-{TOTAL_KITS}."
    )


# --------------------------------------------------------------------------
# Doğrulama
# --------------------------------------------------------------------------

def validate(kit: dict) -> list[str]:
    """Kitapçık üretilmeden önce çalışan sert kontroller. Hata listesi döndürür."""
    errors: list[str] = []
    warnings: list[str] = []

    no = kit.get("kit")
    if not isinstance(no, int) or not 1 <= no <= TOTAL_KITS:
        errors.append(f"kit alanı 1-{TOTAL_KITS} arasında bir tam sayı olmalı: {no!r}")

    questions = kit.get("questions", [])
    numbers = [q.get("no") for q in questions]
    expected = list(range(FIRST_Q, LAST_Q + 1))
    if numbers != expected:
        errors.append(
            f"Soru numaraları {FIRST_Q}-{LAST_Q} arasında ve sırayla olmalı "
            f"(beklenen {len(expected)} soru, gelen {len(numbers)})."
        )

    by_no = {q.get("no"): q for q in questions}
    for sec in SECTIONS:
        for n in range(sec["first"], sec["last"] + 1):
            q = by_no.get(n)
            if q is None:
                continue
            if q.get("group") != sec["group"]:
                errors.append(
                    f"Soru {n}: bölüm {sec['group']} olmalı, {q.get('group')!r} bulundu."
                )

    for q in questions:
        n = q.get("no")
        if q.get("level") not in ("B2", "C1"):
            errors.append(f"Soru {n}: level B2 ya da C1 olmalı, {q.get('level')!r} bulundu.")
        for field in ("why", "trap_note"):
            if not (q.get(field) or "").strip():
                errors.append(f"Soru {n}: {field} boş olamaz (Soru Açıklamaları zorunlu).")

        if q.get("group") == "odd_one_out":
            sentences = q.get("sentences") or []
            if len(sentences) != 5:
                errors.append(f"Soru {n}: anlam bütünlüğünü bozan cümle sorusu 5 cümle ister.")
            if q.get("answer") not in ROMAN:
                errors.append(f"Soru {n}: answer I-V arasında bir Roma rakamı olmalı.")
            if q.get("trap") not in ROMAN:
                errors.append(f"Soru {n}: trap I-V arasında bir Roma rakamı olmalı.")
            continue

        options = q.get("options") or {}
        order = q.get("order") or []
        if len(options) != 5:
            errors.append(f"Soru {n}: tam olarak 5 seçenek olmalı, {len(options)} bulundu.")
        if sorted(order) != sorted(options):
            errors.append(f"Soru {n}: order alanı seçeneklerin bir permütasyonu değil.")
        if q.get("answer") not in options:
            errors.append(f"Soru {n}: answer bir seçenek anahtarı olmalı.")
        if q.get("trap") not in options:
            errors.append(f"Soru {n}: trap bir seçenek anahtarı olmalı.")
        if q.get("answer") == q.get("trap"):
            errors.append(f"Soru {n}: tuzak seçenek doğru cevapla aynı olamaz.")

    for p in kit.get("passages", []):
        if not (p.get("text") or "").strip():
            errors.append(f"Parça {p.get('first')}-{p.get('last')}: metin boş.")

    covered = set()
    for p in kit.get("passages", []):
        covered.update(range(p.get("first", 0), p.get("last", -1) + 1))
    paragraph_qs = {q["no"] for q in questions if q.get("group") == "paragraph"}
    if paragraph_qs - covered:
        errors.append(f"Parçası olmayan paragraf soruları: {sorted(paragraph_qs - covered)}")

    for p in kit.get("passages", []):
        first, last = p.get("first"), p.get("last")
        types = {by_no[n].get("qtype") for n in range(first, last + 1) if n in by_no}
        missing = [t for t in PARAGRAPH_QTYPES if t not in types]
        if missing:
            errors.append(
                f"Parça {first}-{last}: her paragrafta bulunması gereken soru tipleri "
                f"eksik: {', '.join(missing)}."
            )

    vocab = kit.get("vocabulary", [])
    if len(vocab) < MIN_VOCAB:
        errors.append(f"Kelime Röntgeni en az {MIN_VOCAB} madde ister, {len(vocab)} bulundu.")
    if any(len(row) != 2 or not row[0].strip() or not row[1].strip() for row in vocab):
        errors.append("Kelime Röntgeni satırları [İngilizce, Türkçe] biçiminde olmalı.")

    # Kelime Röntgeni yalnızca testte geçen sözcükleri listelemeli.
    blob = _kit_text(kit).lower()
    for en, _tr in vocab:
        head = re.sub(r"\s*\(.*?\)", "", en).strip().lower()
        # Parantez içi çoğu kez sözcüğün testteki çekimli hâlini taşır
        # ("sweep through (swept through)"), o yüzden o da aranır.
        forms = [head] + [f.strip().lower() for f in re.findall(r"\((.*?)\)", en)]
        if head and not any(f and f in blob for f in forms):
            stem = head.split()[0]
            if not re.search(r"\b" + re.escape(stem[: max(4, len(stem) - 3)]), blob):
                warnings.append(f"Kelime Röntgeni: '{en}' testte bulunamadı.")

    for w in warnings:
        print(f"  uyarı: {w}", file=sys.stderr)
    return errors


def _kit_text(kit: dict) -> str:
    """Kitapçıkta basılan tüm İngilizce metin (kelime kontrolü için)."""
    parts = [p.get("text", "") for p in kit.get("passages", [])]
    parts += [s["instr"] for s in SECTIONS]
    for q in kit.get("questions", []):
        parts.append(q.get("stem", ""))
        parts.extend(q.get("sentences", []) or [])
        for who, line in (q.get("dialogue") or []):
            parts.extend((who, line))  # konuşan kişinin adı da kitapçıkta basılır
        parts.extend((q.get("options") or {}).values())
    return "\n".join(parts)


def level_split(kit: dict) -> tuple[int, int]:
    b2 = sum(1 for q in kit["questions"] if q["level"] == "B2")
    c1 = sum(1 for q in kit["questions"] if q["level"] == "C1")
    return b2, c1


def answer_letter(q: dict) -> str:
    """Sorunun doğru cevabının basılı kitapçıktaki harfi."""
    if q["group"] == "odd_one_out":
        return LETTERS[ROMAN.index(q["answer"])]
    return LETTERS[q["order"].index(q["answer"])]


def trap_letter(q: dict) -> str:
    if q["group"] == "odd_one_out":
        return LETTERS[ROMAN.index(q["trap"])]
    return LETTERS[q["order"].index(q["trap"])]


# --------------------------------------------------------------------------
# docx yardımcıları
# --------------------------------------------------------------------------

def _el(tag: str, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn("w:" + k), str(v))
    return el


#: WordprocessingML, alt öğelerin sırasını şemayla sabitler; yanlış sırada
#: eklenen bir öğe belgeyi tümüyle açılamaz hâle getirir.
TBL_PR_ORDER = ("w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual",
                "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
                "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout",
                "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")

P_PR_ORDER = ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
              "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
              "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
              "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
              "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
              "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
              "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
              "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")

R_PR_ORDER = ("w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps",
              "w:smallCaps", "w:strike", "w:dstrike", "w:outline", "w:shadow", "w:emboss",
              "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden",
              "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz", "w:szCs",
              "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText",
              "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang")


def _insert_ordered(parent, child, order):
    """Şema sırasını koruyarak ekler: kendinden sonra gelen ilk öğenin önüne."""
    names = [qn(tag) for tag in order]
    position = names.index(child.tag)
    for existing in parent:
        if existing.tag in names and names.index(existing.tag) > position:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child


def _replace_ordered(parent, child, order):
    """Aynı adlı öğe varsa değiştirir, yoksa sıraya uygun biçimde ekler."""
    existing = parent.find(child.tag)
    if existing is not None:
        parent.replace(existing, child)
        return child
    return _insert_ordered(parent, child, order)


def _shade(xml_parent, fill: str):
    xml_parent.append(_el("w:shd", val="clear", color="auto", fill=fill))


def set_cols(section, num: int, space_cm: float = 0.6, sep: bool = True):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(int(space_cm * 567)))
    cols.set(qn("w:sep"), "1" if (sep and num > 1) else "0")
    cols.set(qn("w:equalWidth"), "1")


def flow(doc, cols: int, new_page: bool = False):
    """Sütun düzenini değiştirmek için yeni bir bölüm açar."""
    section = doc.add_section(WD_SECTION.NEW_PAGE if new_page else WD_SECTION.CONTINUOUS)
    set_cols(section, cols)
    return section


def para(container, text="", *, size=9.5, bold=False, italic=False, color=INK,
         align=None, left=0.0, hanging=0.0, right=0.0, before=0, after=2,
         line=1.02, font=FONT, keep_next=False):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(left)
    pf.right_indent = Cm(right)
    if hanging:
        pf.first_line_indent = Cm(-hanging)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_together = True
    if keep_next:
        pf.keep_with_next = True
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return p


def run_in(p, text, *, size=9.5, bold=False, italic=False, color=INK, font=FONT):
    run = p.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return run


def _table_style(table, *, pct=100, border=RULE, sz=6, kinds=("top", "left", "bottom", "right"),
                 margins=(50, 100, 60, 100)):
    tbl_pr = table._tbl.tblPr
    _replace_ordered(tbl_pr, _el("w:tblW", w=int(pct * 50), type="pct"), TBL_PR_ORDER)
    borders = OxmlElement("w:tblBorders")
    for kind in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if kind in kinds:
            borders.append(_el(f"w:{kind}", val="single", sz=sz, space=0, color=border))
        else:
            borders.append(_el(f"w:{kind}", val="none", sz=0, space=0, color="auto"))
    _replace_ordered(tbl_pr, borders, TBL_PR_ORDER)
    mar = OxmlElement("w:tblCellMar")
    for kind, value in zip(("top", "left", "bottom", "right"), margins):
        mar.append(_el(f"w:{kind}", w=value, type="dxa"))
    _replace_ordered(tbl_pr, mar, TBL_PR_ORDER)
    table.autofit = False


def box(doc, *, pct=100, border=RULE, sz=6, fill=None, margins=(50, 100, 60, 100)):
    """Tek hücreli çerçeveli kutu döndürür; içine paragraf eklenir."""
    table = doc.add_table(rows=1, cols=1)
    _table_style(table, pct=pct, border=border, sz=sz, margins=margins)
    cell = table.cell(0, 0)
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    if fill:
        _shade(cell._tc.get_or_add_tcPr(), fill)
    return cell


def hold_together(cell):
    """Kutunun içeriğini bir arada tutar: soru sütun sonunda ikiye bölünmesin."""
    paragraphs = cell.paragraphs
    for paragraph in paragraphs[:-1]:
        paragraph.paragraph_format.keep_with_next = True
    if paragraphs:
        paragraphs[-1].paragraph_format.keep_with_next = False


def spacer(doc, pt=6, keep_next=False):
    para(doc, "", size=1, after=pt, keep_next=keep_next)


def add_field(paragraph, instr, placeholder="1", size=8, bold=False, color=NAVY):
    fld = _el("w:fldSimple", instr=instr)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_pr.append(_el("w:rFonts", ascii=FONT, hAnsi=FONT))
    if bold:
        r_pr.append(OxmlElement("w:b"))
    r_pr.append(_el("w:color", val=str(color)))
    r_pr.append(_el("w:sz", val=int(size * 2)))
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = placeholder
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


# --------------------------------------------------------------------------
# Kitapçık düzeni
# --------------------------------------------------------------------------

JUSTIFY = WD_ALIGN_PARAGRAPH.JUSTIFY
CENTER = WD_ALIGN_PARAGRAPH.CENTER
RIGHT = WD_ALIGN_PARAGRAPH.RIGHT
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

CONTENT_CM = 18.0


class Layout:
    """Sütun sayısını izler; gereksiz bölüm sonu eklenmesini önler."""

    def __init__(self, doc):
        self.doc = doc
        self.cols = 1

    def need(self, cols: int, new_page: bool = False):
        if new_page or cols != self.cols:
            flow(self.doc, cols, new_page=new_page)
            self.cols = cols


def setup_page(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.7)


def _p_border(paragraph, side="bottom", sz=8, color="1B335F"):
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bdr.append(_el(f"w:{side}", val="single", sz=sz, space=3, color=color))
    _replace_ordered(p_pr, bdr, P_PR_ORDER)


def _bare_cell(cell, text_align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = text_align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    return p


def attach_header_footer(section, kit):
    """Sekme durakları yerine çerçevesiz tablo: hizalama şablondan bağımsız olur."""
    kit_no = kit["kit"]
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Cm(CONTENT_CM))
    _table_style(table, pct=100, border=NAVY_HEX, sz=8, kinds=("bottom",),
                 margins=(0, 0, 40, 0))
    for cell, width in zip(table.rows[0].cells, (Cm(12), Cm(6))):
        cell.width = width
    run_in(_bare_cell(table.cell(0, 0)), f"YDS DENEME — KIT {kit_no:02d}",
           size=9, bold=True, color=NAVY)
    run_in(_bare_cell(table.cell(0, 1), RIGHT), "ENGLISH", size=11, bold=True, color=NAVY)
    _drop_first_paragraph(header)
    para(header, "", size=1, after=0)

    footer = section.footer
    table = footer.add_table(rows=1, cols=3, width=Cm(CONTENT_CM))
    _table_style(table, pct=100, border=NAVY_HEX, sz=8, kinds=("top",),
                 margins=(40, 0, 0, 0))
    for cell, width in zip(table.rows[0].cells, (Cm(7), Cm(4), Cm(7))):
        cell.width = width
    run_in(_bare_cell(table.cell(0, 0)), "Go on to the next page.",
           size=8, italic=True, color=NAVY)
    add_field(_bare_cell(table.cell(0, 1), CENTER), " PAGE ", size=8, bold=True, color=NAVY)
    run_in(_bare_cell(table.cell(0, 2), RIGHT), f"Kit {kit_no} / {TOTAL_KITS}",
           size=8, color=NAVY)
    _drop_first_paragraph(footer)
    para(footer, "", size=1, after=0)


def _drop_first_paragraph(container):
    first = container.paragraphs[0]
    first._p.getparent().remove(first._p)


def section_bar(doc, title: str):
    cell = box(doc, pct=100, border=BAR_FILL, sz=6, fill=BAR_FILL, margins=(60, 150, 60, 150))
    para(cell, title, size=10.5, bold=True, color=WHITE, after=0, line=1.0)
    spacer(doc, 4)


def instruction_box(doc, first: int, last: int, text: str, pct=62):
    cell = box(doc, pct=pct, border=NAVY_HEX, sz=6, margins=(60, 110, 60, 110))
    p = para(cell, "", size=9, after=0, line=1.05)
    run_in(p, f"{first}-{last}: ", size=9, bold=True, color=NAVY)
    run_in(p, text, size=9, bold=True, color=NAVY)
    spacer(doc, 5)


NAVY_HEX = "1B335F"


def passage_block(doc, text: str):
    cell = box(doc, pct=100, border=RULE_SOFT, sz=4, margins=(90, 150, 90, 150))
    para(cell, text, size=9.5, align=JUSTIFY, line=1.1, after=0)
    spacer(doc, 6)


def _stem_paragraph(cell, number: int, indent=0.8):
    p = para(cell, "", left=indent, hanging=indent, after=3, line=1.05)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(indent))
    run_in(p, f"{number}.", bold=True, color=NAVY)
    run_in(p, "\t")
    return p


def _options(cell, q, indent=1.55, hang=0.62, size=9.5):
    for idx, key in enumerate(q["order"]):
        p = para(cell, "", left=indent, hanging=hang, after=2, line=1.03)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(indent))
        run_in(p, f"{LETTERS[idx]})", size=size)
        run_in(p, "\t", size=size)
        run_in(p, q["options"][key], size=size)


def render_question(doc, q):
    group = q["group"]
    cell = box(doc, pct=100, border=RULE, sz=4, margins=(70, 110, 80, 110))

    if group == "dialogue":
        first = True
        for who, line in q["dialogue"]:
            p = para(cell, "", left=0.8, hanging=0.8 if first else 0.0,
                     after=0, before=0 if first else 3, line=1.05)
            if first:
                p.paragraph_format.tab_stops.add_tab_stop(Cm(0.8))
                run_in(p, f"{q['no']}.", bold=True, color=NAVY)
                run_in(p, "\t")
                first = False
            run_in(p, f"{who}:")
            body = para(cell, "", left=1.4, hanging=0.4, after=0, line=1.05)
            body.paragraph_format.tab_stops.add_tab_stop(Cm(1.4))
            run_in(body, "–")
            run_in(body, "\t")
            run_in(body, line, bold=True)
        spacer(cell, 3)
        _options(cell, q)

    elif group == "odd_one_out":
        p = _stem_paragraph(cell, q["no"])
        for idx, sentence in enumerate(q["sentences"]):
            run_in(p, f"({ROMAN[idx]}) ", bold=True)
            run_in(p, sentence + (" " if idx < 4 else ""))
        opt = para(cell, "", left=0.9, after=0, before=4, line=1.0)
        step = (CONTENT_CM / 2 - 1.6) / 5
        for idx, roman in enumerate(ROMAN):
            opt.paragraph_format.tab_stops.add_tab_stop(Cm(0.9 + step * idx))
            if idx:
                run_in(opt, "\t")
            run_in(opt, f"{LETTERS[idx]}) {roman}")

    else:
        p = _stem_paragraph(cell, q["no"])
        run_in(p, q["stem"], bold=True)
        p.paragraph_format.space_after = Pt(6)
        _options(cell, q)

    hold_together(cell)
    spacer(doc, 5)


# --------------------------------------------------------------------------
# Kapak ve kapanış sayfaları
# --------------------------------------------------------------------------

COVER_RULES = [
    "Bu testte {count} soru vardır; toplam cevaplama süresi {minutes} dakikadır.",
    "Cevaplarınızı, kitapçığın sonundaki cevap anahtarında yer alan "
    "“Benim cevabım” sütununa kurşun kalemle yazınız.",
    "Değiştirmek istediğiniz bir cevabı yumuşak bir silgiyle temizleyiniz; "
    "cevap anahtarına karalama yapmayınız.",
    "Her sorunun yalnızca bir doğru cevabı vardır. Birden fazla işaretleme "
    "yapılan sorular değerlendirilmeye alınmaz.",
    "Süre bitmeden testi tamamlarsanız, kalan süreyi cevaplarınızı gözden "
    "geçirmek için kullanınız.",
    "Bu kitapçık bir ÖSYM yayını değildir; sınavın biçimi örnek alınarak "
    "hazırlanmış bir deneme çalışmasıdır.",
]


def render_cover(doc, kit):
    kit_no = kit["kit"]
    level = level_for_kit(kit_no)
    b2, c1 = level_split(kit)
    count = len(kit["questions"])
    minutes = round(count * 180 / 80 / 5) * 5

    spacer(doc, 26)
    cell = box(doc, pct=100, border=NAVY_HEX, sz=12, margins=(180, 200, 180, 200))
    para(cell, f"YDS DENEME — {kit['topic']} · {kit['date']}", size=20, bold=True,
         color=NAVY, align=CENTER, after=8, line=1.0)
    para(cell, f"Kit {kit_no} / {TOTAL_KITS} — Level: {level}", size=13, bold=True,
         color=INK, align=CENTER, after=0, line=1.0)

    spacer(doc, 14)
    para(doc, f"Kit No: {kit_no} / {TOTAL_KITS}", size=11.5, bold=True, color=INK,
         align=CENTER, after=4)
    para(doc, f"Seviye: B2 ({b2} soru) / C1 ({c1} soru)", size=11.5, bold=True,
         color=INK, align=CENTER, after=4)
    para(doc, f"TEST OF ENGLISH — {count} soru · {minutes} dakika", size=10.5,
         color=NAVY, align=CENTER, after=0)

    spacer(doc, 20)
    cell = box(doc, pct=100, border=RULE, sz=6, margins=(120, 170, 130, 170))
    para(cell, "SINAVDA UYULACAK KURALLAR", size=10.5, bold=True, color=NAVY, after=6)
    for idx, rule in enumerate(COVER_RULES, start=1):
        p = para(cell, "", size=9.5, left=0.75, hanging=0.75, after=4, line=1.1)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
        run_in(p, f"{idx}.", size=9.5, bold=True, color=NAVY)
        run_in(p, "\t", size=9.5)
        run_in(p, rule.format(count=count, minutes=minutes), size=9.5)

    spacer(doc, 16)
    cell = box(doc, pct=100, border=RULE_SOFT, sz=4, margins=(90, 150, 90, 150))
    para(cell, "SORU DAĞILIMI", size=10, bold=True, color=NAVY, after=5)
    for sec in SECTIONS:
        total = sec["last"] - sec["first"] + 1
        p = para(cell, "", size=9.5, after=2, line=1.05)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(7.0))
        p.paragraph_format.tab_stops.add_tab_stop(Cm(10.5), WD_TAB_ALIGNMENT.RIGHT)
        run_in(p, sec["tr"], size=9.5, bold=True, color=NAVY)
        run_in(p, "\t", size=9.5)
        run_in(p, f"{sec['first']}–{sec['last']}", size=9.5)
        run_in(p, "\t", size=9.5)
        run_in(p, f"{total} soru", size=9.5)


def render_explanations(doc, layout, kit):
    layout.need(1)  # gövdenin son iki sütununu dengeler
    layout.need(1, new_page=True)
    para(doc, f"SORU AÇIKLAMALARI — Kit {kit['kit']}", size=15, bold=True, color=NAVY, after=3)
    para(doc, "Her soru için: doğru cevabın gerekçesi ve en çok işaretlenen yanlış "
              "seçenekteki tuzak.", size=9, italic=True, color=NAVY, after=8)
    layout.need(2)
    for q in kit["questions"]:
        p = para(doc, "", size=8.5, left=0.0, after=4, line=1.05, align=JUSTIFY)
        run_in(p, f"Q{q['no']} — Doğru: {answer_letter(q)}. ", size=8.5, bold=True, color=NAVY)
        run_in(p, q["why"] + " ", size=8.5)
        run_in(p, f"Tuzak ({trap_letter(q)}): ", size=8.5, bold=True, color=NAVY)
        run_in(p, q["trap_note"], size=8.5)


def _cell_text(cell, text, *, size=9, bold=False, color=INK, align=CENTER, fill=None):
    if fill:
        _shade(cell._tc.get_or_add_tcPr(), fill)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    if text:
        run_in(p, text, size=size, bold=bold, color=color)
    return p


def render_answer_key(doc, layout, kit):
    layout.need(1, new_page=True)
    para(doc, f"CEVAP ANAHTARI — Kit {kit['kit']}", size=15, bold=True, color=NAVY, after=3)
    para(doc, "“Benim cevabım” sütunu boş bırakılmıştır; testi çözerken kendi "
              "cevaplarınızı buraya yazınız.", size=9, italic=True, color=NAVY, after=8)

    questions = kit["questions"]
    rows_per_block, blocks = 18, 3
    table = doc.add_table(rows=rows_per_block + 1, cols=4 * blocks)
    _table_style(table, pct=100, border=RULE, sz=4,
                 kinds=("top", "left", "bottom", "right", "insideH", "insideV"),
                 margins=(30, 60, 30, 60))
    widths = [Cm(1.05), Cm(1.15), Cm(2.7), Cm(1.1)] * blocks
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width

    headers = ["Soru", "Doğru", "Benim cevabım", "✓ / ✗"]
    for block in range(blocks):
        for offset, label in enumerate(headers):
            _cell_text(table.cell(0, block * 4 + offset), label, size=8.5, bold=True,
                       color=NAVY, fill="EAEEF4")

    for index, q in enumerate(questions):
        block, row = divmod(index, rows_per_block)
        _cell_text(table.cell(row + 1, block * 4), str(q["no"]), size=9)
        _cell_text(table.cell(row + 1, block * 4 + 1), answer_letter(q), size=9, bold=True,
                   color=NAVY)
        _cell_text(table.cell(row + 1, block * 4 + 2), "", size=9)
        _cell_text(table.cell(row + 1, block * 4 + 3), "", size=9)

    spacer(doc, 12)
    cell = box(doc, pct=100, border=RULE_SOFT, sz=4, margins=(90, 150, 90, 150))
    para(cell, "ANSWER KEY", size=10, bold=True, color=NAVY, after=4)
    # Bölünmez tire: soru numarası ile cevabı satır sonunda ayrılmaz.
    line = "   ".join(f"{q['no']}\u2011{answer_letter(q)}" for q in questions)
    para(cell, line, size=9.5, after=0, line=1.3)


def render_vocabulary(doc, layout, kit):
    layout.need(1, new_page=True)
    para(doc, f"KELİME RÖNTGENİ — Kit {kit['kit']}", size=15, bold=True, color=NAVY, after=3)
    para(doc, f"Testte geçen C1+ sözcük, kalıp ve eş dizimler; alfabetik sıra ile "
              f"({len(kit['vocabulary'])} madde).", size=9, italic=True, color=NAVY, after=8)

    entries = sorted(kit["vocabulary"], key=lambda row: row[0].lower())
    half = (len(entries) + 1) // 2
    columns = [entries[:half], entries[half:]]
    rows = max(len(columns[0]), len(columns[1]))

    table = doc.add_table(rows=rows + 1, cols=4)
    _table_style(table, pct=100, border=RULE_SOFT, sz=4,
                 kinds=("top", "left", "bottom", "right", "insideH", "insideV"),
                 margins=(20, 70, 20, 70))
    widths = [Cm(4.3), Cm(4.7), Cm(4.3), Cm(4.7)]
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    for pair in (0, 1):
        _cell_text(table.cell(0, pair * 2), "English word / phrase", size=8.5, bold=True,
                   color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT, fill="EAEEF4")
        _cell_text(table.cell(0, pair * 2 + 1), "Türkçe karşılığı", size=8.5, bold=True,
                   color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT, fill="EAEEF4")

    for pair, column in enumerate(columns):
        for index, (en, tr) in enumerate(column):
            _cell_text(table.cell(index + 1, pair * 2), en, size=8.5,
                       align=WD_ALIGN_PARAGRAPH.LEFT)
            _cell_text(table.cell(index + 1, pair * 2 + 1), tr, size=8.5,
                       align=WD_ALIGN_PARAGRAPH.LEFT)


# --------------------------------------------------------------------------
# Gövde ve giriş noktası
# --------------------------------------------------------------------------

def close_footer(section, kit):
    """Son sayfanın altbilgisi ileri yönlendirmez."""
    section.footer.is_linked_to_previous = False
    table = section.footer.add_table(rows=1, cols=3, width=Cm(CONTENT_CM))
    _table_style(table, pct=100, border=NAVY_HEX, sz=8, kinds=("top",),
                 margins=(40, 0, 0, 0))
    for cell, width in zip(table.rows[0].cells, (Cm(7), Cm(4), Cm(7))):
        cell.width = width
    run_in(_bare_cell(table.cell(0, 0)), "TEST BİTTİ.", size=8, bold=True, color=NAVY)
    add_field(_bare_cell(table.cell(0, 1), CENTER), " PAGE ", size=8, bold=True, color=NAVY)
    run_in(_bare_cell(table.cell(0, 2), RIGHT), f"Kit {kit['kit']} / {TOTAL_KITS}",
           size=8, color=NAVY)
    _drop_first_paragraph(section.footer)
    para(section.footer, "", size=1, after=0)


def render_body(doc, layout, kit):
    by_no = {q["no"]: q for q in kit["questions"]}
    passages = {(p["first"], p["last"]): p["text"] for p in kit["passages"]}

    for sec in SECTIONS:
        layout.need(1)
        section_bar(doc, sec["tr"])

        if sec["group"] == "paragraph":
            for (first, last), text in sorted(passages.items()):
                layout.need(1)
                instruction_box(doc, first, last, sec["instr"])
                passage_block(doc, text)
                layout.need(2)
                for n in range(first, last + 1):
                    render_question(doc, by_no[n])
            continue

        # Diyaloglar tam sayfa genişliğinde basılır; diğer bölümler iki sütun.
        # Yönerge kutusu sütunun içine girer, böylece ilk sorudan kopmaz.
        columns = 1 if sec["group"] == "dialogue" else 2
        layout.need(columns)
        instruction_box(doc, sec["first"], sec["last"], sec["instr"],
                        pct=100 if columns == 1 else 96)
        for n in range(sec["first"], sec["last"] + 1):
            render_question(doc, by_no[n])


def build(kit: dict, out_path: Path) -> Path:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(2)

    setup_page(doc.sections[0])
    set_cols(doc.sections[0], 1)

    render_cover(doc, kit)

    layout = Layout(doc)
    layout.need(1, new_page=True)
    attach_header_footer(doc.sections[-1], kit)

    render_body(doc, layout, kit)
    render_explanations(doc, layout, kit)
    render_answer_key(doc, layout, kit)
    render_vocabulary(doc, layout, kit)
    close_footer(doc.sections[-1], kit)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="YDS deneme kitapçığı üretir.")
    parser.add_argument("kit", type=Path, help="kit JSON dosyası, örn. kits/kit-01.json")
    parser.add_argument("--out", type=Path, default=Path(__file__).parent / "dist",
                        help="çıktı klasörü (varsayılan: dist/)")
    args = parser.parse_args(argv)

    kit = json.loads(args.kit.read_text(encoding="utf-8"))

    errors = validate(kit)
    if errors:
        print(f"{args.kit}: {len(errors)} kural ihlali", file=sys.stderr)
        for error in errors:
            print(f"  - {error}", file=sys.stderr)
        return 1

    level = level_for_kit(kit["kit"])
    b2, c1 = level_split(kit)
    out = build(kit, args.out / f"YDS-DENEME-Kit-{kit['kit']:02d}.docx")

    print(f"Kit {kit['kit']} / {TOTAL_KITS} — Level: {level}")
    print(f"  sorular : {len(kit['questions'])} ({FIRST_Q}-{LAST_Q})")
    print(f"  seviye  : B2 ({b2} soru) / C1 ({c1} soru)")
    print(f"  kelime  : {len(kit['vocabulary'])} madde")
    print(f"  cevaplar: " + " ".join(
        f"{q['no']}-{answer_letter(q)}" for q in kit["questions"][:8]) + " ...")
    print(f"  çıktı   : {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
