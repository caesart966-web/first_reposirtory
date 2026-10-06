"""Сборка итоговой таблицы: all.json (исходные файлы) + out/batch_*.json (поиск)."""
import json, glob, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else 'result.xlsx'
src = json.load(open('all.json'))
def score(r):
    return 2 * (r.get('status', 'Не установлено') != 'Не установлено') + bool(r.get('email'))

# out/batch_*.json — первый проход, out/r*.json — следующие ходы; берём лучший результат по ИНН
res = {}
for f in sorted(glob.glob('out/batch_*.json')) + sorted(glob.glob('out/r*.json')):
    for r in json.load(open(f)):
        if (r.get('searches') or 0) > 0 and (r['inn'] not in res or score(r) >= score(res[r['inn']])):
            res[r['inn']] = r

GROUP = {
    'Действует': 'Действует',
    'В процессе ликвидации': 'Закрывается',
    'Банкротство': 'Закрывается',
    'Предстоящее исключение': 'Закрывается',
    'Ликвидирована': 'Не существует',
    'Исключена ФНС': 'Не существует',
    'Реорганизована': 'Не существует',
    'ИП прекратил деятельность': 'Не существует',
    'Не установлено': 'Не найдено',
    'Ещё не проверялась': 'Не проверялась',
}
ORDER = ['Действует', 'Закрывается', 'Не существует', 'Не найдено', 'Не проверялась']
FILL = {'Действует': 'E2F0D9', 'Закрывается': 'FFF2CC', 'Не существует': 'F8D7DA',
        'Не найдено': 'EDEDED', 'Не проверялась': 'FFFFFF'}

# объединяем дубли по ИНН
rows = {}
for x in src:
    inn = x['ИНН']
    n = x['N п/п']
    file_ = '1' if x['_src'].startswith('файл 1') else '2'
    if 'Статус права' in x:
        m = {'Действует': 'член СРО', 'Прекращено': 'прекращено',
             'Приостановлено': 'приостановлено', '': 'нет данных'}.get(x['Статус права'], x['Статус права'])
    else:
        d = x.get('Дата прекращения членства', '')
        m = f'прекращено {d}' if d else 'член СРО'
    where = f'файл {file_}, № {n}'
    if inn in rows:
        rows[inn]['where'] += '; ' + where
        rows[inn]['sro'] += '; ' + m
        continue
    rows[inn] = dict(
        inn=inn, ogrn=x['ОГРН/ОГРНИП'],
        name=x['Сокращенное наименование'] or x['Полное наименование'],
        phone=x.get('Контактные телефоны', ''),
        addr=x.get('юр адрес') or x.get('Адрес места нахождения юридического лица', ''),
        head=x.get('ФИО') or x.get('Фамилия, имя, отчество (при наличии) для ИП', ''),
        sro=m, where=where,
    )

missing = [i for i in rows if i not in res]
recs = []
for inn, r in rows.items():
    s = res.get(inn, {})
    st = s.get('status') or 'Ещё не проверялась'
    if st not in GROUP:
        print('!! неизвестный статус', inn, st); st = 'Не установлено'
    email = '; '.join(e.strip().lower() for e in re.split(r'[;,\s]+', s.get('email', '') or '') if '@' in e)
    det = s.get('status_detail', '') or ''
    recheck = ''
    if st not in ('Ещё не проверялась', 'Не установлено'):
        if re.search(r'перепровер|не подтвержд|противореч|косвенн|реорганиз|не указал|не указана страниц|источник[^.;]{0,30}не указ', det, re.I) or not s.get('source'):
            recheck = 'да'
    recs.append(dict(r, group=GROUP[st], status=st, date=s.get('status_date', ''), recheck=recheck,
                     email=email, cur=s.get('current_name', ''), detail=s.get('status_detail', ''),
                     source=s.get('source', '')))
recs.sort(key=lambda r: (ORDER.index(r['group']), r['email'] == '', r['name'].lower()))

HEAD = ['№', 'Итог', 'Статус в ЕГРЮЛ / ЕГРИП', 'Дата события', 'Электронная почта',
        'Наименование (из файла)', 'Текущее наименование', 'ИНН', 'ОГРН / ОГРНИП',
        'Телефон (из файла)', 'Адрес (из файла)', 'Руководитель (из файла)',
        'Членство в СРО (по файлу)', 'Где в исходных файлах', 'Что сказано в источнике',
        'Источник', 'Перепроверить', 'Проверить самому']
WID = [6, 15, 24, 13, 32, 38, 26, 13, 16, 20, 50, 38, 20, 20, 60, 45, 14, 16]
F = Font(name='Arial', size=10)
FB = Font(name='Arial', size=10, bold=True)
FL = Font(name='Arial', size=10, color='0563C1', underline='single')
thin = Side(style='thin', color='BFBFBF')
BRD = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')

def sheet(ws, items):
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='404040')
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = BRD
    for i, r in enumerate(items, 1):
        link = f'https://www.rusprofile.ru/search?query={r["inn"]}'
        ws.append([i, r['group'], r['status'], r['date'], r['email'], r['name'], r['cur'], r['inn'],
                   r['ogrn'], r['phone'], r['addr'], r['head'], r['sro'], r['where'], r['detail'],
                   r['source'], r['recheck'], 'Rusprofile'])
        row = ws.max_row
        fill = PatternFill('solid', fgColor=FILL[r['group']])
        for c in ws[row]:
            c.font = F; c.alignment = WRAP; c.border = BRD
        for col in (2, 3, 4):
            ws.cell(row, col).fill = fill
        ws.cell(row, 2).font = FB
        lc = ws.cell(row, 18); lc.hyperlink = link; lc.font = FL
        if r['recheck']:
            ws.cell(row, 17).font = Font(name='Arial', size=10, bold=True, color='C00000')
        for col in (8, 9):
            ws.cell(row, col).number_format = '@'
    for i, w in enumerate(WID, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'F2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(HEAD))}{max(ws.max_row, 2)}'
    ws.row_dimensions[1].height = 32

wb = Workbook()
sv = wb.active; sv.title = 'Сводка'
wa = wb.create_sheet('Все компании')
sheet(wa, recs)
names = {'Действует': 'Действуют', 'Закрывается': 'Закрываются',
         'Не существует': 'Не существуют', 'Не найдено': 'Не найдены',
         'Не проверялась': 'Ещё не проверены'}
for g in ORDER:
    sheet(wb.create_sheet(names[g]), [r for r in recs if r['group'] == g])

# --- Сводка: формулы по листу «Все компании»
N = wa.max_row
A = f"'Все компании'!$B$2:$B${N}"
S = f"'Все компании'!$C$2:$C${N}"
E = f"'Все компании'!$E$2:$E${N}"
sv['A1'] = 'Строительные компании Ростовской области: статус и электронная почта'
sv['A1'].font = Font(name='Arial', size=14, bold=True)
sv['A2'] = ('Исходные файлы: register-1-500.xlsx и register-501-988.xlsx (реестр Союза «Строители Ростовской '
            'области»). Статус и почта найдены поиском по ИНН в открытых источниках (Rusprofile, Checko, '
            'List-org, ЗаЧестныйБизнес, РБК Компании, Audit-it и др.), проверка — октябрь 2026.')
sv['A2'].font = F; sv['A2'].alignment = WRAP
sv.merge_cells('A2:E2'); sv.row_dimensions[2].height = 54

hdr = ['Итог', 'Что это значит', 'Компаний', 'Из них с почтой', 'Доля от всех']
for i, h in enumerate(hdr, 1):
    c = sv.cell(4, i, h); c.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='404040'); c.border = BRD; c.alignment = WRAP
MEAN = {
    'Действует': 'Компания работает, в реестре налоговой значится действующей',
    'Закрывается': 'Ещё в реестре, но идёт ликвидация, банкротство или налоговая готовит исключение',
    'Не существует': 'Ликвидирована, исключена налоговой как недействующая, присоединена к другой или ИП закрыт',
    'Не найдено': 'Искали, но поиск не нашёл сведений по этому ИНН — проверить вручную по ссылке в последнем столбце',
    'Не проверялась': 'До компании ещё не дошла очередь: лимит поиска — 200 запросов за один ход',
}
r0 = 5
for k, g in enumerate(ORDER):
    r = r0 + k
    sv.cell(r, 1, g).font = FB
    sv.cell(r, 1).fill = PatternFill('solid', fgColor=FILL[g])
    sv.cell(r, 2, MEAN[g]).font = F
    sv.cell(r, 3, f'=COUNTIF({A},A{r})')
    sv.cell(r, 4, f'=COUNTIFS({A},A{r},{E},"?*")')
    sv.cell(r, 5, f'=IF($C${r0+len(ORDER)}=0,0,C{r}/$C${r0+len(ORDER)})')
    sv.cell(r, 5).number_format = '0.0%'
    for col in range(1, 6):
        sv.cell(r, col).border = BRD; sv.cell(r, col).alignment = WRAP
        if col > 2: sv.cell(r, col).font = F
rt = r0 + len(ORDER)
sv.cell(rt, 1, 'Всего').font = FB
sv.cell(rt, 3, f'=SUM(C{r0}:C{rt-1})').font = FB
sv.cell(rt, 4, f'=SUM(D{r0}:D{rt-1})').font = FB
sv.cell(rt, 5, f'=IF(C{rt}=0,0,C{rt}/C{rt})').font = FB
sv.cell(rt, 5).number_format = '0.0%'
for col in range(1, 6): sv.cell(rt, col).border = BRD

rd = rt + 2
sv.cell(rd, 1, 'Подробно по статусам').font = Font(name='Arial', size=11, bold=True)
for i, h in enumerate(['Статус в ЕГРЮЛ / ЕГРИП', 'Итог', 'Компаний', 'Из них с почтой'], 1):
    c = sv.cell(rd + 1, i, h); c.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='404040'); c.border = BRD
for k, st in enumerate(GROUP):
    r = rd + 2 + k
    sv.cell(r, 1, st); sv.cell(r, 2, GROUP[st])
    sv.cell(r, 2).fill = PatternFill('solid', fgColor=FILL[GROUP[st]])
    sv.cell(r, 3, f'=COUNTIF({S},A{r})')
    sv.cell(r, 4, f'=COUNTIFS({S},A{r},{E},"?*")')
    for col in range(1, 5):
        sv.cell(r, col).border = BRD; sv.cell(r, col).font = F

rn = rd + 2 + len(GROUP) + 1
notes = [
    'Как читать:',
    '• «Итог» — главное: действует компания или её уже нет. «Членство в СРО» — отдельная вещь: компания могла выйти из СРО и продолжать работать.',
    '• Почта — та, что опубликована на сайтах-справочниках (в основном из ЕГРЮЛ и госзакупок). У многих компаний почта не публикуется, тогда ячейка пустая.',
    '• В листах компании отсортированы: сначала с почтой, затем по названию. Фильтр стоит на каждом листе.',
    '• «Не найдено» не значит «не существует»: поисковик просто не нашёл страницу этой компании. Ссылка «Rusprofile» в последнем столбце открывает её карточку по ИНН.',
    '• «Перепроверить = да» — статус взят из одного источника без даты, косвенно (по выручке) или компания в процессе реорганизации. Такие строки стоит открыть по ссылке.',
    '• «Дата события» — дата ликвидации, исключения, начала банкротства или ликвидации, если источник её назвал.',
    '• Перед рассылкой проверьте почту выборочно: справочники обновляются с задержкой, адрес мог смениться.',
]
for k, t in enumerate(notes):
    c = sv.cell(rn + k, 1, t); c.font = FB if k == 0 else F
    sv.merge_cells(start_row=rn + k, start_column=1, end_row=rn + k, end_column=5)
    c.alignment = WRAP; sv.row_dimensions[rn + k].height = 30 if k else 16
for col, w in zip('ABCDE', [26, 70, 12, 16, 13]):
    sv.column_dimensions[col].width = w

wb.save(OUT)
print('saved', OUT, 'rows', len(recs), 'missing results', len(missing))
import collections
print(collections.Counter(r['group'] for r in recs), sum(1 for r in recs if r['email']))
