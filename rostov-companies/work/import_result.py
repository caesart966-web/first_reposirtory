"""Результат Checko, присланный пользователем, -> out/checko.json -> полная таблица.

Пакет на компьютере пользователя собирает свою таблицу, но в ней только то,
что было в пакете на момент запуска. Здесь его данные Checko переносятся
в рабочую папку, и таблица собирается заново текущим build.py.

    python3 -I import_result.py checko_results.csv            лучше: это прямой разбор ответов Checko
    python3 -I import_result.py Строители_РО_статус_и_почты.xlsx   если прислана только таблица

Из таблицы берутся строки, где «Откуда статус» = Checko, плюс строки поиска,
к пояснению которых дописан ответ Checko («не найдено»): почта в таблице
уже объединена с поиском — это не страшно, build.py объединит её снова.
"""
import csv
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
STATUSES = {'Действует', 'В процессе ликвидации', 'Банкротство', 'Предстоящее исключение', 'Ликвидирована',
            'Исключена ФНС', 'Реорганизована', 'ИП прекратил деятельность', 'Не установлено'}


def record(inn, status, date, email, phones, site, name, detail):
    return {'inn': inn, 'status': status, 'status_date': date or '', 'status_detail': detail or '',
            'email': email or '', 'current_name': name or '', 'phones': phones or '', 'site': site or '',
            'source': f'Checko API (ЕГРЮЛ/ЕГРИП), https://checko.ru/search?query={inn}',
            'searches': 1, 'via': 'checko'}


def from_csv(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        rows = list(csv.reader(f, delimiter=';'))
    head, out = rows[0], []
    ix = {h: i for i, h in enumerate(head)}
    for r in rows[1:]:
        if not r or not r[ix['ИНН']].strip():
            continue
        g = lambda h: r[ix[h]].strip() if ix.get(h) is not None and ix[h] < len(r) else ''  # noqa: E731
        out.append(record(g('ИНН'), g('Статус'), g('Дата'), g('Почта'), g('Телефоны'), g('Сайт'),
                          g('Наименование'), g('Пояснение')))
    return out


def from_xlsx(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True)['Все компании']
    rows = list(ws.iter_rows(values_only=True))
    ix = {h: i for i, h in enumerate(rows[0])}
    out = []
    for r in rows[1:]:
        g = lambda h: str(r[ix[h]] or '').strip() if h in ix else ''  # noqa: E731
        inn, detail = g('ИНН'), g('Что сказано в источнике')
        if not inn:
            continue
        if g('Откуда статус') == 'Checko':
            out.append(record(inn, g('Статус в ЕГРЮЛ / ЕГРИП'), g('Дата события'), g('Электронная почта'),
                              g('Телефоны (Checko)'), g('Сайт (Checko)'), g('Текущее наименование'), detail))
        elif ' | ' in detail and 'Checko' in detail.split(' | ', 1)[1]:
            # статус из поиска, а Checko ответил без статуса (не найдено и т.п.)
            out.append(record(inn, 'Не установлено', '', '', g('Телефоны (Checko)'), g('Сайт (Checko)'), '',
                              detail.split(' | ', 1)[1]))
    return out


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    path = sys.argv[1]
    recs = from_csv(path) if path.lower().endswith('.csv') else from_xlsx(path)
    os.chdir(HERE)
    known = {x['ИНН'] for x in json.load(open('all.json', encoding='utf-8'))}
    bad = [r['inn'] for r in recs if r['inn'] not in known or r['status'] not in STATUSES]
    if bad:
        sys.exit(f'Не узнаю {len(bad)} строк (чужой ИНН или статус), первые: {bad[:5]} — файл не из этого пакета?')
    old = {}
    if os.path.exists('out/checko.json'):
        old = {r['inn']: r for r in json.load(open('out/checko.json', encoding='utf-8'))}
    old.update({r['inn']: r for r in recs})
    with open('out/checko.json', 'w', encoding='utf-8') as f:
        json.dump(list(old.values()), f, ensure_ascii=False, indent=1)
    print(f'Из «{os.path.basename(path)}»: {len(recs)} компаний Checko; всего в out/checko.json: {len(old)}')


if __name__ == '__main__':
    main()
