"""Выгрузки: заключение и ответ — в DOCX (из markdown), реестр компании — в XLSX."""
from __future__ import annotations

import re
from pathlib import Path


def safe_name(s: str, limit: int = 80) -> str:
    """Имя файла для Windows: без запрещённых знаков, кавычек-ёлочек и лишних пробелов."""
    s = re.sub(r'[\\/:*?"<>|«»„“”]', "", s or "")
    s = re.sub(r"\s+", " ", s).strip(" .")
    return s[:limit] or "без названия"


def report_file(reports_dir: Path, company_name: str, kind: str, number: str | None, date_iso: str, ext: str) -> Path:
    """«ОДО-отчёты/ООО КИТ/ООО КИТ — договор 0145… — заключение — 2026-09-24.docx». Папка компании создаётся."""
    comp = safe_name(company_name, 60)
    parts = [comp]
    if number:
        parts.append(f"договор {safe_name(number, 40)}")
    parts += [kind, date_iso]
    folder = reports_dir / comp
    folder.mkdir(parents=True, exist_ok=True)
    return folder / (" — ".join(parts) + "." + ext)


def md_to_docx(md: str, path: Path, title: str | None = None):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|\s*-", lines[i + 1]):
            header = [c.strip() for c in ln.strip("|").split("|")]
            rows = []
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
            tbl.style = "Table Grid"
            for j, h in enumerate(header):
                tbl.cell(0, j).text = _plain(h)
            for r, row in enumerate(rows, 1):
                for j, c in enumerate(row[:len(header)]):
                    tbl.cell(r, j).text = _plain(c)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            doc.add_heading(_plain(m.group(2)), level=min(len(m.group(1)), 4))
        elif re.match(r"^\s*[-*]\s+", ln):
            doc.add_paragraph(_plain(re.sub(r"^\s*[-*]\s+", "", ln)), style="List Bullet")
        elif re.match(r"^\s*\d+\.\s+", ln):
            doc.add_paragraph(_plain(re.sub(r"^\s*\d+\.\s+", "", ln)), style="List Number")
        elif ln.startswith(">"):
            p = doc.add_paragraph(_plain(ln.lstrip("> ")))
            p.runs[0].italic = True
        elif ln.strip():
            _para_with_bold(doc, ln)
        i += 1
    doc.save(str(path))


def _plain(s: str) -> str:
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"\*(.+?)\*", r"\1", s)
    return s.replace("`", "")


def _para_with_bold(doc, ln: str):
    p = doc.add_paragraph()
    for part in re.split(r"(\*\*.+?\*\*)", ln):
        if part.startswith("**") and part.endswith("**"):
            p.add_run(_plain(part)).bold = True
        elif part:
            p.add_run(_plain(part))


def register_xlsx(company: dict, rows: list[dict], totals: dict, path: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Реестр"
    ws.append([f"Реестр договоров: {company['name']}" + (f" (ИНН {company['inn']})" if company.get("inn") else "")])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    head = ["№ договора", "Дата", "Заказчик", "Цена, ₽", "Принято по актам, ₽", "Остаток, ₽", "Способ", "Членство", "В ОДО", "В совокупный размер, ₽", "Решение", "Кто решил", "Уведомление в СРО"]
    ws.append(head)
    for c in ws[3]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append([r["number"], r["date"], r["customer"], r["price"], r["executed"], r["remaining"], r["procurement_ru"], r["membership_ru"], r["odo_ru"], r["odo_amount"], r["decision_ru"], r["decided_by"], r["notified_ru"]])
    ws.append([])
    ws.append(["Итого в совокупный размер по КФ ОДО", None, None, None, None, None, None, None, None, totals["odo_sum"]])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.cell(ws.max_row, 10).font = Font(bold=True)
    ws.append(["Договоров всего / в ОДО", totals["n_all"], totals["n_odo"]])
    ws.append(["Расчёт: остаток = цена − принято по актам (ч. 7 ст. 55.13 ГрК РФ); в совокупный размер входят конкурентные договоры, по которым требуется членство (ч. 3 ст. 55.8, ст. 60.1 ГрК РФ)."])
    for col, w in zip("ABCDEFGHIJKLM", (26, 12, 40, 16, 18, 16, 14, 12, 8, 20, 22, 14, 16)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=4, max_row=3 + len(rows)):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for idx in (3, 4, 5, 9):
            row[idx].number_format = "#,##0.00"
    wb.save(str(path))
