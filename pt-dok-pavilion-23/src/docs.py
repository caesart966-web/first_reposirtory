"""Текстовые документы в рамке по ГОСТ Р 21.101 (PDF, reportlab): спецификация (ГОСТ 21.110), ведомость трубопроводов,
кабельный журнал, ВОР, расчёты. Плюс XLSX спецификации."""
import json, os, math
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.pdfbase.pdfmetrics import stringWidth

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'out')
from ru import ru
from stamp import NAMES, DATE
from meta import clean_pdf, clean_xlsx
import zipfile, re
_FONTS = os.path.join(HERE, '..', 'fonts')
pdfmetrics.registerFont(TTFont('F', os.path.join(_FONTS, 'osifont-lgpl3fe.ttf')))
pdfmetrics.registerFont(TTFont('FB', os.path.join(_FONTS, 'osifont-lgpl3fe.ttf')))
pdfmetrics.registerFont(TTFont('FI', os.path.join(_FONTS, 'osifont-lgpl3fe.ttf')))


class RuCanvas(canvas.Canvas):
    def drawString(self, x, y, text, *a, **kw):
        return super().drawString(x, y, ru(text), *a, **kw)

    def drawCentredString(self, x, y, text, *a, **kw):
        return super().drawCentredString(x, y, ru(text), *a, **kw)

    def drawRightString(self, x, y, text, *a, **kw):
        return super().drawRightString(x, y, ru(text), *a, **kw)
SHIFR = 'ОПР-01/24/2024-ДП1-ПТ'
OBJ = '«Современный коммерческий оптово-продовольственный рынок (ОПР) по адресу: Донецкая область, г. Макеевка, Горняцкий район»'
BLD = 'Док-павильон №23, №23.2 Рыба-Мясо. Пожаротушение (ВПВ, АУПТ, АУПП ВРУ)'
ORG = 'ООО «Технология»'

def wrap(text, font, size, width):
    out = []
    for para in ru(str(text)).split('\n'):
        words = para.split(' ')
        line = ''
        for w in words:
            t = (line + ' ' + w).strip()
            if stringWidth(t, font, size) <= width:
                line = t
            else:
                if line:
                    out.append(line)
                while stringWidth(w, font, size) > width and len(w) > 1:
                    # жёсткий перенос длинного слова
                    k = len(w)
                    while k > 1 and stringWidth(w[:k], font, size) > width:
                        k -= 1
                    out.append(w[:k]); w = w[k:]
                line = w
        out.append(line)
    return out

class Doc:
    def __init__(self, path, code, title, pagesize='A3L'):
        self.ps = landscape(A3) if pagesize == 'A3L' else A4
        self.W, self.H = self.ps[0] / mm, self.ps[1] / mm
        self.c = RuCanvas(path, pagesize=self.ps)
        self.path = path
        self.code, self.title = code, title
        self.page = 0
        self.pages_total = None
    def frame(self, first):
        c = self.c
        self.page += 1
        c.setLineWidth(0.7)
        c.rect(20 * mm, 5 * mm, (self.W - 25) * mm, (self.H - 10) * mm)
        # боковые графы
        y = 5
        for s, h in [('Инв. № подл.', 25), ('Подп. и дата', 35), ('Взам. инв. №', 25)]:
            c.setLineWidth(0.5); c.rect(8 * mm, y * mm, 12 * mm, h * mm); c.line(13 * mm, y * mm, 13 * mm, (y + h) * mm)
            c.saveState(); c.translate(11.6 * mm, (y + h / 2) * mm); c.rotate(90); c.setFont('F', 6.5); c.drawCentredString(0, 0, s); c.restoreState()
            y += h
        x = (self.W - 5 - 185)
        if first:
            self.stamp3(x, 5)
            return 5 + 55
        self.stamp6(x, 5)
        return 5 + 15
    def stamp3(self, x, y):
        c = self.c
        def L(x1, y1, x2, y2, w=0.5):
            c.setLineWidth(w); c.line((x + x1) * mm, (y + y1) * mm, (x + x2) * mm, (y + y2) * mm)
        c.setLineWidth(0.7); c.rect(x * mm, y * mm, 185 * mm, 55 * mm)
        for i in range(1, 11):
            L(0, i * 5, 65, i * 5, 0.5 if i in (5, 6) else 0.25)
        for cx in (10, 20, 30, 40, 55):
            L(cx, 25, cx, 55)
        for cx in (20, 40, 55):
            L(cx, 0, cx, 25)
        L(65, 0, 65, 55, 0.7); L(65, 40, 185, 40, 0.7); L(65, 30, 185, 30); L(65, 15, 185, 15, 0.7)
        L(135, 0, 135, 30, 0.7); L(135, 25, 185, 25); L(150, 15, 150, 30); L(165, 15, 165, 30)
        c.setFont('F', 6.5)
        for s, cx, w in [('Изм.', 0, 10), ('Кол.уч', 10, 10), ('Лист', 20, 10), ('№ док.', 30, 10), ('Подп.', 40, 15), ('Дата', 55, 10)]:
            c.drawCentredString((x + cx + w / 2) * mm, (y + 26.5) * mm, s)
        for i, s in enumerate(['Разраб.', 'Пров.', '', 'Н.контр.', 'ГИП']):
            c.drawString((x + 1) * mm, (y + 21.5 - i * 5) * mm, s)
            if NAMES.get(s):
                c.drawString((x + 21) * mm, (y + 21.5 - i * 5) * mm, NAMES[s])
                c.drawCentredString((x + 60) * mm, (y + 21.5 - i * 5) * mm, DATE)
        c.setFont('FB', 13); c.drawCentredString((x + 125) * mm, (y + 45.5) * mm, self.code)
        c.setFont('F', 6.3)
        for i, l in enumerate(wrap(OBJ, 'F', 6.3, 116 * mm)):
            c.drawCentredString((x + 125) * mm, (y + 36.5 - i * 2.6) * mm, l)
        c.setFont('F', 7)
        for i, l in enumerate(wrap(BLD, 'F', 7, 66 * mm)):
            c.drawCentredString((x + 100) * mm, (y + 25 - i * 3.0) * mm, l)
        c.setFont('F', 6.5)
        for s, cx, w in [('Стадия', 135, 15), ('Лист', 150, 15), ('Листов', 165, 20)]:
            c.drawCentredString((x + cx + w / 2) * mm, (y + 26.5) * mm, s)
        c.setFont('F', 10)
        c.drawCentredString((x + 142.5) * mm, (y + 18.5) * mm, 'Р')
        c.drawCentredString((x + 157.5) * mm, (y + 18.5) * mm, str(self.page))
        if self.pages_total:
            c.drawCentredString((x + 175) * mm, (y + 18.5) * mm, str(self.pages_total))
        c.setFont('F', 8)
        lines = wrap(self.title, 'F', 8, 66 * mm)
        for i, l in enumerate(lines):
            c.drawCentredString((x + 100) * mm, (y + 7.5 + (len(lines) - 1) * 1.6 - i * 3.3) * mm, l)
        c.setFont('F', 10); c.drawCentredString((x + 160) * mm, (y + 6.5) * mm, ORG)
    def stamp6(self, x, y):
        c = self.c
        c.setLineWidth(0.7); c.rect(x * mm, y * mm, 185 * mm, 15 * mm)
        for cx in (10, 20, 30, 40, 55, 65, 175):
            c.setLineWidth(0.5); c.line((x + cx) * mm, y * mm, (x + cx) * mm, (y + 15) * mm)
        c.line((x + 175) * mm, (y + 8) * mm, (x + 185) * mm, (y + 8) * mm)
        for i in (1, 2):
            c.setLineWidth(0.25); c.line(x * mm, (y + i * 5) * mm, (x + 65) * mm, (y + i * 5) * mm)
        c.setFont('F', 6.5)
        for s, cx, w in [('Изм.', 0, 10), ('Кол.уч', 10, 10), ('Лист', 20, 10), ('№ док.', 30, 10), ('Подп.', 40, 15), ('Дата', 55, 10)]:
            c.drawCentredString((x + cx + w / 2) * mm, (y + 1.5) * mm, s)
        c.setFont('FB', 12); c.drawCentredString((x + 120) * mm, (y + 5.5) * mm, self.code)
        c.setFont('F', 6.5); c.drawCentredString((x + 180) * mm, (y + 10.5) * mm, 'Лист')
        c.setFont('F', 10); c.drawCentredString((x + 180) * mm, (y + 2.5) * mm, str(self.page))
    def _ensure(self):
        if getattr(self, 'y', None) is None:
            self.y_bottom = self.frame(self.page == 0)
            self.y = self.H - 8
    def new_page(self):
        if getattr(self, 'y', None) is not None:
            self.c.showPage()
        self.y_bottom = self.frame(self.page == 0)
        self.y = self.H - 8
    def table(self, cols, rows, header_h=14, font=7.5, min_h=8, title=None, group_key='group'):
        """cols: [(заголовок, ширина мм, выравнивание L/C/R)]; rows: списки ячеек или dict(group=...) — строка-подзаголовок на всю ширину"""
        c = self.c
        self._ensure()
        x0 = 20
        W = sum(w for _, w, _ in cols)
        def header():
            c.setLineWidth(0.6); c.rect(x0 * mm, (self.y - header_h) * mm, W * mm, header_h * mm)
            xx = x0
            for t, w, _ in cols:
                c.line(xx * mm, self.y * mm, xx * mm, (self.y - header_h) * mm)
                ls = wrap(t, 'F', font - 0.5, (w - 2) * mm)
                for i, l in enumerate(ls):
                    c.setFont('F', font - 0.5)
                    c.drawCentredString((xx + w / 2) * mm, (self.y - header_h / 2 + (len(ls) - 1) * 1.5 - i * 3.2 - 1) * mm, l)
                xx += w
            self.y -= header_h
        if self.y - header_h - min_h < self.y_bottom:
            self.new_page()
        header()
        for r in rows:
            if isinstance(r, dict):
                ls = wrap(r[group_key], 'FB', font, (W - 2) * mm)
                rh = max(min_h, len(ls) * font * 0.42 + 3)
            else:
                wrapped = [wrap(cell, 'F', font, (w - 2) * mm) for (t, w, _), cell in zip(cols, r)]
                rh = max(min_h, max(len(wl) for wl in wrapped) * font * 0.42 + 3)
            if self.y - rh < self.y_bottom:
                self.new_page(); header()
            y = self.y
            if isinstance(r, dict):
                c.setLineWidth(0.25); c.line(x0 * mm, y * mm, x0 * mm, (y - rh) * mm); c.line((x0 + W) * mm, y * mm, (x0 + W) * mm, (y - rh) * mm)
                c.setFont('FB', font)
                for k, l in enumerate(ls):
                    c.drawString((x0 + 1.5) * mm, (y - 3.5 - k * font * 0.42) * mm, l)
            else:
                xx = x0
                for i, (t, w, al) in enumerate(cols):
                    c.setLineWidth(0.25); c.line(xx * mm, y * mm, xx * mm, (y - rh) * mm)
                    c.setFont('F', font)
                    for k, l in enumerate(wrapped[i]):
                        yy = (y - 3.5 - k * font * 0.42) * mm
                        if al == 'C':
                            c.drawCentredString((xx + w / 2) * mm, yy, l)
                        elif al == 'R':
                            c.drawRightString((xx + w - 1) * mm, yy, l)
                        else:
                            c.drawString((xx + 1) * mm, yy, l)
                    xx += w
                c.line(xx * mm, y * mm, xx * mm, (y - rh) * mm)
            self.y -= rh
            c.line(x0 * mm, self.y * mm, (x0 + W) * mm, self.y * mm)
        self.y -= 4
        return self.y
    def text(self, paragraphs, font=9, width=None, start_first=True, line_gap=1.35):
        c = self.c
        self._ensure()
        width = width or (self.W - 30)
        for kind, txt in paragraphs:
            f, s = {'h1': ('FB', font + 3), 'h2': ('FB', font + 1), 'p': ('F', font), 'pre': ('F', font - 1), 'i': ('FI', font)}[kind]
            ls = wrap(txt, f, s, width * mm)
            if kind in ('h1', 'h2'):
                self.y -= 2
            for l in ls:
                if self.y - s * 0.3528 * line_gap < self.y_bottom + 2:
                    self.new_page()
                c.setFont(f, s); c.drawString(25 * mm, (self.y - s * 0.3528) * mm, l)
                self.y -= s * 0.3528 * line_gap
            self.y -= 1.2
        return self.y
    def save(self):
        self.c.save()
        clean_pdf(self.path, f'{self.code}. {self.title}')

def two_pass(build_fn):
    """Сначала считаем страницы, затем пишем «Листов»."""
    d = build_fn(None)
    n = d.page
    d2 = build_fn(n)
    return d2

def spec_pdf():
    S = json.load(open(os.path.join(HERE, 'spec.json')))
    cols = [('Поз.', 14, 'C'), ('Наименование и техническая характеристика', 140, 'L'), ('Тип, марка, обозначение документа, опросного листа', 58, 'L'),
            ('Код продукции', 22, 'C'), ('Поставщик', 42, 'L'), ('Ед. изм.', 14, 'C'), ('Кол.', 18, 'C'), ('Масса ед., кг', 18, 'C'), ('Примечание', 69, 'L')]
    rows = []
    for r in S['rows']:
        if 'group' in r:
            rows.append(dict(group=r['group']))
        else:
            rows.append([str(r['pos']), r['name'], r['type'] + ((', ' + r['code']) if r['code'] else ''), '', r['maker'], r['unit'],
                         str(r['qty']), str(r['mass']) if r['mass'] != '' else '', r['note']])
    def build(total):
        d = Doc(os.path.join(OUT, f'{SHIFR.replace("/", "_")}.С_спецификация.pdf'), SHIFR + '.С', 'Спецификация оборудования, изделий и материалов')
        d.pages_total = total
        d.table(cols, rows)
        d.save()
        return d
    two_pass(build)

def spec_xlsx():
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    S = json.load(open(os.path.join(HERE, 'spec.json')))
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Спецификация'
    hdr = ['Поз.', 'Наименование и техническая характеристика', 'Тип, марка, обозначение документа, опросного листа', 'Код продукции',
           'Поставщик', 'Ед. изм.', 'Кол.', 'Масса ед., кг', 'Примечание']
    def put(ws_, vals):
        ws_.append([ru(v) if isinstance(v, str) else v for v in vals])
    ws.append([SHIFR + '.С — Спецификация оборудования, изделий и материалов (ГОСТ 21.110-2013)'])
    ws.append([])
    ws.append(hdr)
    thin = Side(style='thin'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, w in enumerate([7, 70, 38, 12, 26, 8, 9, 10, 30]):
        ws.column_dimensions[chr(65 + i)].width = w
    for cell in ws[3]:
        cell.font = Font(bold=True); cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center'); cell.border = B
    for r in S['rows']:
        if 'group' in r:
            ws.append(['', r['group']])
            ws[ws.max_row][1].font = Font(bold=True)
            ws[ws.max_row][1].fill = PatternFill('solid', fgColor='EEEEEE')
        else:
            put(ws, [r['pos'], r['name'], r['type'] + ((', ' + r['code']) if r['code'] else ''), '', r['maker'], r['unit'], r['qty'],
                       r['mass'] if r['mass'] != '' else None, r['note']])
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical='top'); cell.border = B
    ws['A1'].font = Font(bold=True, size=12)
    ws.freeze_panes = 'A4'
    # лист: трубы по системам
    B_ = json.load(open(os.path.join(HERE, 'bom.json')))
    ws2 = wb.create_sheet('Трубопроводы')
    ws2.append(['Система', 'Ду', 'Длина, м', 'С запасом 5 %, м'])
    for k, v in B_['length'].items():
        s, dn = k.split('|'); ws2.append([s, int(dn), v, math.ceil(v * 1.05)])
    ws3 = wb.create_sheet('Кабели')
    ws3.append(['№', 'Откуда', 'Куда', 'Марка', 'Длина, м', 'Прим.'])
    for c in json.load(open(os.path.join(HERE, 'cables.json'))):
        put(ws3, [c['n'], c['frm'], c['to'], c['mark'], c['L'], c['note']])
    for w_, col in zip([8, 45, 55, 30, 10, 30], 'ABCDEF'):
        ws3.column_dimensions[col].width = w_
    clean_xlsx(wb, SHIFR + '.С. Спецификация оборудования, изделий и материалов')
    xp = os.path.join(OUT, f'{SHIFR.replace("/", "_")}.С_спецификация.xlsx')
    wb.save(xp)
    # в docProps/app.xml библиотека пишет своё имя — убираем
    with zipfile.ZipFile(xp) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(xp + '.tmp', 'w', zipfile.ZIP_DEFLATED) as z:
        for i, data in items:
            if i.filename == 'docProps/app.xml':
                data = re.sub(rb'<Application>.*?</Application>|<AppVersion>.*?</AppVersion>', b'', data)
            z.writestr(i, data)
    os.replace(xp + '.tmp', xp)

def vt_pdf():
    """Ведомость трубопроводов, арматуры и оборудования (по системам)."""
    S = json.load(open(os.path.join(HERE, 'spec.json')))
    B_ = json.load(open(os.path.join(HERE, 'bom.json')))
    P = S['POS']
    rows = [dict(group='Трубопроводы (длины без запаса)')]
    names = {'В2': 'В2 — ВПВ', 'В21': 'В21 — АУПТ, секция 1', 'В22': 'В22 — АУПТ, секция 2'}
    from spec_data import DN_SZ, PIPE_NAME
    for k, v in B_['length'].items():
        s, dn = k.split('|'); dn = int(dn)
        rows.append([names[s], f'Ду{dn} ({DN_SZ[dn]})', PIPE_NAME[dn][1], 'м', f'{v:.1f}', ''])
    rows.append(dict(group='Арматура'))
    Q = {r['pos']: r['qty'] for r in S['rows'] if r.get('pos')}
    arm = [('Насосная: вводы, обвод, Н1, Н2, после УУ', 'Затвор дисковый Ду150 с контролем положения', P['bfly150'], Q[P['bfly150']]),
           ('Насосная: опробование, отвод ВПВ, ГМ-80', 'Затвор дисковый Ду100 с контролем положения', P['bfly100'], Q[P['bfly100']]),
           ('Насосная: вводы, обвод, Н1, Н2', 'Клапан обратный Ду150', P['chk150'], Q[P['chk150']]),
           ('Насосная: патрубки ГМ-80', 'Клапан обратный Ду100', P['chk100'], Q[P['chk100']]),
           ('Насосная, жокей-насос', 'Кран шаровой Ду32 / клапан обратный Ду32', f"{P['ball32']}, {P['chk32']}", f"{Q[P['ball32']]} / {Q[P['chk32']]}"),
           ('Насосная, дренаж', 'Кран шаровой Ду50', P['ball50'], Q[P['ball50']]),
           ('Насосная, дренажные насосы', 'Кран шаровой Ду40 / клапан обратный Ду40', f"{P['ball40']}, {P['chk40']}", f"{Q[P['ball40']]} / {Q[P['chk40']]}"),
           ('Кольцо ВПВ (подвал)', 'Затвор Ду100 секционирующий', P['ringv'], Q[P['ringv']]),
           ('Шкафы ПК', 'Клапан пожарный КПЛМ 65-1', P['kpk'], B_['pk'] * 2),
           ('Концы питающих В21, В22', 'Узел промывки (кран Ду50)', P['flush'], 4), ('Диктующие точки', 'Узел опробования (кран Ду25)', P['test'], 2),
           ('Верхние точки В22/В21', 'Воздухоотводчик Ду15', P['vent'], 4)]
    for a, b, p, q in arm:
        rows.append([a, b, f'поз. {p}', 'шт.', str(q), ''])
    rows.append(dict(group='Оборудование'))
    PM = json.load(open(os.path.join(HERE, 'calc_final.json')))['pump']
    NS = 'Насосная (подвал)'
    eq = [(NS, f"Насосы Н1, Н2 (Q={PM['Q_m3h']} м³/ч, H={PM['H']} м, {PM['N']:g} кВт)", P['pump'], 2), (NS, 'Жокей-насос Н3', P['jockey'], 1),
          (NS, 'Бак мембранный 100 л', P['tank'], 1), (NS, 'Узел управления УУ-С150 (УУ-1, УУ-2)', P['uu'], 2),
          (NS, 'Насосы дренажные ДН1, ДН2, шкаф управления ШУД', f"{P['drain_pump']}, {P['shud']}", '2 / 1'),
          (NS, 'Шкафы ШКП-30 / ШКП-4', f"{P['shkp30']}, {P['shkp4']}", '2 / 1'),
          (NS, 'Поток-3Н, С2000-КДЛ-2И, РИП-24 в шкафу автоматики', f"{P['potok']}, {P['kdl']}, {P['rip']}", '1/1/1'),
          ('Подвал, 1 и 2 этажи', 'Шкафы пожарные ШПК-320 с 2 ПК', P['cab'], B_['pk']), ('Секция 1', 'Оросители СВО0-РНо0,42', P['spr_dn'], B_['spr']['DN']),
          ('Секция 2', 'Оросители СВО0-РВо0,42', P['spr_up'], B_['spr']['UP']), ('Пом. 06', 'МПП (тушение по объёму)', P['mpp'], 4),
          ('У пом. 06', 'С2000-АСПТ, РИП-24', f"{P['aspt']}, {P['rip_p']}", '1 / 1')]
    for a, b, p, q in eq:
        rows.append([a, b, f'поз. {p}', 'шт.', str(q), ''])
    cols = [('Система / место', 75, 'L'), ('Наименование', 135, 'L'), ('Обозначение / поз. спецификации', 70, 'L'), ('Ед.', 15, 'C'), ('Кол.', 25, 'C'), ('Прим.', 75, 'L')]
    def build(total):
        d = Doc(os.path.join(OUT, f'{SHIFR.replace("/", "_")}.ВТ_ведомость_трубопроводов_арматуры_оборудования.pdf'), SHIFR + '.ВТ', 'Ведомость трубопроводов, арматуры и оборудования')
        d.pages_total = total
        d.table(cols, rows); d.save(); return d
    two_pass(build)

def kj_pdf():
    C = json.load(open(os.path.join(HERE, 'cables.json')))
    rows = [[c['n'], c['frm'], c['to'], c['mark'], str(c['L']), c['note']] for c in C]
    cols = [('Обозначение кабеля', 25, 'C'), ('Начало', 95, 'L'), ('Конец', 105, 'L'), ('Марка, число и сечение жил', 60, 'L'), ('Длина, м', 20, 'C'), ('Примечание', 90, 'L')]
    def build(total):
        d = Doc(os.path.join(OUT, f'{SHIFR.replace("/", "_")}.КЖ_кабельный_журнал.pdf'), SHIFR + '.КЖ', 'Кабельный журнал')
        d.pages_total = total
        d.table(cols, rows); d.save(); return d
    two_pass(build)

def vor_pdf():
    S = json.load(open(os.path.join(HERE, 'spec.json')))
    rows = []
    n = 0
    work = {
        'шт.': 'Монтаж', 'компл.': 'Монтаж', 'м': 'Прокладка', 'кг': 'Нанесение'
    }
    for r in S['rows']:
        if 'group' in r:
            rows.append(dict(group=r['group'])); continue
        n += 1
        w = work.get(r['unit'], 'Монтаж')
        nm = r['name']
        if r['unit'] == 'м' and nm.startswith('Труба'):
            act = 'Прокладка трубопровода: ' + nm.replace('Труба ', '').lower() + ', с креплением и окраской'
        elif r['unit'] == 'м' and nm.startswith('Кабель'):
            act = 'Прокладка кабеля ' + nm.replace('Кабель ', '')
        elif r['unit'] == 'кг':
            act = 'Огнезащита/окраска: ' + nm
        else:
            act = 'Монтаж: ' + nm
        rows.append([str(n), act, r['unit'], str(r['qty']), f"поз. {r['pos']}"])
    extra = [('Промывка трубопроводов АУПТ и ВПВ', 'система', '3'), ('Гидравлические испытания трубопроводов', 'система', '3'),
             ('Пусконаладочные работы насосной станции, автоматики, АУПП', 'комплекс', '1'), ('Комплексное опробование и сдача систем', 'комплекс', '1')]
    rows.append(dict(group='Испытания и наладка'))
    for a, u, q in extra:
        n += 1; rows.append([str(n), a, u, q, ''])
    cols = [('№', 14, 'C'), ('Наименование работ', 250, 'L'), ('Ед. изм.', 25, 'C'), ('Кол.', 30, 'C'), ('Примечание', 76, 'L')]
    def build(total):
        d = Doc(os.path.join(OUT, f'{SHIFR.replace("/", "_")}.ВОР_ведомость_объемов_работ.pdf'), SHIFR + '.ВОР', 'Ведомость объёмов работ')
        d.pages_total = total
        d.table(cols, rows); d.save(); return d
    two_pass(build)

if __name__ == '__main__':
    spec_pdf(); spec_xlsx(); vt_pdf(); kj_pdf(); vor_pdf()
    print('ok')
